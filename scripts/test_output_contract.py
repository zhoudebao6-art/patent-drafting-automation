import copy,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
import batch_tool as b
import inspect_docx as inspector
import output_contract as contract
import flat_delivery as flat
import parallel_batch as queue
import test_parallel_batch as fixtures


class OutputContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def report(self,description='保存对应记录。',claim='1.一种方法，其特征在于，按权利要求限定的对象处理。',title='一种处理系统'):
        doc=Document();doc.add_paragraph('说明书摘要');doc.add_paragraph('内容摘要。')
        doc.add_paragraph('权利要求书');doc.add_paragraph(claim)
        doc.add_paragraph(title);doc.add_paragraph('技术领域');doc.add_paragraph('本发明涉及记录处理。')
        doc.add_paragraph('具体实施方式')
        p=doc.add_paragraph()
        # Terms deliberately cross runs; no single-run word search is valid.
        for ch in description:p.add_run(ch)
        path=self.root/'scratch.docx';doc.save(path)
        r=inspector.inspect(path);inspector.precheck(r,'draft')
        return r

    def test_claim_terms_remain_valid_and_exact_title_is_used(self):
        report=self.report()
        result=contract.check(report,'一种处理系统')
        self.assertEqual(result['delivery_filename'],'一种处理系统.docx')
        self.assertFalse(report['document_contract_evidence']['description_word_hits'])

    def test_global_terms_and_description_only_terms_cross_runs_are_rejected(self):
        for text in ['从而保留记录。','记录一一对应。','记录一一相应。','记录一一进行校验。','根据权利要求处理。','其特征在于保存记录。']:
            with self.subTest(text=text),self.assertRaises(b.BatchError):
                contract.check(self.report(text),'一种处理系统')
        for text in ['1.处理数据，从而保存结果。','1.将记录一一匹配。']:
            with self.subTest(claim=text),self.assertRaises(b.BatchError):
                contract.check(self.report(claim=text),'一种处理系统')

    def test_yiyi_nonmodifier_requires_position_bound_independent_finding(self):
        r=self.report('统一一套记录格式。')
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统')
        hit=r['document_contract_evidence']['global_word_hits'][0]
        f=dict(hit,rule_id='WORDING-04',classification='not_modifier',status='closed',reason='统一与一套属于两个词，非逐个配对修饰。')
        contract.check(r,'一种处理系统',[f])
        f['offset']+=1
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',[f])

    def test_known_yiyi_modifier_cannot_be_waived_as_nonmodifier(self):
        r=self.report('记录一一对应。');hit=r['document_contract_evidence']['global_word_hits'][0]
        f=dict(hit,rule_id='WORDING-04',classification='not_modifier',status='closed',reason='假豁免')
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',[f])

    def test_title_mismatch_missing_title_and_changed_suffix_rejected(self):
        r=self.report()
        for name in ['处理系统','一种处理系统_001','一种处理系统修改版']:
            with self.subTest(name=name),self.assertRaises(b.BatchError):contract.check(r,name)
        r['document_contract_evidence']['title']=None
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统')

    def test_exact_filename_does_not_sanitize_truncate_or_add_case_suffix(self):
        title='一种'+('长名称'*40)+'方法'
        m={'policy':{'output_contract':contract.MODE},'projects':[{'id':'001','title':title}]}
        self.assertEqual(flat.names(m)['001'],title+'.docx')
        for bad in ['名称/分支',' CON','CON','名称.']:
            with self.subTest(bad=bad),self.assertRaises(b.BatchError):contract.filename(bad)
        m['projects'].append({'id':'002','title':title})
        with self.assertRaises(b.BatchError):flat.names(m)

    def test_legacy_collision_suffix_remains_unchanged(self):
        m={'projects':[{'id':'001','title':'同名'},{'id':'002','title':'同名'}]}
        self.assertEqual(flat.names(m),{'001':'同名_001.docx','002':'同名_002.docx'})

    def test_disabled_gate_does_not_require_new_evidence_on_old_batches(self):
        with patch('prepare_handoff.verify_prepared',side_effect=AssertionError('must not inspect')):
            self.assertIsNone(contract.verify({'policy':{}},{},{},{}))

    def test_policy_cannot_enable_contract_without_frozen_preparation(self):
        p=copy.deepcopy(fixtures.POLICY);p['output_contract']=contract.MODE
        with self.assertRaises(b.BatchError):queue.validate_policy(p)

    def test_prepared_gate_uses_existing_report_without_reparsing(self):
        report=self.report();report_file=self.root/'inspection.json';b.atomic(report_file,report)
        handoff_file=self.root/'handoff.json';b.atomic(handoff_file,{'input_inspection':{'inspection':str(report_file)}})
        attempt={'prepared_handoff':{'path':str(handoff_file)},'input_artifact_sha256':report['sha256']}
        verified={'handoff':{'input_inspection':{'inspection':str(report_file)}},'input_report':report}
        with patch('prepare_handoff.verify_prepared',return_value=verified) as seal,patch.object(inspector,'inspect',side_effect=AssertionError('duplicate parse')):
            result=contract.verify({'policy':{'output_contract':contract.MODE}},{'title':'一种处理系统'},attempt,{'findings':[]})
        seal.assert_called_once()
        self.assertEqual(result['artifact_sha256'],report['sha256'])

    def test_global_terms_in_header_are_not_missed(self):
        self.report()
        path=self.root/'scratch.docx';doc=Document(path)
        doc.sections[0].header.paragraphs[0].text='从而保留记录'
        doc.save(path);report=inspector.inspect(path);inspector.precheck(report,'draft')
        self.assertTrue(any(h['part'].startswith('word/header') for h in report['document_contract_evidence']['global_word_hits']))
        with self.assertRaises(b.BatchError):contract.check(report,'一种处理系统')


class QueueContractTests(unittest.TestCase):
    """Real queue/cache/file transitions; model dispatch and renders are fixtures."""
    reserve=fixtures.ParallelTests.reserve
    receipt=fixtures.ParallelTests.receipt

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.src=self.root/'input';self.src.mkdir()
        doc=Document();doc.add_paragraph('源技术内容');doc.save(self.src/'source.docx')
        policy=copy.deepcopy(fixtures.POLICY)
        policy.update(worker_preparation='frozen_stage_packet_v1',output_contract=contract.MODE,
                      delivery={'folder_prefix':'专利撰写完成','layout':'flat_final_files','timing':'after_each_acceptance'})
        policy_path=self.root/'policy.json';b.atomic(policy_path,policy)
        selection={'projects':[{'title':'一种处理系统','source_files':['source.docx'],'primary_source':'source.docx'}]}
        with patch.object(queue,'POLICY',policy_path):
            self.manifest=queue.init(self.src,self.root/'work',selection,self.root/'desktop')
        self.batch=self.manifest.parent

    def review_attempt(self,text='保存对应记录。',title='一种处理系统'):
        draft=self.reserve(1);receipt=self.receipt(draft)
        doc=Document();doc.add_paragraph('权利要求书')
        doc.add_paragraph('1.一种处理系统，其特征在于，依照权利要求所限定的记录。')
        doc.add_paragraph(title);doc.add_paragraph('技术领域');doc.add_paragraph('本发明涉及记录处理。')
        doc.add_paragraph('具体实施方式');doc.add_paragraph(text)
        doc.save(receipt['artifact']);receipt['artifact_sha256']=b.sha(receipt['artifact'])
        queue.record(self.manifest,draft['case_id'],draft['attempt']['id'],receipt)
        review=self.reserve(1)
        return review,self.receipt(review,'pass')

    def test_final_acceptance_and_exact_name_delivery_reuse_sealed_report(self):
        review,receipt=self.review_attempt()
        with patch.object(inspector,'inspect',side_effect=AssertionError('duplicate parse')):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
            dest=Path(queue.deliver(self.manifest))
        self.assertEqual([p.name for p in dest.glob('*.docx')],['一种处理系统.docx'])
        case=b.read(self.manifest)['projects'][0]
        self.assertEqual(case['status'],'delivered')
        self.assertEqual(case['receipts'][-1]['document_contract']['document_title'],'一种处理系统')

    def test_wrong_word_or_title_cannot_be_accepted_or_delivered(self):
        review,receipt=self.review_attempt('根据权利要求处理，从而保存记录。')
        with self.assertRaisesRegex(b.BatchError,'wording contract'):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
        self.assertNotEqual(b.read(self.manifest)['projects'][0]['status'],'accepted')
        with self.assertRaises(b.BatchError):queue.deliver(self.manifest)
        self.assertFalse(list((self.root/'desktop').rglob('*.docx')))

    def test_wrong_formal_title_cannot_be_accepted(self):
        review,receipt=self.review_attempt(title='另一个标题')
        with self.assertRaisesRegex(b.BatchError,'title must exactly match'):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)

    def test_modified_prepared_report_cannot_forge_acceptance(self):
        review,receipt=self.review_attempt()
        prepared=b.read(review['attempt']['prepared_handoff']['path'])
        report_path=Path(prepared['input_inspection']['inspection'])
        report=b.read(report_path);report['document_contract_evidence']['title']['text']='伪造'
        b.atomic(report_path,report)
        with self.assertRaises(b.BatchError):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)

    def test_delivery_rechecks_contract_even_after_acceptance(self):
        review,receipt=self.review_attempt()
        queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
        m=b.read(self.manifest);m['projects'][0]['receipts'][-1]['document_contract']['document_title']='伪造'
        b.atomic(self.manifest,m)
        with self.assertRaises(b.BatchError):queue.deliver(self.manifest)
        self.assertFalse(list((self.root/'desktop').rglob('*.docx')))


if __name__=='__main__':unittest.main()
