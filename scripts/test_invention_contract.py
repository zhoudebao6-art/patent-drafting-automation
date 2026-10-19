"""Observable claim projection, tampering and same-artifact acceptance boundaries."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml import parse_xml
import authoring_toolkit as t
import batch_tool as batch
import inspect_docx as scan
import invention_contract as inv
import output_contract as final
from test_authoring_toolkit import minimal_spec

TITLE='一种记录处理方法'
CLAIMS=['1.'+TITLE+'，其特征在于，包括：获取记录，保存所述记录；',
        '仅在所述记录有效时输出所述记录，否则不参与更新。',
        '2.根据权利要求1所述的'+TITLE+'，其特征在于，对所述记录核验，保留所述记录的原标识。']


class InventionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.path=self.root/'fixture.docx'

    def doc(self,claims=None,solution=None,purpose=None,transition=None):
        claims=claims or CLAIMS
        if solution is None:
            solution=[inv.claim_head(p)[2] if p[0].isdigit() else p for p in claims]
        lines=['说明书摘要','摘要。','权利要求书',*claims,TITLE,'技术领域','本发明涉及记录处理。',
               '背景技术','背景。','发明内容',purpose or inv.purpose(TITLE),transition or inv.TRANSITION,
               *solution,inv.BENEFITS,'1.保留记录。','附图说明','具体实施方式','实施内容。']
        doc=Document()
        for line in lines:doc.add_paragraph(line)
        return doc

    def report(self,doc=None):
        (doc or self.doc()).save(self.path)
        report=scan.inspect(self.path);scan.precheck(report,'draft');return report

    def accept(self,report,findings=()):
        return final.check(report,TITLE,findings,strict_invention=True)

    def test_valid_multiline_claims_keep_every_word(self):
        r=self.report();self.assertFalse(r['invention_evidence']['issues'])
        self.assertEqual(len(r['invention_evidence']['mappings']),3)
        self.assertEqual(self.accept(r)['invention_contract'],inv.MODE)

    def test_exact_phrase_required_in_purpose_not_elsewhere(self):
        doc=self.doc(purpose='本发明的目的在于提供'+TITLE+'，以解决记录缺失问题。')
        doc.add_paragraph(inv.purpose(TITLE))
        with self.assertRaisesRegex(batch.BatchError,'fixed_purpose'):self.accept(self.report(doc))
        with self.assertRaisesRegex(batch.BatchError,'fixed_transition'):
            self.accept(self.report(self.doc(transition='为实现上述目的，本发明采用如下技术方案：')))

    def test_deleted_shusuo_negations_conditions_and_substitutions_rejected(self):
        original=[inv.claim_head(CLAIMS[0])[2],CLAIMS[1],inv.claim_head(CLAIMS[2])[2]]
        for old,new in [('所述',''),('仅',''),('否则','同时'),('不参与','参与'),('原标识','新标识')]:
            with self.subTest(old=old),self.assertRaisesRegex(batch.BatchError,'solution_text_changed'):
                self.accept(self.report(self.doc(solution=[p.replace(old,new) for p in original])))

    def test_missing_reordered_split_or_extra_paragraphs_rejected(self):
        original=[inv.claim_head(CLAIMS[0])[2],CLAIMS[1],inv.claim_head(CLAIMS[2])[2]]
        for changed in [original[:-1],list(reversed(original)),original+['附加功能。'],[original[0]+original[1],original[2]]]:
            with self.subTest(changed=changed),self.assertRaises(batch.BatchError):self.accept(self.report(self.doc(solution=changed)))

    def test_cross_run_text_and_deleted_revision_read_as_accepted_view(self):
        doc=self.doc()
        p=next(p for p in doc.paragraphs if p.text==inv.purpose(TITLE))
        p.clear()
        for c in inv.purpose(TITLE):p.add_run(c)
        p._p.append(parse_xml('<w:del xmlns:w="'+t.NS['w']+'"><w:r><w:delText>旧句。</w:delText></w:r></w:del>'))
        self.accept(self.report(doc))

    def test_same_plain_text_with_different_math_structure_is_rejected(self):
        doc=self.doc()
        a=doc.paragraphs[4];b=next(p for p in doc.paragraphs if p.text.startswith('仅在') and p is not a)
        # Select the actual second occurrence, independent of python-docx wrappers.
        candidates=[p for p in doc.paragraphs if p.text==CLAIMS[1]];a,b=candidates
        m=t.NS['m']
        a._p.append(parse_xml('<m:oMath xmlns:m="'+m+'"><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:oMath>'))
        b._p.append(parse_xml('<m:oMath xmlns:m="'+m+'"><m:r><m:t>xi</m:t></m:r></m:oMath>'))
        with self.assertRaisesRegex(batch.BatchError,'solution_math_changed'):self.accept(self.report(doc))

    def test_located_reference_rewrite_is_not_blanket_waiver(self):
        claims=CLAIMS+['3.一种记录系统，其特征在于，所述模块执行权利要求1所述的方法。']
        solution=[inv.claim_head(p)[2] if p[0].isdigit() else p for p in CLAIMS]+['一种记录系统，所述模块获取记录，仅在记录有效时输出记录，否则不参与更新。']
        r=self.report(self.doc(claims=claims,solution=solution));ev=r['invention_evidence']
        self.assertEqual(len(ev['review_candidates']),1)
        with self.assertRaisesRegex(batch.BatchError,'located independent review'):self.accept(r)
        finding=dict(ev['review_candidates'][0],rule_id='STYLE-02',classification='claim_solution_mapping',status='closed',artifact_sha256=r['sha256'],reason='当前独立检查按被引用方法逐项比对直接叙述。')
        self.accept(r,[finding])
        for key in ['artifact_sha256','source_text_sha256','target_text_sha256','source_paragraph']:
            with self.subTest(key=key),self.assertRaises(batch.BatchError):self.accept(r,[dict(finding,**{key:'wrong'})])
        bad=self.report(self.doc(solution=[inv.claim_head(CLAIMS[0])[2].replace('所述',''),CLAIMS[1],inv.claim_head(CLAIMS[2])[2]]))
        with self.assertRaisesRegex(batch.BatchError,'solution_text_changed'):self.accept(bad,[finding])

    def test_missing_or_wrong_hash_evidence_rejected_and_legacy_is_unchanged(self):
        r=self.report()
        for key,value in [('schema_version',True),('complete',False),('artifact_sha256','f'*64),('mappings',[])]:
            copy=deepcopy(r);copy['invention_evidence'][key]=value
            with self.subTest(key=key),self.assertRaises(batch.BatchError):self.accept(copy)
        r.pop('invention_evidence');self.assertNotIn('invention_contract',final.check(r,TITLE))
        with self.assertRaises(batch.BatchError):self.accept(r)

    def test_nested_textless_object_cannot_hide_additional_solution_content(self):
        from PIL import Image
        image=self.root/'tiny.png';Image.new('RGB',(4,4),'black').save(image)
        doc=self.doc();table=doc.add_table(rows=1,cols=1)
        table.cell(0,0).paragraphs[0].add_run().add_picture(str(image))
        intro=next(p for p in doc.paragraphs if p.text==inv.BENEFITS)
        intro._p.addprevious(table._tbl)
        with self.assertRaisesRegex(batch.BatchError,'nested_solution_requires_explicit_mapping'):self.accept(self.report(doc))

    def test_new_policy_reuses_sealed_report_without_zip_reparse(self):
        r=self.report();verified={'handoff':{'input_inspection':{'inspection':'sealed.json'}},'input_report':r}
        with patch('prepare_handoff.verify_prepared',return_value=verified),patch.object(scan,'inspect',side_effect=AssertionError('duplicate parse')):
            answer=final.verify({'policy':{'output_contract':final.MODE,'invention_contract':inv.MODE}},{'title':TITLE},{'input_artifact_sha256':r['sha256']},{})
            self.assertEqual(answer['invention_contract'],inv.MODE)

    def test_helper_changes_invalidate_existing_cache(self):
        self.doc().save(self.path)
        first=scan.run_inspection(self.path,self.root/'cache','draft',True,True)
        self.assertEqual(first['cache_status'],'miss')
        self.assertEqual(scan.run_inspection(self.path,self.root/'cache','draft',True,True)['cache_status'],'hit')
        with patch.object(scan,'extractor_dependencies',return_value={'invention_contract.py':'b'*64}):
            self.assertEqual(scan.run_inspection(self.path,self.root/'cache','draft',True,True)['cache_status'],'miss')

    def test_policy_flag_requires_frozen_contract(self):
        import parallel_batch as queue
        from test_parallel_batch import POLICY
        with self.assertRaisesRegex(batch.BatchError,'Invention contract requires'):queue.validate_policy(dict(POLICY,invention_contract=inv.MODE))
        with self.assertRaisesRegex(batch.BatchError,'Invalid invention'):queue.validate_policy(dict(POLICY,invention_contract='unknown'))

    def spec(self):
        spec=minimal_spec();spec['invention_from_claims']=True
        spec['sections'][2]['blocks']=[dict(id='claim'+str(i),type='paragraph',role='body',text=p) for i,p in enumerate(CLAIMS)]
        spec['sections'][3]['blocks'] += [dict(id='invention',type='paragraph',role='heading',text='发明内容'),
                                        dict(id='benefit_intro',type='paragraph',role='body',text=inv.BENEFITS),
                                        dict(id='benefit',type='paragraph',role='benefit',text='保留记录。')]
        return spec

    def test_generator_uses_current_claims_without_mutating_input(self):
        spec=self.spec();original=deepcopy(spec)
        context=t.validate(self.root,spec);t.render_document(context,self.path);t.verify_document(context,self.path)
        self.assertEqual(spec,original)
        r=scan.inspect(self.path);scan.precheck(r,'draft');inv.validate_evidence(r,[])
        spec['sections'][2]['blocks'][1]['text']=CLAIMS[1].replace('原','原')+'保留原时间。'
        changed=t.validate(self.root,spec)
        self.assertTrue(any(p.endswith('保留原时间。') for key,p in changed['block_texts'].items() if key.startswith('iv_')))

    def test_claim_intro_requires_chinese_comma_with_or_without_mirroring(self):
        for mirror,punctuation in ((True,'：'),(False,':'),(False,','),(True,'；')):
            with self.subTest(mirror=mirror,punctuation=punctuation):
                spec=self.spec()
                if not mirror:spec.pop('invention_from_claims')
                block=spec['sections'][2]['blocks'][0]
                text=block.pop('text').replace('其特征在于，','其特征在于'+punctuation)
                block['runs']=[{'type':'text','text':c} for c in text]
                original=deepcopy(spec)
                with self.assertRaisesRegex(t.AuthoringError,'Chinese comma'):
                    t.validate(self.root,spec)
                self.assertEqual(spec,original)

        spec=self.spec();context=t.validate(self.root,spec)
        self.assertIn('其特征在于，包括：',context['block_texts']['claim0'])

    def test_manual_colon_claim_is_rejected_by_existing_final_gate(self):
        doc=self.doc()
        for p in doc.paragraphs:
            if p.text==CLAIMS[0]:
                p.text=p.text.replace('其特征在于，','其特征在于：')
        report=self.report(doc)
        self.assertTrue(any(i['code']=='claim_head_unresolved' for i in report['invention_evidence']['issues']))
        with self.assertRaises(batch.BatchError):self.accept(report)

    def test_wrong_intro_in_continuation_cannot_use_mapping_waiver(self):
        claims=list(CLAIMS)
        claims[1]='其特征在于：'+claims[1]
        report=self.report(self.doc(claims=claims))
        evidence=report['invention_evidence']
        self.assertTrue(any(i['code']=='claim_intro_punctuation' for i in evidence['issues']))
        findings=[dict(row,rule_id='STYLE-02',classification='claim_solution_mapping',status='closed',
                       artifact_sha256=report['sha256'],reason='Located technical correspondence')
                  for row in evidence['review_candidates']]
        with self.assertRaisesRegex(ValueError,'claim_intro_punctuation'):
            inv.validate_evidence(report,findings)

    def test_cross_run_headers_and_multiple_independent_categories(self):
        spec=self.spec()
        for b in spec['sections'][2]['blocks']:
            b['runs']=[{'type':'text','text':c} for c in b.pop('text')]
        for n,name in [(3,'一种记录系统'),(4,'一种存储介质')]:
            spec['sections'][2]['blocks'].append(dict(id='category'+str(n),type='paragraph',role='body',text=f'{n}.{name}，其特征在于，保存所述记录。'))
        context=t.validate(self.root,spec)
        values=[v for k,v in context['block_texts'].items() if k.startswith('iv_')]
        self.assertIn('优选的，对所述记录核验，保留所述记录的原标识。',values)
        self.assertIn('一种记录系统，保存所述记录。',values)
        self.assertIn('一种存储介质，保存所述记录。',values)

    def test_generator_refuses_existing_plan_and_unsupported_objects(self):
        spec=self.spec();spec['sections'][3]['blocks'].insert(-2,dict(id='old_plan',type='paragraph',role='body',text='旧方案。'))
        with self.assertRaisesRegex(t.AuthoringError,'Leave the solution region empty'):t.validate(self.root,spec)
        spec=self.spec();spec['sections'][2]['blocks'][1]=dict(id='math',type='formula',formula_id='f1')
        with self.assertRaises(t.AuthoringError):t.validate(self.root,spec)


if __name__=='__main__':unittest.main()
