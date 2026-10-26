"""Five uploaded templates through source-neutral synthetic structural fixtures."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
import authoring_toolkit as toolkit
import batch_tool as batch
import claims_contract as cc
import inspect_docx as scan
import invention_contract as inv
import output_contract as final
from test_authoring_toolkit import minimal_spec, MATH


TITLE = '记录处理方法'


def claims_for(family='method', title=TITLE):
    themes = {
        'method': {1: '一种记录处理方法'},
        'system': {1: '一种记录处理系统'},
        'method_system': {1: '一种记录处理方法', 10: '一种记录处理系统'},
        'system_method': {1: '一种记录处理系统', 10: '一种记录处理方法'},
        'three_independent': {1: '一种记录处理方法', 9: '一种记录处理系统', 10: '一种计算机可读存储介质'},
    }[family]
    rows = []
    for n in range(1, 11):
        theme = themes.get(n, '根据权利要求1所述的'+themes[1][2:])
        head = f'{n}.{theme}，其特征在于，'
        prefix = head+'获取记录数据，核对所述记录数据，保存所述记录数据，读取所述记录数据；'
        bounds = (150, 300) if n in themes else cc.dependent_band(title, n)
        target = sum(bounds)//2
        # Padding is solely a unit-test boundary fixture, NEVER drafting advice.
        text = prefix+'存'*(target-cc.visible_count(prefix)-1)+'。'
        rows.append(dict(number=n, text=text, paragraphs=[n]))
    return rows


class ClaimTemplatesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def report(self, rows=None, formula=False, bad_purpose=False, rule_version=76):
        rows = claims_for() if rows is None else rows
        doc = Document()
        for line in ['说明书摘要', '摘要内容。', '权利要求书']:
            doc.add_paragraph(line)
        for row in rows:
            for line in row['text'].split('\n'):
                p = doc.add_paragraph(line)
            if formula and row['number']==1:
                p._p.append(parse_xml(MATH))
        for line in [TITLE, '技术领域', '记录处理。', '背景技术', '原有记录。', '发明内容',
                     inv.purpose(('一种' if bad_purpose else '')+TITLE), inv.TRANSITION]:
            doc.add_paragraph(line)
        for row in rows:
            first, *rest = row['text'].split('\n')
            # Invalid claim punctuation is deliberately retained in the claims
            # fixture; keep its separate description structurally constructible.
            projected=inv.claim_head(first.replace('其特征在于：','其特征在于，'))[2]
            if rule_version >= 79 and cc.standard_storage_medium(row,rows):
                # Description prose must not mechanically copy claim numbers.
                # This fixture exercises only mechanical checks, not source or
                # semantic claim-to-description acceptance.
                projected=projected[:projected.index('实现权利要求')]+'实现记录处理方法。'
            doc.add_paragraph(projected)
            for line in rest:doc.add_paragraph(line)
        for line in [inv.BENEFITS, '1.记录保存。', '附图说明', '图1为本发明的流程图。', '具体实施方式', '记录内容。']:
            doc.add_paragraph(line)
        path = self.root/'sample.docx'
        doc.save(path)
        result = scan.inspect(path)
        scan.precheck(result, 'draft', rule_version=rule_version)
        return result, path

    def spec(self, rows=None):
        spec = minimal_spec()
        spec.update(project_title=TITLE, protected_names=[])
        spec['sections'][0]['blocks'][0]['text']='本发明涉及记录技术领域，公开了'+TITLE+'。'
        spec['sections'][2]['blocks'] = [
            dict(id=f'c{r["number"]}', type='paragraph', role='body', text=r['text']) for r in (rows or claims_for())]
        spec['sections'][3]['blocks'][0]['text']=TITLE
        return spec

    def test_five_families_and_illegal_numbering_or_category_combinations(self):
        for family in ['method', 'system', 'method_system', 'system_method', 'three_independent']:
            with self.subTest(family=family):
                ev = cc.evaluate(claims_for(family), TITLE)
                self.assertEqual(ev['family'], family)
                self.assertEqual(ev['issues'], [])
                self.assertIn('source', ev['semantic_review'])
        rows=claims_for('three_independent')
        rows[9]['text']=rows[9]['text'].replace('一种计算机可读存储介质','一种记录处理系统')
        self.assertEqual(cc.evaluate(rows,TITLE)['family'],'three_independent')
        for mutate in [
            lambda rows: rows.pop(),
            lambda rows: rows[8].update(text=rows[8]['text'].replace('根据权利要求1所述的', '一种')),
            lambda rows: rows[0].update(text=rows[0]['text'].replace('方法', '装置')),
            lambda rows: rows[-1].update(number=9),
        ]:
            rows=claims_for();mutate(rows)
            self.assertTrue(cc.evaluate(rows, TITLE)['issues'])

    def test_all_open_intervals_and_counting_convention(self):
        self.assertEqual(cc.visible_count(' 10. A B，\n数据。 '), 6)
        self.assertEqual(cc.dependent_band(TITLE, 2), cc.dependent_band(TITLE, 2))
        seen=set()
        for n in range(1, 90):seen.add(cc.dependent_band(TITLE, n))
        self.assertEqual(seen, set(cc.BANDS))
        for band in cc.BANDS:
            number=next(n for n in range(2,100) if cc.dependent_band(TITLE,n)==band)
            for count,okay in [(band[0],False),(band[0]+1,True),(band[1]-1,True),(band[1],False)]:
                prefix=f'{number}.根据权利要求1所述的记录处理方法，其特征在于，'
                row=dict(number=number,text=prefix+'存'*(count-cc.visible_count(prefix)-1)+'。',paragraphs=[1])
                ev=cc.evaluate([row], TITLE)
                self.assertEqual(not any(h['code']=='claim_length' for h in ev['issues']),okay)
        for count in (150,151,299,300):
            row=claims_for()[0];prefix='1.一种记录处理方法，其特征在于，'
            row['text']=prefix+'存'*(count-cc.visible_count(prefix)-1)+'。'
            self.assertEqual(any(h['code']=='claim_length' for h in cc.evaluate([row],TITLE)['issues']),count in (150,300))

    def test_single_stop_cross_paragraphs_and_decimal_not_sentence(self):
        for fragment,invalid in [('记录0.5；',False),('记录。\n继续；',True),('Node.js；',False),('.NET；',False)]:
            rows=claims_for();rows[0]['text']=rows[0]['text'].replace('获取记录数据，',fragment)
            ev=cc.evaluate(rows,TITLE)
            self.assertEqual(any(h['code']=='claim_terminal_stop' for h in ev['issues']),invalid)
        rows=claims_for();rows[0]['text']=rows[0]['text'].replace('保存','\n保存',1)
        report,_=self.report(rows)
        self.assertEqual(report['claim_template_evidence']['issues'],[])
        self.assertEqual(len(report['claims_detected'][0]['paragraphs']),2)
        self.assertFalse(report['claims_detected'][-1]['text'].endswith(TITLE))

    def test_prohibited_objects_and_vague_refs_not_computer_keyword(self):
        for fragment,code in [('该系统','claim_vague_reference'),('二者','claim_vague_reference'),('x=2','claim_text_formula')]:
            rows=claims_for();rows[0]['text']=rows[0]['text'].replace('获取记录数据',fragment)
            self.assertIn(code,[h['code'] for h in cc.evaluate(rows,TITLE)['issues']])
        ev=cc.evaluate(claims_for('three_independent'),TITLE)
        self.assertFalse(any(h['text']=='计算' for h in ev['review_candidates']))
        report,_=self.report(formula=True)
        self.assertIn('claim_math_object',[h['code'] for h in report['claim_template_evidence']['issues']])
        rows=claims_for();body=[dict(p=1,media=[{'type':'image'}])]
        self.assertTrue(cc.evaluate(rows,TITLE,body)['review_candidates'])
        self.assertEqual(cc.math_kind('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>α</m:t></m:r></m:oMath>'),'symbol')

    def test_old_medium_exception_and_proposal_scope(self):
        rows=claims_for('three_independent')
        rows[-1]['text']='10.一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至7任一项所述的方法。'
        report,path=self.report(rows)
        self.assertEqual(report['claims_detected'][-1]['antecedent_count_evidence']['required_valid_minimum'],3)
        self.assertIn('claim_length',[h['code'] for h in report['claim_template_evidence']['issues']])
        old=scan.inspect(path);scan.precheck(old,'draft',75)
        # The legacy helper retains its exact wording behavior.
        self.assertTrue(scan._standard_storage_medium_claim(rows[-1],list(range(1,11))))
        self.assertNotIn('claim_template_evidence',old)
        source=scan.inspect(path);scan.precheck(source,'proposal',76)
        self.assertNotIn('claim_template_evidence',source)

    def test_rule79_screenshot_wording_actual_method_ranges_and_old_versions(self):
        prefix='10.一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求'
        for reference in ('1至7任一项','1至8任一项','2至8任一项','1'):
            rows=claims_for('three_independent')
            rows[-1]['text']=prefix+reference+'所述的方法。'
            with self.subTest(reference=reference):
                report,_=self.report(rows,rule_version=79)
                medium=report['claims_detected'][-1]
                self.assertEqual(medium['text'],rows[-1]['text'])
                self.assertEqual(medium['antecedent_count_evidence']['required_valid_minimum'],2)
                self.assertIn('not_automatic_pass',medium['antecedent_count_evidence']['status'])
                self.assertEqual(report['claim_template_evidence']['issues'],[])
                self.assertIsNone(report['claim_template_evidence']['claims'][-1]['exclusive_bounds'])
                toolkit.validate(self.root,self.spec(rows),79)
                final.check(report,TITLE,strict_claim_templates=True,claim_rule_version=79)
                for version in (76,77,78):
                    old,_=self.report(rows,rule_version=version)
                    self.assertEqual(old['claims_detected'][-1]['antecedent_count_evidence']['required_valid_minimum'],3)
                    self.assertIn('claim_length',[h['code'] for h in old['claim_template_evidence']['issues']])

    def test_rule79_fixed_medium_rejects_changed_wording_and_wrong_references(self):
        standard='10.一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至8任一项所述的方法。'
        variants={
            'missing_storage':standard.replace('其上存储有计算机程序，',''),
            'extra_quota_text':standard.replace('方法。','方法，核对所述计算机程序。'),
            'wrong_punctuation':standard.replace('其特征在于，','其特征在于：'),
            'system_in_range':standard.replace('1至8','1至9'),
            'self':standard.replace('1至8','1至10'),
            'future':standard.replace('1至8','1至11'),
            'descending':standard.replace('1至8','8至1'),
            'singleton_range':standard.replace('1至8','1至1'),
            'leading_zero':standard.replace('1至8','01至8'),
        }
        for name,text in variants.items():
            rows=claims_for('three_independent');rows[-1]['text']=text
            with self.subTest(case=name):
                report,_=self.report(rows,rule_version=79)
                self.assertEqual(report['claims_detected'][-1]['antecedent_count_evidence']['required_valid_minimum'],3)
                self.assertIn('storage_medium_wording_or_method_range',[h['code'] for h in report['claim_template_evidence']['issues']])
                with self.assertRaises(batch.BatchError):
                    final.check(report,TITLE,[{'rule_id':'STYLE-08','status':'closed','reason':'No blanket waiver'}],strict_claim_templates=True,claim_rule_version=79)
        for mutation in ('missing','wrong_method_category','wrong_family'):
            rows=claims_for('three_independent');rows[-1]['text']=standard
            if mutation=='missing':rows.pop(3)
            elif mutation=='wrong_method_category':rows[3]['text']=rows[3]['text'].replace('记录处理方法','记录处理系统')
            else:rows[8]['text']=rows[8]['text'].replace('一种记录处理系统','根据权利要求1所述的记录处理方法')
            with self.subTest(case=mutation):
                self.assertFalse(cc.standard_storage_medium(rows[-1],rows))

    def test_rule79_other_third_categories_keep_general_limits_and_source_choice(self):
        for theme in ('一种记录处理设备','一种记录处理装置'):
            rows=claims_for('three_independent')
            rows[-1]['text']=rows[-1]['text'].replace('一种计算机可读存储介质',theme)
            ev=cc.evaluate(rows,TITLE,rule_version=79)
            with self.subTest(theme=theme):
                self.assertEqual(ev['issues'],[])
                self.assertEqual(ev['claims'][-1]['exclusive_bounds'],[150,300])
                rows[-1]['text']='10.'+theme+'，其特征在于，存储计算机程序。'
                report,_=self.report(rows,rule_version=79)
                self.assertEqual(report['claims_detected'][-1]['antecedent_count_evidence']['required_valid_minimum'],3)
                self.assertIn('claim_length',[h['code'] for h in report['claim_template_evidence']['issues']])

    def test_rule79_cache_build_and_sealed_gate_bind_current_wording_contract(self):
        rows=claims_for('three_independent')
        rows[-1]['text']='10.一种计算机可读存储介质，其上存储有计算机程序，其特征在于，所述计算机程序被处理器执行时实现权利要求1至8任一项所述的方法。'
        _,path=self.report(rows,rule_version=79)
        old=scan.run_inspection(path,self.root/'cache79',reuse=True,rule_version=78)
        new=scan.run_inspection(path,self.root/'cache79',reuse=True,rule_version=79)
        self.assertNotEqual(old['inspection'],new['inspection'])
        self.assertEqual(new['cache_status'],'miss')
        stale=json.loads(Path(old['inspection']).read_text(encoding='utf-8'))
        with self.assertRaises(batch.BatchError):
            final.check(stale,TITLE,strict_claim_templates=True,claim_rule_version=79)
        spec_path=self.root/'medium.json'
        spec_path.write_text(json.dumps(self.spec(rows),ensure_ascii=False),encoding='utf-8')
        with patch.object(scan,'tool_rule_version',return_value=79):
            built=toolkit.build(self.root,spec_path.name,'medium-build')
            report=json.loads(Path(built['inspection']['inspection']).read_text(encoding='utf-8'))
            cc.require_valid(report,TITLE,79)
            toolkit.verify_bundle(self.root,'medium-build')
            with patch.object(scan,'inspect',side_effect=AssertionError('No repeat DOCX parse')):
                final.check(report,TITLE,strict_claim_templates=True,claim_rule_version=79)

    def test_rule80_sealed_gate_reuses_marker_evidence_and_rejects_enabled_properties(self):
        report,_=self.report(rule_version=80)
        attempt={'input_artifact_sha256':report['sha256']}
        manifest={'rule_version':80,'policy':{'output_contract':final.MODE}}
        verified={'handoff':{'input_inspection':{'inspection':'sealed'}},'input_report':report}
        with patch('prepare_handoff.verify_prepared',return_value=verified),patch.object(scan,'inspect',side_effect=AssertionError('No repeat DOCX parse')):
            final.verify(manifest,{'title':TITLE},attempt,{'findings':[]})
            report['page_flow_evidence']['pagination_markers']['violations'].append(
                {'part':'word/document.xml','path':'paragraph[1]/pPr/keepNext'})
            with self.assertRaisesRegex(batch.BatchError,'pagination marker'):
                final.verify(manifest,{'title':TITLE},attempt,{'findings':[]})

    def test_rule81_plain_paper_sealed_gate_and_versioned_cache(self):
        report,path=self.report(rule_version=81)
        attempt={'input_artifact_sha256':report['sha256']}
        manifest={'rule_version':81,'policy':{'output_contract':final.MODE}}
        verified={'handoff':{'input_inspection':{'inspection':'sealed'}},'input_report':report}
        original=deepcopy(report['background_evidence'])
        with patch('prepare_handoff.verify_prepared',return_value=verified),patch.object(scan,'inspect',side_effect=AssertionError('No repeat DOCX parse')):
            final.verify(manifest,{'title':TITLE},attempt,{'findings':[]})
            for kind in ('missing','stale','incomplete','hit','candidate'):
                with self.subTest(kind=kind):
                    report['background_evidence']=deepcopy(original)
                    e=report['background_evidence']
                    if kind=='missing':report.pop('background_evidence')
                    elif kind=='stale':e['artifact_sha256']='0'*64
                    elif kind=='incomplete':e['extraction_complete']=False
                    elif kind=='hit':e['hits'].append({'part':'word/document.xml','path':'background[1]'})
                    else:e['review_candidates'].append({'part':'word/document.xml','paragraph':0,'path':'table[1]'})
                    with self.assertRaisesRegex(batch.BatchError,'plain-paper'):
                        final.verify(manifest,{'title':TITLE},attempt,{'findings':[]})
            candidate=report['background_evidence']['review_candidates'][0]
            finding=dict(candidate,rule_id='FORMAT-02',classification='plain_paper_reviewed',status='closed',artifact_sha256=report['sha256'],reason='Fixture: conditional fill is inactive; final page plain.')
            final.verify(manifest,{'title':TITLE},attempt,{'findings':[finding]})
            for key,value in [('artifact_sha256','0'*64),('path','table[2]'),('reason','')]:
                bad=dict(finding);bad[key]=value
                with self.subTest(key=key),self.assertRaises(batch.BatchError):
                    final.verify(manifest,{'title':TITLE},attempt,{'findings':[bad]})
            report['background_evidence']['hits'].append(candidate)
            with self.assertRaises(batch.BatchError):final.verify(manifest,{'title':TITLE},attempt,{'findings':[finding]})
        self.assertNotEqual(scan.cache_identity(path,'draft',False,79),scan.cache_identity(path,'draft',False,80))
        self.assertNotEqual(scan.cache_identity(path,'draft',False,80),scan.cache_identity(path,'draft',False,81))
        self.assertEqual(scan.cache_identity(path,'proposal',False,80),scan.cache_identity(path,'proposal',False,81))

    def test_title_prefix_exception_is_only_claims_and_exact_solution_slots(self):
        report,_=self.report()
        self.assertFalse(any(h['code']=='title_reference_prefix' for h in report['invention_evidence']['issues']))
        bad,_=self.report(bad_purpose=True)
        self.assertIn('title_reference_prefix',[h['code'] for h in bad['invention_evidence']['issues']])
        self.assertIsNotNone(inv.name_prefix_issue('一种一种'+TITLE+'，',TITLE,True))
        self.assertIsNotNone(inv.name_prefix_issue(inv.purpose('一种'+TITLE),TITLE))
        with self.assertRaisesRegex(toolkit.AuthoringError,'title_reference_prefix'):
            toolkit.validate(self.root,self.spec(),75)
        toolkit.validate(self.root,self.spec(),76)
        spec=self.spec();spec['invention_from_claims']=True
        spec['sections'][3]['blocks'] += [
            dict(id='iv',type='paragraph',role='heading',text='发明内容'),
            dict(id='be',type='paragraph',role='body',text=inv.BENEFITS)]
        spec['sections'][2]['blocks'].insert(1,dict(id='cont',type='paragraph',role='body',text='一种'+TITLE+'，保存记录。'))
        with self.assertRaisesRegex(ValueError,'title_reference_prefix'):
            inv.expand(spec,76)

    def test_generation_and_sealed_gate_use_same_evidence_without_reparse(self):
        spec=self.spec()
        spec['invention_from_claims']=True
        spec['sections'][3]['blocks'] += [
            dict(id='iv',type='paragraph',role='heading',text='发明内容'),
            dict(id='be',type='paragraph',role='body',text=inv.BENEFITS)]
        toolkit.validate(self.root,spec,76)
        rows=claims_for();rows[0]['text']='1.一种记录处理方法，其特征在于，保存记录。'
        with self.assertRaisesRegex(toolkit.AuthoringError,'claim_length'):
            toolkit.validate(self.root,self.spec(rows),76)
        report,_=self.report()
        with patch.object(scan,'inspect',side_effect=AssertionError('No second DOCX parse')):
            final.check(report,TITLE,strict_claim_templates=True)
        for mutation in ['missing','stale','changed','issues']:
            altered=deepcopy(report)
            if mutation=='missing':altered.pop('claim_template_evidence')
            elif mutation=='stale':altered['claim_template_evidence']['artifact_sha256']='0'*64
            elif mutation=='changed':altered['claims_detected'][0]['text']+='变更'
            else:altered['claim_template_evidence']['issues'].append({'code':'claim_length'})
            with self.subTest(mutation=mutation), self.assertRaises(batch.BatchError):
                final.check(altered,TITLE,[{'rule_id':'STYLE-01','status':'closed','reason':'no waiver'}],strict_claim_templates=True)

    def test_rule76_cache_separates_old_reports_and_reuses_current(self):
        _,path=self.report()
        one=scan.run_inspection(path,self.root/'cache',reuse=True,rule_version=75)
        two=scan.run_inspection(path,self.root/'cache',reuse=True,rule_version=76)
        self.assertNotEqual(one['inspection'],two['inspection'])
        self.assertEqual(two['cache_status'],'miss')
        with patch.object(scan,'inspect',side_effect=AssertionError('repeat parse')):
            self.assertEqual(scan.run_inspection(path,self.root/'cache',reuse=True,rule_version=76)['cache_status'],'hit')
        self.assertEqual(scan.cache_identity(path,'proposal',False,75),scan.cache_identity(path,'proposal',False,76))

    def test_real_rule76_build_readback_and_final_gate(self):
        path=self.root/'content.json'
        path.write_text(json.dumps(self.spec(),ensure_ascii=False),encoding='utf-8')
        with patch.object(scan,'tool_rule_version',return_value=76):
            built=toolkit.build(self.root,path.name,'built')
            report=json.loads(Path(built['inspection']['inspection']).read_text(encoding='utf-8'))
            self.assertEqual(report['claim_template_evidence']['issues'],[])
            cc.require_valid(report,TITLE)
            toolkit.verify_bundle(self.root,'built')
            with patch('prepare_handoff.verify_prepared',return_value={
                'handoff':{'input_inspection':{'inspection':'already verified'}},'input_report':report}):
                final.verify({'rule_version':76,'policy':{'output_contract':final.MODE}},
                             {'title':TITLE},{'input_artifact_sha256':report['sha256']},{})

    def test_verbatim_feature_exception_is_position_and_source_bound(self):
        rows=claims_for()
        rows[0]['text']=rows[0]['text'].replace('获取记录数据，','配置为记录数据一一对应，')
        report,_=self.report(rows)
        hits=report['punctuation_evidence']['preferred_wording_hits']+[
            h for h in report['document_contract_evidence']['global_word_hits'] if h['text']=='一一']
        self.assertTrue(hits)
        source='1'*64
        findings=[dict(h,rule_id='WORDING-07',classification='source_verbatim_feature',status='closed',
                       artifact_sha256=report['sha256'],source_sha256=source,source_locator='synthetic P1',
                       registry_id='test-feature',reason='Synthetic location-bound test, not an actual source acceptance.')
                  for h in hits]
        self.assertTrue(all(cc.is_feature_hit(report,h) for h in hits))
        final.check(report,TITLE,findings,strict_prose=True,source_hashes=[source],strict_claim_templates=True)
        for altered,hashes,new in [(findings,[],True),([], [source],True),(findings,[source],False)]:
            with self.subTest(hashes=hashes,new=new),self.assertRaises(batch.BatchError):
                final.check(report,TITLE,altered,strict_prose=True,source_hashes=hashes,strict_claim_templates=new)
        first=report['claims_detected'][0]
        p=next(p for p in report['parts'][0]['paragraphs'] if p['p']==first['paragraphs'][0])
        header=dict(part='word/document.xml',paragraph=p['p'],offset=2,text='一种')
        self.assertFalse(cc.is_feature_hit(report,header))
        purpose=next(p for p in report['parts'][0]['paragraphs'] if p['text'].startswith('本发明的目的'))
        self.assertFalse(cc.is_feature_hit(report,dict(part='word/document.xml',paragraph=purpose['p'],offset=0,text='本发明')))
        forged=deepcopy(findings)
        forged[0]['offset']+=1
        with self.assertRaises(batch.BatchError):
            final.check(report,TITLE,forged,strict_prose=True,source_hashes=[source],strict_claim_templates=True)

    def test_feature_locator_converts_native_math_placeholder_offsets(self):
        rows=claims_for()
        rows[0]['text']=rows[0]['text'].replace('获取记录数据，','配置为记录数据，')
        _,path=self.report(rows)
        doc=Document(path)
        symbol='<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub></m:oMath>'
        for p in doc.paragraphs:
            if '配置为' not in p.text:continue
            before,after=p.text.split('配置为',1)
            p.text=before;p._p.append(parse_xml(symbol));p.add_run('配置为'+after)
        doc.save(path)
        report=scan.inspect(path);scan.precheck(report,'draft',76)
        hits=[h for h in report['punctuation_evidence']['preferred_wording_hits'] if h['text']=='配置为']
        self.assertEqual(len(hits),2)
        self.assertTrue(all(h['offset_projection']=='plain_text_with_OMML_boundary_placeholder' for h in hits))
        self.assertTrue(all(cc.is_feature_hit(report,h) for h in hits))
        self.assertFalse(cc.is_feature_hit(report,dict(hits[0],offset_projection='unknown')))


if __name__=='__main__':
    unittest.main()
