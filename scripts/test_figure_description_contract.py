from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from PIL import Image
import authoring_toolkit as toolkit
import batch_tool as batch
import inspect_docx as scan
import output_contract as final
import parallel_batch as queue
import figure_description_contract as fig
from test_authoring_toolkit import minimal_spec
from test_parallel_batch import POLICY


class FigureDescriptionTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.path=self.root/'case.docx'

    def spec(self,titles):
        spec=minimal_spec();im=self.root/'source.png';Image.new('RGB',(40,30),'white').save(im)
        spec['sources'].append(dict(id='image',file=im.name,sha256=toolkit.sha(im)))
        blocks=spec['sections'][3]['blocks']
        blocks.append(dict(id='fig_heading',type='paragraph',role='heading',text='附图说明'))
        for n,title in enumerate(titles,1):
            fid='figure'+str(n)
            spec['figures'].append(dict(id=fid,source_id='image',sha256=toolkit.sha(im),number=n,title=title,width_cm=8))
            blocks.append(dict(id='desc'+str(n),type='figure_description',figure_id=fid))
            spec['sections'][4]['blocks'].append(dict(id='pic'+str(n),type='figure',figure_id=fid))
            if n==1:spec['sections'][1]['blocks'].append(dict(id='abstract',type='figure',figure_id=fid))
        blocks.append(dict(id='impl_heading',type='paragraph',role='heading',text='具体实施方式'))
        blocks.append(dict(id='impl',type='paragraph',role='body',text='这是普通说明，图名不在此处自动补字。'))
        return spec

    def build(self,titles):
        spec=self.spec(titles);original=deepcopy(spec)
        ctx=toolkit.validate(self.root,spec);toolkit.render_document(ctx,self.path);toolkit.verify_document(ctx,self.path)
        self.assertEqual(spec,original)
        self.assertEqual([f['title'] for f in ctx['figures'].values()],titles)
        self.assertEqual([f['number'] for f in ctx['figures'].values()],list(range(1,len(titles)+1)))
        return ctx,Document(self.path)

    def report(self,doc):
        doc.save(self.path);before=self.path.read_bytes()
        report=scan.inspect(self.path);scan.precheck(report,'draft')
        self.assertEqual(before,self.path.read_bytes())
        return report

    def accept(self,r):
        return final.check(r,minimal_spec()['project_title'],strict_figures=True)

    def test_all_names_have_one_suffix_and_correct_ending(self):
        titles=['允许进气上限','预测比较','剩余温度响应','流程框图','示意图']
        ctx,doc=self.build(titles);r=self.report(doc);self.accept(r)
        items=r['figure_description_evidence']['items']
        self.assertEqual([x['name'] for x in items],['允许进气上限图','预测比较图','剩余温度响应图','流程框图','示意图'])
        self.assertEqual([x['punctuation'] for x in items],['；']*4+['。'])
        for p in doc.paragraphs:
            if p.text.startswith('图') and '为本发明的' in p.text:
                self.assertTrue(all(str(run.font.color.rgb)=='000000' for run in p.runs))
        # Only the final figures have standalone captions; no abstract caption.
        self.assertEqual([p.text for p in doc.paragraphs if p.text in {f'图{n}' for n in range(1,6)}],['图1','图2','图3','图4','图5'])
        self.assertEqual(ctx['block_texts']['impl'],'这是普通说明，图名不在此处自动补字。')

    def test_missing_suffix_across_runs_rejected(self):
        _,doc=self.build(['温度响应']);p=next(p for p in doc.paragraphs if p.text.startswith('图1为'))
        p.clear()
        for c in '图1为本发明的温度响应。':p.add_run(c)
        with self.assertRaisesRegex(batch.BatchError,'missing_tu'):self.accept(self.report(doc))

    def test_multiple_entries_in_one_paragraph_checked_individually(self):
        _,doc=self.build(['温度响应']);p=next(p for p in doc.paragraphs if p.text.startswith('图1为'))
        p.text='图1为本发明的上限；图2为本发明的比较图。'
        r=self.report(doc)
        self.assertEqual(len(r['figure_description_evidence']['items']),2)
        with self.assertRaisesRegex(batch.BatchError,'missing_tu'):self.accept(r)

    def test_bad_suffix_or_input_punctuation_is_not_silently_rewritten(self):
        for title in ['温度响应图图','温度响应图。','温度响应；']:
            with self.subTest(title=title),self.assertRaises(toolkit.AuthoringError):toolkit.validate(self.root,self.spec([title]))
        _,doc=self.build(['温度响应']);p=next(p for p in doc.paragraphs if p.text.startswith('图1为'))
        p.text='图1为本发明的温度响应图图。'
        with self.assertRaisesRegex(batch.BatchError,'duplicate_tu'):self.accept(self.report(doc))

    def test_punctuation_and_unparsed_text_cannot_hide_missing_suffix(self):
        for text in ['图1为本发明的温度响应图；','图1为本发明的温度响应', '温度响应。']:
            _,doc=self.build(['温度响应']);p=next(p for p in doc.paragraphs if p.text.startswith('图1为'));p.text=text
            with self.subTest(text=text),self.assertRaises(batch.BatchError):self.accept(self.report(doc))

    def test_empty_figure_section_does_not_invent_a_figure(self):
        _,doc=self.build([]);r=self.report(doc);self.accept(r)
        self.assertEqual(r['figure_description_evidence']['items'],[])
        self.assertFalse(doc.inline_shapes)

    def test_scope_and_artifact_must_resolve_legacy_gate_stays_optional(self):
        _,doc=self.build(['温度响应']);r=self.report(doc)
        for key,value in [('complete',False),('schema_version',True),('artifact_sha256','f'*64),('items',None)]:
            changed=deepcopy(r);changed['figure_description_evidence'][key]=value
            with self.subTest(key=key),self.assertRaises(batch.BatchError):self.accept(changed)
        r.pop('figure_description_evidence');final.check(r,minimal_spec()['project_title'])
        with self.assertRaises(batch.BatchError):self.accept(r)
        next(p for p in doc.paragraphs if p.text=='附图说明').text='无法识别的栏目'
        with self.assertRaises(batch.BatchError):self.accept(self.report(doc))

    def test_new_policy_validated_and_sealed_report_reused(self):
        with self.assertRaisesRegex(batch.BatchError,'Invalid figure description'):
            queue.validate_policy(dict(POLICY,figure_description_contract='unknown'))
        with self.assertRaisesRegex(batch.BatchError,'requires'):
            queue.validate_policy(dict(POLICY,figure_description_contract=fig.MODE))
        _,doc=self.build(['温度响应']);r=self.report(doc)
        prepared={'handoff':{'input_inspection':{'inspection':'sealed.json'}},'input_report':r}
        with patch('prepare_handoff.verify_prepared',return_value=prepared),patch.object(scan,'inspect',side_effect=AssertionError('duplicate extraction')):
            got=final.verify({'policy':{'output_contract':final.MODE,'figure_description_contract':fig.MODE}},
                {'title':minimal_spec()['project_title']},{'input_artifact_sha256':r['sha256']},{})
            self.assertEqual(got['figure_description_contract'],fig.MODE)


if __name__=='__main__':unittest.main()
