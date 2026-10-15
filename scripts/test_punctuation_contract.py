"""No model calls or visual claims: source-safe punctuation and final-gate tests."""
import copy,hashlib,unittest,zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from unittest.mock import patch
from docx import Document
from docx.oxml import parse_xml
import batch_tool as b
import inspect_docx as scan
import output_contract as contract
import parallel_batch as queue
import test_output_contract as base

W=scan.NS['w'];M=scan.NS['m'];A=scan.NS['a']


class PunctuationTests(unittest.TestCase):
    setUp=base.OutputContractTests.setUp
    report=base.OutputContractTests.report

    def reread(self,edit):
        self.report();p=self.root/'scratch.docx';doc=Document(p);edit(doc);doc.save(p)
        before=b.sha(p);r=scan.inspect(p);scan.precheck(r,'draft');self.assertEqual(before,b.sha(p))
        return r

    def test_unicode_quotes_brackets_and_compatibility_enclosures_are_blocked(self):
        for ch in '"\'＂＇“”‘’„‟‹›«»「」『』﹁﹂❛❞❟❠🙶🙷🙸()（）[]［］{}｛｝【】〈〉⟨⟩⏞⏟⎛⎜⎝⎡⎨⎶⑴⒜㈠':
            with self.subTest(ch=ch):
                self.assertTrue(scan.forbidden_punctuation(ch))
                with self.assertRaisesRegex(b.BatchError,'quotation/bracket'):
                    contract.check(self.report('保存'+ch+'记录。'),'一种处理系统',strict_prose=True)

    def test_math_primes_comparisons_and_summation_parts_are_not_brackets(self):
        r=self.report('x′≤y″，x‴>0，x‵<y‶，x⁗≠y，Σ⎲⎳。')
        contract.check(r,'一种处理系统',strict_prose=True)
        self.assertFalse(r['punctuation_evidence']['hits'])

    def test_real_nested_word_text_header_footer_and_run_splits_are_scanned(self):
        def edit(doc):
            p=doc.add_paragraph();p.add_run('“');p.add_run('对象');p.add_run('”')
            doc.add_table(rows=1,cols=1).cell(0,0).text='[数据]'
            doc.sections[0].header.paragraphs[0].text='「页眉」'
            doc.sections[0].footer.paragraphs[0].text='（页脚）'
            doc.add_paragraph()._p.append(parse_xml(f'<w:r xmlns:w="{W}"><w:pict><w:txbxContent><w:p><w:r><w:t>【框内】</w:t></w:r></w:p></w:txbxContent></w:pict></w:r>'))
        r=self.reread(edit);hits=r['punctuation_evidence']['hits']
        self.assertEqual(len(hits),10)
        self.assertEqual({h['part'] for h in hits},{'word/document.xml','word/header1.xml','word/footer1.xml'})
        self.assertTrue(all(h['paragraph']>0 and h['path'] for h in hits))

    def math_hits(self,body):
        root=ET.fromstring(f'<w:document xmlns:w="{W}" xmlns:m="{M}"><w:body><w:p><m:oMath>{body}</m:oMath></w:p></w:body></w:document>')
        return scan.punctuation_evidence(root,'word/document.xml')['formula_protected_candidates']

    def test_implicit_math_fences_empty_properties_and_matrix_boundary(self):
        e='<m:e><m:r><m:t>x</m:t></m:r></m:e>'
        self.assertEqual([h['text'] for h in self.math_hits('<m:d>'+e+'</m:d>')],['(',')'])
        self.assertEqual([h['text'] for h in self.math_hits('<m:d><m:dPr><m:begChr/></m:dPr>'+e+'</m:d>')],[')'])
        self.assertFalse(self.math_hits('<m:d><m:dPr><m:begChr m:val=""/><m:endChr/></m:dPr>'+e+'</m:d>'))
        self.assertEqual([h['text'] for h in self.math_hits('<m:d><m:dPr><m:begChr m:val="&lt;"/><m:endChr m:val="&gt;"/></m:dPr>'+e+'</m:d>')],['<','>'])
        self.assertFalse(self.math_hits('<m:m><m:mr>'+e+'</m:mr></m:m>'))
        self.assertEqual([h['text'] for h in self.math_hits('<m:groupChr>'+e+'</m:groupChr>')],['⏟'])
        self.assertFalse(self.math_hits('<m:groupChr><m:groupChrPr><m:chr m:val="→"/></m:groupChrPr>'+e+'</m:groupChr>'))

    def test_source_math_punctuation_is_preserved_and_does_not_block_final_gate(self):
        def edit(doc):
            doc.add_paragraph()._p.append(parse_xml(f'<m:oMath xmlns:m="{M}"><m:d><m:e><m:r><m:t>x</m:t></m:r></m:e></m:d></m:oMath>'))
        r=self.reread(edit)
        contract.check(r,'一种处理系统',strict_prose=True)
        self.assertFalse(r['punctuation_evidence']['hits'])
        self.assertIn('math_delimiter',{h['representation'] for h in r['punctuation_evidence']['formula_protected_candidates']})
        self.assertIn('<',r['formulas'][0]['omml_xml'])

    def test_text_formula_requires_source_bound_character_evidence_not_paragraph_waiver(self):
        r=self.report('使用f(x)与“记录”计算。');hits=r['punctuation_evidence']['hits']
        digest='a'*64
        findings=[dict(h,rule_id='WORDING-06',classification='source_formula',status='closed',
                       artifact_sha256=r['sha256'],source_sha256=digest,source_locator='源公式1',registry_id='F1',reason='原公式的函数参数括号')
                  for h in hits if h['text'] in '()']
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',findings,True,[digest])
        r=self.report('使用f(x)计算。')
        findings=[dict(h,rule_id='WORDING-06',classification='source_formula',status='closed',
                       artifact_sha256=r['sha256'],source_sha256=digest,source_locator='源公式1',registry_id='F1',reason='原公式括号')
                  for h in r['punctuation_evidence']['hits']]
        contract.check(r,'一种处理系统',findings,True,[digest])
        for hashes in ([],['b'*64]):
            with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',findings,True,hashes)
        findings[0]['offset']+=1
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',findings,True,[digest])

    def test_font_markers_field_code_and_deleted_text_do_not_create_fake_hits(self):
        def edit(doc):
            doc.add_paragraph()._p.append(parse_xml(f'''<w:r xmlns:w="{W}"><w:sym w:font="Symbol" w:char="F061"/><w:instrText>MERGEFIELD "a(b)"</w:instrText><w:fldChar w:fldCharType="begin"/></w:r>'''))
            doc.add_paragraph()._p.append(parse_xml(f'<w:del xmlns:w="{W}"><w:r><w:delText>（已删除）</w:delText></w:r></w:del>'))
        r=self.reread(edit)
        self.assertFalse(r['punctuation_evidence']['hits'])
        self.assertTrue(r['punctuation_evidence']['review_candidates'])
        self.assertIn('[w:sym',str(r['parts']))

    def test_chart_drawingml_vml_and_note_parts_are_scanned(self):
        self.report();p=self.root/'scratch.docx'
        extra={
            'word/charts/chart7.xml':'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:v>（图表）</c:v></c:chartSpace>',
            'word/diagrams/data7.xml':f'<a:txBody xmlns:a="{A}"><a:p><a:r><a:t>“图形”</a:t></a:r></a:p></a:txBody>',
            'word/drawings/vmlDrawing7.vml':'<xml xmlns:v="urn:schemas-microsoft-com:vml"><v:textpath string="[艺术字]"/></xml>'}
        for name in ['footnotes','endnotes','comments']:
            extra['word/'+name+'.xml']=f'<w:{name} xmlns:w="{W}"><w:p><w:r><w:t>【内容】</w:t></w:r></w:p></w:{name}>'
        with zipfile.ZipFile(p,'a') as z:
            for name,xml in extra.items():z.writestr(name,xml)
        r=scan.inspect(p);scan.precheck(r,'draft')
        self.assertEqual({h['part'] for h in r['punctuation_evidence']['hits']},set(extra))

    def test_numbering_used_and_style_inherited_labels_block_unused_definition_does_not(self):
        numbering=ET.fromstring(f'<w:numbering xmlns:w="{W}"><w:abstractNum w:abstractNumId="99"><w:lvl w:ilvl="0"><w:lvlText w:val="%1）"/></w:lvl></w:abstractNum><w:num w:numId="99"><w:abstractNumId w:val="99"/></w:num></w:numbering>')
        p=ET.fromstring(f'<w:p xmlns:w="{W}"><w:pPr><w:numPr><w:numId w:val="99"/></w:numPr></w:pPr></w:p>')
        self.assertEqual([h['text'] for h in scan.numbered_punctuation(p,numbering,{})],['）'])
        style=ET.fromstring(f'<w:style xmlns:w="{W}" w:styleId="List"><w:pPr><w:numPr><w:numId w:val="99"/></w:numPr></w:pPr></w:style>')
        p=ET.fromstring(f'<w:p xmlns:w="{W}"><w:pPr><w:pStyle w:val="List"/></w:pPr></w:p>')
        self.assertTrue(scan.numbered_punctuation(p,numbering,{'List':style}))
        self.assertFalse(scan.numbered_punctuation(p,numbering,{}))

    def test_incomplete_evidence_and_forbidden_formal_title_rejected_legacy_preserved(self):
        r=self.report('保存（记录）。');contract.check(r,'一种处理系统')
        del r['punctuation_evidence']
        with self.assertRaisesRegex(b.BatchError,'missing or incomplete'):
            contract.check(r,'一种处理系统',strict_prose=True)
        self.assertEqual(contract.filename('一种处理系统（甲）'),'一种处理系统（甲）.docx')
        r=self.report(title='一种处理系统（甲）')
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统（甲）',strict_prose=True)

    def test_preferred_phrases_cross_runs_preserve_other_nouns_and_new_accepted_text(self):
        for phrase,replacement in scan.PREFERRED_WORDING.items():
            r=self.report(phrase+'关联数据模型。')
            matches=r['punctuation_evidence']['preferred_wording_hits']
            self.assertEqual(len(matches),1)
            self.assertEqual(matches[0]['replacement'],replacement)
            with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',strict_prose=True)
        clean=self.report('优选的，关联数据模型用于处理文中各记录，文中各标识都相对应；配置参数保持原值。')
        contract.check(clean,'一种处理系统',strict_prose=True)
        self.assertFalse(clean['punctuation_evidence']['preferred_wording_hits'])

    def test_protected_technical_name_needs_exact_source_and_artifact_binding(self):
        r=self.report('使用源文件标识进行索引。');h=r['punctuation_evidence']['preferred_wording_hits'][0];digest='a'*64
        f=dict(h,rule_id='WORDING-07',classification='protected_technical_name',status='closed',
               artifact_sha256=r['sha256'],source_sha256=digest,source_locator='术语定义',registry_id='T1',reason='源文件标识为完整技术名词')
        contract.check(r,'一种处理系统',[f],True,[digest])
        f['artifact_sha256']='stale'
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',[f],True,[digest])

    def test_defined_technical_name_protection_covers_only_exact_yiyi_hits(self):
        r=self.report('读取一一对应映射表。');digest='a'*64
        hits=r['document_contract_evidence']['global_word_hits']+r['punctuation_evidence']['preferred_wording_hits']
        findings=[dict(h,rule_id='WORDING-07',classification='protected_technical_name',status='closed',
                    artifact_sha256=r['sha256'],source_sha256=digest,source_locator='完整术语定义',registry_id='T2',reason='一一对应映射表为来源完整定义名称') for h in hits]
        contract.check(r,'一种处理系统',findings,True,[digest])
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',findings[1:],True,[digest])
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',findings,False,[digest])

    def test_formula_string_literals_preserved_without_exempting_adjacent_prose(self):
        r=self.report('计算hash("正文一一对应权利要求从而")。');digest='c'*64
        hits=(r['punctuation_evidence']['hits']+r['punctuation_evidence']['preferred_wording_hits']+
              r['document_contract_evidence']['global_word_hits']+r['document_contract_evidence']['description_word_hits'])
        findings=[dict(h,rule_id='WORDING-06',classification='source_formula',status='closed',
                       artifact_sha256=r['sha256'],source_sha256=digest,source_locator='公式2',registry_id='F2',reason='原式字符串常量保持原样') for h in hits]
        contract.check(r,'一种处理系统',findings,True,[digest])
        r['punctuation_evidence']['preferred_wording_hits'].append(dict(r['punctuation_evidence']['preferred_wording_hits'][0],offset=99))
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',findings,True,[digest])

    def test_native_formula_strings_protected_without_waiving_same_paragraph_prose(self):
        def edit(doc):
            p=doc.add_paragraph('计算')
            p._p.append(parse_xml(f'<m:oMath xmlns:m="{M}"><m:r><m:t>hash("一一对应权利要求从而")</m:t></m:r></m:oMath>'))
        r=self.reread(edit);contract.check(r,'一种处理系统',strict_prose=True)
        self.assertTrue(all(h['formula_protected_candidate'] for h in r['document_contract_evidence']['global_word_hits']))
        path=self.root/'scratch.docx';doc=Document(path);doc.paragraphs[-1].add_run('一一对应');doc.save(path)
        r=scan.inspect(path);scan.precheck(r,'draft')
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',strict_prose=True)

    def test_omml_separates_plain_text_phrases_and_drawing_bullet_is_not_formula(self):
        def edit(doc):
            p=doc.add_paragraph('配')
            p._p.append(parse_xml(f'<m:oMath xmlns:m="{M}"><m:r><m:t>x</m:t></m:r></m:oMath>'))
            p.add_run('置为')
            p._p.append(parse_xml(f'<a:p xmlns:a="{A}"><a:pPr><a:buChar char="（"/></a:pPr><a:r><a:t>图中文字</a:t></a:r></a:p>'))
        r=self.reread(edit)
        self.assertFalse(r['punctuation_evidence']['preferred_wording_hits'])
        hit=next(h for h in r['punctuation_evidence']['hits'] if h['representation']=='drawing_bullet')
        digest='d'*64;f=dict(hit,rule_id='WORDING-06',classification='source_formula',status='closed',
            artifact_sha256=r['sha256'],source_sha256=digest,source_locator='假公式',registry_id='F3',reason='试图豁免自动编号')
        with self.assertRaises(b.BatchError):contract.check(r,'一种处理系统',[f],True,[digest])

    def test_breaks_and_display_objects_do_not_join_nonadjacent_phrases(self):
        for middle in ['<w:tab/>','<w:br/>','<w:cr/>','<a:br/>','<w:sym w:font="Symbol" w:char="F061"/>','<a:blip/>','<o:OLEObject/>','<v:imagedata/>']:
            root=ET.fromstring(f'<w:document xmlns:w="{W}" xmlns:a="{A}" xmlns:o="{scan.NS["o"]}" xmlns:v="{scan.NS["v"]}"><w:body><w:p><w:r><w:t>配</w:t>{middle}<w:t>置为</w:t></w:r></w:p></w:body></w:document>')
            with self.subTest(middle=middle):self.assertFalse(scan.punctuation_evidence(root,'word/document.xml')['preferred_wording_hits'])


class PunctuationQueueTests(unittest.TestCase):
    reserve=base.QueueContractTests.reserve
    receipt=base.QueueContractTests.receipt
    review_attempt=base.QueueContractTests.review_attempt

    def setUp(self):
        original=queue.init
        def init(*args,**kwargs):
            policy=b.read(queue.POLICY);policy['prose_contract']=contract.PROSE_MODE
            path=Path(args[0]).parent/'v19-policy.json';b.atomic(path,policy)
            with patch.object(queue,'POLICY',path):return original(*args,**kwargs)
        with patch.object(queue,'init',side_effect=init):base.QueueContractTests.setUp(self)

    def test_new_queue_rejects_and_cannot_deliver_quoted_output(self):
        review,receipt=self.review_attempt('保存“关联”记录。')
        with self.assertRaisesRegex(b.BatchError,'quotation/bracket'):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
        with self.assertRaises(b.BatchError):queue.deliver(self.manifest)
        self.assertFalse(list((self.root/'desktop').rglob('*.docx')))

    def test_clean_output_passes_without_second_parse_and_delivery_binds_new_mode(self):
        review,receipt=self.review_attempt()
        with patch.object(scan,'inspect',side_effect=AssertionError('reparsed')):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
            dest=Path(queue.deliver(self.manifest))
        self.assertTrue((dest/'一种处理系统.docx').is_file())
        r=b.read(self.manifest)['projects'][0]['receipts'][-1]
        self.assertEqual(r['document_contract']['prose_contract'],contract.PROSE_MODE)

    def test_delivery_refuses_tampered_punctuation_acceptance(self):
        review,receipt=self.review_attempt();queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
        m=b.read(self.manifest);del m['projects'][0]['receipts'][-1]['document_contract']['prose_contract'];b.atomic(self.manifest,m)
        with self.assertRaises(b.BatchError):queue.deliver(self.manifest)
        self.assertFalse(list((self.root/'desktop').rglob('*.docx')))


if __name__=='__main__':unittest.main()
