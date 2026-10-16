"""Located pagination evidence and sealed v22 acceptance, not page-fill guesses."""
import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.enum.section import WD_SECTION_START
import batch_tool as b
import inspect_docx as scan
import output_contract as contract
import parallel_batch as queue
import test_inspect_docx as xml_fixtures
import test_parallel_batch as queue_fixtures


FALSE_FLOW = ''.join('<w:'+k+' w:val="0"/>' for k in ('keepLines','keepNext','pageBreakBefore','widowControl'))
TITLE = '一种记录处理方法'


def paragraph(text='', props='', content=''):
    return '<w:p><w:pPr>'+props+'</w:pPr><w:r><w:t>'+text+'</w:t></w:r>'+content+'</w:p>'


class FlowEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def report(self, body='', flow=FALSE_FLOW, abstract_label='', styles='', abstract_content=None, abstract_label_props='', final_content=None):
        sections=[paragraph('摘要内容。',FALSE_FLOW),paragraph('',FALSE_FLOW, abstract_content or '<w:r><w:drawing/></w:r>')+
                  (paragraph(abstract_label,FALSE_FLOW+abstract_label_props) if abstract_label else ''),
                  paragraph('1.一种记录处理方法，其特征在于，取得记录。',FALSE_FLOW),
                  paragraph(TITLE,FALSE_FLOW)+paragraph('技术领域',FALSE_FLOW)+
                  paragraph('本发明涉及记录处理。',FALSE_FLOW)+paragraph('具体实施方式',FALSE_FLOW)+
                  (body or paragraph('保存更新后的记录。',flow)),
                  final_content if final_content is not None else paragraph('',FALSE_FLOW,'<w:r><w:drawing/></w:r>')+paragraph('图1',FALSE_FLOW)]
        headers=['说明书摘要','摘要附图','权利要求书','说明书','说明书附图']
        rels=[];parts={};complete=''
        for i,(section,header) in enumerate(zip(sections,headers),1):
            rid='h'+str(i);rels.append(f'<Relationship Id="{rid}" Type="header" Target="header{i}.xml"/>')
            parts[f'word/header{i}.xml']=f'<w:hdr xmlns:w="{scan.NS["w"]}">'+paragraph(header)+'</w:hdr>'
            sp=f'<w:sectPr><w:headerReference w:type="default" r:id="{rid}"/><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1417" w:bottom="850"/><w:docGrid w:type="lines" w:linePitch="312"/></w:sectPr>'
            complete+=section+(paragraph('',sp) if i<5 else sp)
        parts['word/_rels/document.xml.rels']='<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'+''.join(rels)+'</Relationships>'
        if styles:parts['word/styles.xml']=f'<w:styles xmlns:w="{scan.NS["w"]}">'+styles+'</w:styles>'
        report=scan.inspect(xml_fixtures.InspectionTests().make_docx(self.root,complete,parts))
        scan.precheck(report,'draft')
        return report

    def body_row(self, report, text='保存更新后的记录。'):
        p=next(p for p in report['parts'][0]['paragraphs'] if p['text']==text)
        return next(r for r in report['page_flow_evidence']['paragraphs'] if r['paragraph']==p['p'])

    def check(self, report, findings=(), strict=True):
        return contract.check(report,TITLE,findings,strict_layout=strict)

    def test_inherited_true_and_explicit_false_have_actual_sources(self):
        styles='<w:style w:styleId="Base"><w:pPr><w:keepLines/></w:pPr></w:style><w:style w:styleId="Child"><w:basedOn w:val="Base"/></w:style>'
        flow='<w:pStyle w:val="Child"/><w:keepNext w:val="false"/><w:pageBreakBefore w:val="off"/><w:widowControl w:val="0"/>'
        report=self.report(flow=flow,styles=styles);row=self.body_row(report)
        self.assertEqual(row['effective']['keepLines'],{'value':True,'source':'style:Base','raw':'1'})
        self.assertFalse(row['effective']['keepNext']['value'])
        with self.assertRaisesRegex(b.BatchError,'page flow'):self.check(report)
        report=self.report(flow=flow+'<w:keepLines w:val="0"/>',styles=styles)
        self.assertEqual(self.body_row(report)['effective']['keepLines']['source'],'direct')
        self.check(report)

    def test_document_defaults_and_specification_default_are_distinguished(self):
        styles='<w:docDefaults><w:pPrDefault><w:pPr>'+FALSE_FLOW+'</w:pPr></w:pPrDefault></w:docDefaults>'
        report=self.report(flow='',styles=styles)
        self.assertEqual(self.body_row(report)['effective']['widowControl']['source'],'docDefaults')
        self.check(report)
        report=self.report(flow='');row=self.body_row(report)
        self.assertEqual(row['effective']['keepLines']['value'],False)
        self.assertEqual(row['effective']['widowControl'],{'value':True,'source':'ooxml_default','raw':None})
        with self.assertRaises(b.BatchError):self.check(report)

    def test_missing_or_cyclic_style_does_not_silently_certify_defaults(self):
        for styles in ['', '<w:style w:styleId="Broken"><w:basedOn w:val="Broken"/></w:style>']:
            report=self.report(flow='<w:pStyle w:val="Broken"/>',styles=styles)
            row=self.body_row(report)
            self.assertIsNone(row['effective']['keepLines']['value'])
            self.assertTrue(row['unresolved'])
            with self.assertRaisesRegex(b.BatchError,'located review'):self.check(report)

    def test_formula_image_ole_field_and_symbol_are_not_empty_body_paragraphs(self):
        payloads=['<m:oMath/>','<w:r><w:drawing/></w:r>','<w:r><w:object><o:OLEObject/></w:object></w:r>',
                  '<w:fldSimple w:instr="PAGE"/>','<w:r><w:sym w:font="Symbol" w:char="03B1"/></w:r>']
        body=''.join(paragraph('',FALSE_FLOW,p) for p in payloads)+paragraph('',FALSE_FLOW)
        report=self.report(body=body)
        rows=[r for r in report['page_flow_evidence']['paragraphs'] if r['section']=='description']
        content=[r for r in rows if r['content']['objects']]
        self.assertEqual(len(content),5)
        self.assertTrue(all(r['content']['kind']!='empty' for r in content))
        self.assertEqual(sum(r['content']['kind']=='empty' and not r['content']['section_boundary'] for r in rows),1)
        self.assertFalse(any(r['classification']=='ordinary' for r in content))

    def test_sample_pattern_all_keep_lines_fails_but_reference_flow_passes(self):
        report=self.report(flow=FALSE_FLOW.replace('<w:keepLines w:val="0"/>','<w:keepLines w:val="1"/>'))
        self.assertTrue(any(i['code']=='PAGE_FLOW_REVIEW' for i in report['issues']))
        with self.assertRaises(b.BatchError):self.check(report)
        self.check(self.report())

    def test_external_equation_label_requires_located_review_but_inline_prose_cannot_bypass_flow(self):
        math='<m:oMath><m:r><m:t>x=y</m:t></m:r></m:oMath>'
        keep=FALSE_FLOW.replace('<w:keepLines w:val="0"/>','<w:keepLines w:val="1"/>')
        report=self.report(body=paragraph('(1)',keep,math))
        ev=report['page_flow_evidence']
        self.assertFalse(ev['ordinary_violations'])
        candidate=next(c for c in ev['review_candidates'] if c['reason']=='complex_body_content')
        finding=dict(candidate,rule_id='FORMAT-03',classification='page_flow_reviewed',status='closed',
                     artifact_sha256=report['sha256'],reason='已核对这是完整原式及其外置公式号，保持整式同页。')
        with self.assertRaises(b.BatchError):self.check(report)
        self.check(report,[finding])
        report=self.report(body=paragraph('通过该变量说明普通文字内容。',keep,math))
        self.assertTrue(report['page_flow_evidence']['ordinary_violations'])
        with self.assertRaises(b.BatchError):self.check(report,[finding])

    def test_abstract_extra_label_is_rejected_and_final_caption_is_preserved(self):
        report=self.report(abstract_label='图1 整体流程图')
        ev=report['page_flow_evidence']
        self.assertEqual([r['text'] for r in ev['abstract_labels']],['图1 整体流程图'])
        self.assertEqual([r['text'] for r in ev['final_captions']],['图1'])
        with self.assertRaisesRegex(b.BatchError,'abstract figure'):self.check(report)
        self.check(self.report())

    def test_abstract_label_cannot_be_hidden_by_heading_style_or_outline(self):
        styles='<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>'
        for props in ['<w:pStyle w:val="Heading1"/>','<w:outlineLvl w:val="0"/>']:
            with self.subTest(props=props):
                report=self.report(abstract_label='图1 整体流程图',styles=styles,abstract_label_props=props)
                self.assertEqual(len(report['page_flow_evidence']['abstract_labels']),1)
                with self.assertRaisesRegex(b.BatchError,'abstract figure'):self.check(report)

    def test_known_heading_and_display_formula_can_remain_together(self):
        keep=FALSE_FLOW.replace('<w:keepLines w:val="0"/>','<w:keepLines/>').replace('<w:keepNext w:val="0"/>','<w:keepNext/>')
        report=self.report(body=paragraph('与现有技术相比，本发明的有益效果在于：',keep)+
                           paragraph('',keep,'<m:oMath><m:r><m:t>x=y</m:t></m:r></m:oMath>')+
                           paragraph('保存更新后的记录。',FALSE_FLOW))
        self.assertFalse(report['page_flow_evidence']['ordinary_violations'])
        self.check(report)
        # Custom direct-format headings have no semantic role in OOXML. They
        # need the original located review instead of a false ordinary-body ban.
        report=self.report(body='<w:p><w:pPr>'+keep+'</w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>实施例甲</w:t></w:r></w:p>')
        self.assertFalse(report['page_flow_evidence']['ordinary_violations'])
        self.assertTrue(report['page_flow_evidence']['review_candidates'])
        with self.assertRaisesRegex(b.BatchError,'located review'):self.check(report)

    def test_text_inside_abstract_picture_is_not_an_external_caption(self):
        report=self.report(abstract_content='<w:r><w:pict><w:txbxContent>'+paragraph('图中技术节点')+'</w:txbxContent></w:pict></w:r>')
        ev=report['page_flow_evidence']
        nested=next(p for p in report['parts'][0]['paragraphs'] if p['text']=='图中技术节点')
        self.assertFalse(nested['body_direct'])
        self.assertFalse(ev['abstract_labels'])
        self.assertTrue(any(r['paragraph']==nested['p'] and r['reason']=='complex_body_content' for r in ev['review_candidates']))

    def test_unknown_structure_requires_current_position_bound_existing_review(self):
        report=self.report(body=paragraph('读取域结果。',FALSE_FLOW,'<w:fldSimple w:instr="TEST"/>'))
        candidate=next(c for c in report['page_flow_evidence']['review_candidates'] if c['reason']=='complex_body_content')
        finding=dict(candidate,rule_id='FORMAT-03',classification='page_flow_reviewed',status='closed',
                     artifact_sha256=report['sha256'],reason='已核对本段域结果及分页，属于保留的复杂对象。')
        self.check(report,[finding])
        for key,value in [('artifact_sha256','wrong'),('paragraph',0),('path','body/w:p[999]')]:
            bad=dict(finding);bad[key]=value
            with self.subTest(key=key),self.assertRaises(b.BatchError):self.check(report,[bad])

    def test_legacy_contract_does_not_require_v22_evidence(self):
        report=self.report();report.pop('page_flow_evidence')
        self.check(report,strict=False)
        with self.assertRaises(b.BatchError):self.check(report)

    def test_sealed_report_is_reused_without_reparsing_and_keeps_layout_marker(self):
        report=self.report();report_file=self.root/'inspection.json';b.atomic(report_file,report)
        handoff_file=self.root/'handoff.json';b.atomic(handoff_file,{'input_inspection':{'inspection':str(report_file)}})
        attempt={'prepared_handoff':{'path':str(handoff_file)},'input_artifact_sha256':report['sha256']}
        manifest={'policy':{'output_contract':contract.MODE,'layout_contract':contract.LAYOUT_MODE}}
        verified={'handoff':{'input_inspection':{'inspection':str(report_file)}},'input_report':report}
        with patch('prepare_handoff.verify_prepared',return_value=verified),patch.object(scan,'inspect',side_effect=AssertionError('duplicate parse')):
            result=contract.verify(manifest,{'title':TITLE},attempt,{'findings':[]})
        self.assertEqual(result['layout_contract'],contract.LAYOUT_MODE)
        attempt['input_artifact_sha256']='wrong'
        with patch('prepare_handoff.verify_prepared',return_value=verified),self.assertRaises(b.BatchError):
            contract.verify(manifest,{'title':TITLE},attempt,{'findings':[]})

    def test_layout_policy_needs_valid_flag_and_existing_sealed_contract(self):
        p=copy.deepcopy(queue_fixtures.POLICY);p['layout_contract']='wrong'
        with self.assertRaises(b.BatchError):queue.validate_policy(p)
        p['layout_contract']=contract.LAYOUT_MODE
        with self.assertRaises(b.BatchError):queue.validate_policy(p)
        p.update(output_contract=contract.MODE,worker_preparation='frozen_stage_packet_v1',
                 delivery={'folder_prefix':'专利撰写完成','layout':'flat_final_files','timing':'after_each_acceptance'})
        queue.validate_policy(p)


class QueueFlowContractTests(unittest.TestCase):
    reserve=queue_fixtures.ParallelTests.reserve
    receipt=queue_fixtures.ParallelTests.receipt

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.src=self.root/'input';self.src.mkdir()
        doc=Document();doc.add_paragraph('源技术内容');doc.save(self.src/'source.docx')
        policy=copy.deepcopy(queue_fixtures.POLICY)
        policy.update(worker_preparation='frozen_stage_packet_v1',output_contract=contract.MODE,
                      layout_contract=contract.LAYOUT_MODE,
                      delivery={'folder_prefix':'专利撰写完成','layout':'flat_final_files','timing':'after_each_acceptance'})
        policy_path=self.root/'policy.json';b.atomic(policy_path,policy)
        selection={'projects':[{'title':TITLE,'source_files':['source.docx'],'primary_source':'source.docx'}]}
        with patch.object(queue,'POLICY',policy_path):
            self.manifest=queue.init(self.src,self.root/'work',selection,self.root/'desktop')
        self.batch=self.manifest.parent

    def review_attempt(self,bad_flow=False,abstract_caption=False):
        draft=self.reserve(1);receipt=self.receipt(draft);doc=Document()
        for i,(header,texts) in enumerate([
                ('说明书摘要',['摘要内容。']),('摘要附图',['图1'] if abstract_caption else []),
                ('权利要求书',['1.一种记录处理方法，其特征在于，取得记录。']),
                ('说明书',[TITLE,'技术领域','本发明涉及记录处理。','具体实施方式','保存更新后的记录。']),
                ('说明书附图',['图1'])]):
            section=doc.sections[0] if i==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.header.is_linked_to_previous=False;section.header.paragraphs[0].text=header
            for text in texts:
                p=doc.add_paragraph(text);pf=p.paragraph_format
                pf.keep_together=bad_flow and text=='保存更新后的记录。'
                pf.keep_with_next=False;pf.page_break_before=False;pf.widow_control=False
        doc.save(receipt['artifact']);receipt['artifact_sha256']=b.sha(receipt['artifact'])
        queue.record(self.manifest,draft['case_id'],draft['attempt']['id'],receipt)
        review=self.reserve(1);return review,self.receipt(review,'pass')

    def test_same_artifact_acceptance_and_delivery_reuse_v22_evidence(self):
        review,receipt=self.review_attempt()
        with patch.object(scan,'inspect',side_effect=AssertionError('duplicate parse')):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
            dest=Path(queue.deliver(self.manifest))
        self.assertEqual([p.name for p in dest.glob('*.docx')],[TITLE+'.docx'])
        saved=b.read(self.manifest)['projects'][0]['receipts'][-1]['document_contract']
        self.assertEqual(saved['layout_contract'],contract.LAYOUT_MODE)

    def test_flow_or_abstract_label_prevents_acceptance(self):
        review,receipt=self.review_attempt(bad_flow=True)
        with self.assertRaisesRegex(b.BatchError,'page flow'):
            queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
        self.assertNotEqual(b.read(self.manifest)['projects'][0]['status'],'accepted')
        with self.assertRaises(b.BatchError):queue.deliver(self.manifest)

    def test_delivery_rechecks_frozen_layout_result(self):
        review,receipt=self.review_attempt();queue.record(self.manifest,review['case_id'],review['attempt']['id'],receipt)
        m=b.read(self.manifest);m['projects'][0]['receipts'][-1]['document_contract']['layout_contract']='wrong'
        b.atomic(self.manifest,m)
        with self.assertRaises(b.BatchError):queue.deliver(self.manifest)


if __name__=='__main__':unittest.main()
