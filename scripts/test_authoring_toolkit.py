"""Source preservation, real DOCX readback, and no-overwrite regression tests."""
from copy import deepcopy
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from lxml import etree as ET
from PIL import Image

import authoring_toolkit as toolkit


def minimal_spec():
    return {
        'schema_version': 1, 'case_id': 'demo', 'project_title': '文稿组件示例名称',
        'protected_names': [{'text': '文稿组件示例名称', 'block_ids': ['title', 'summary']}],
        'sources': [], 'formulas': [], 'figures': [],
        'formula_policy': {'source_inventory_complete': True, 'source_main_count': 0,
                           'non_omml_formulas': 'none', 'no_formula_reason': '非技术排版样例，无公式来源'},
        'sections': [
            {'kind': 'abstract', 'blocks': [{'id': 'summary', 'type': 'paragraph', 'role': 'body',
                                          'text': '文稿组件示例名称用于展示段落格式。此页仅为排版测试。'}]},
            {'kind': 'abstract_figure', 'blocks': []},
            {'kind': 'claims', 'blocks': [
                {'id': 'claim1', 'type': 'paragraph', 'role': 'body', 'text': '1.此处是非技术段落示例，仅核对字体与首行缩进；'},
                {'id': 'claim1b', 'type': 'paragraph', 'role': 'body', 'text': '此处是同项后续自然段，不构成专利技术内容。'}]},
            {'kind': 'description', 'blocks': [
                {'id': 'title', 'type': 'paragraph', 'role': 'project_title', 'text': '文稿组件示例名称'},
                {'id': 'heading', 'type': 'paragraph', 'role': 'heading', 'text': '技术领域'},
                {'id': 'body', 'type': 'paragraph', 'role': 'body', 'runs': [
                    {'type': 'text', 'text': '本页检查普通文字和孤立字符'},
                    {'type': 'symbol_reference', 'text': 'g'},
                    {'type': 'text', 'text': '的正体与基线显示。'}]}]},
            {'kind': 'figures', 'blocks': []}
        ]
    }


MATH = '''<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
 <m:sSub><m:e><m:r><m:rPr><m:scr m:val="script"/><m:sty m:val="b"/></m:rPr>
 <w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><w:sz w:val="28"/></w:rPr><m:t>ℬ</m:t></m:r></m:e>
 <m:sub><m:r><m:t>g</m:t></m:r></m:sub></m:sSub>
 <m:r><m:t>=</m:t></m:r><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e>
 <m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>'''


class ToolkitTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='authoring-中文-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.spec = minimal_spec()

    def save_spec(self, spec=None):
        p = self.root / '内容.json'
        p.write_text(json.dumps(spec or self.spec, ensure_ascii=False), encoding='utf-8')
        return p

    def build(self, out='v1'):
        self.save_spec()
        return toolkit.build(self.root, '内容.json', out)

    def add_math(self):
        document = Document()
        document.add_paragraph('仅用于检查源数学对象，不表示发明方案。')
        document.add_paragraph()._p.append(ET.fromstring(MATH))
        source = self.root / '原始公式.docx'
        document.save(source)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'].append(inv['source'])
        row = inv['omml_objects'][0]
        row.pop('linear_text_for_navigation_only')
        row['category'] = 'main_formula'
        self.spec['formulas'].append(row)
        self.spec['formula_policy']['source_main_count'] = 1
        self.spec['sections'][3]['blocks'].append({'id': 'equation', 'type': 'formula', 'formula_id': row['id']})
        return source

    def add_figure(self):
        image = self.root / '图.png'
        Image.new('RGB', (600, 240), 'white').save(image)
        self.spec['sources'].append({'id': 'figure_source', 'file': image.name, 'sha256': toolkit.sha(image)})
        self.spec['figures'].append({'id': 'fig1', 'source_id': 'figure_source', 'sha256': toolkit.sha(image),
                                     'number': 1, 'title': '非技术版式示例图', 'width_cm': 12})
        self.spec['sections'][1]['blocks'].append({'id': 'abstract_fig', 'type': 'figure', 'figure_id': 'fig1'})
        self.spec['sections'][3]['blocks'].append({'id': 'figure_description', 'type': 'figure_description', 'figure_id': 'fig1'})
        self.spec['sections'][3]['blocks'].append({'id': 'figure_reference', 'type': 'paragraph', 'role': 'body',
                                                   'runs': [{'type': 'text', 'text': '请参考'},
                                                            {'type': 'figure_ref', 'figure_id': 'fig1'},
                                                            {'type': 'text', 'text': '，此图仅为测试夹具。'}]})
        self.spec['sections'][4]['blocks'].append({'id': 'final_fig', 'type': 'figure', 'figure_id': 'fig1'})
        return image

    def rewrite_docx(self, path, modify):
        with zipfile.ZipFile(path) as archive:
            parts = {n: archive.read(n) for n in archive.namelist()}
        modify(parts)
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

    def test_no_formulas_five_sections_and_real_headers_footers(self):
        result = self.build()
        actual = toolkit.verify_bundle(self.root, 'v1')
        self.assertEqual(result['counts']['omml_occurrences'], 0)
        self.assertEqual(actual['counts']['sections'], 5)
        self.assertEqual([s['header_text'] for s in actual['sections']], list(toolkit.HEADERS))
        self.assertEqual([s['page_field'] for s in actual['sections']], [False, False, False, True, True])
        self.assertEqual(actual['semantic_review'], 'not_run')
        inspection = json.loads(Path(result['inspection']['inspection']).read_text(encoding='utf-8'))
        self.assertEqual(inspection['sha256'], result['output_sha256'])
        self.assertEqual(inspection['input'], str(self.root / 'v1' / 'draft.docx'))

    def test_plain_text_fonts_indents_claim_continuation_and_protected_title(self):
        self.build()
        document = Document(self.root / 'v1' / 'draft.docx')
        title = next(p for p in document.paragraphs if p.text == self.spec['project_title'])
        self.assertTrue(title.runs[0].bold)
        self.assertIn('组件', title.text)
        self.assertNotIn('模块', title.text)
        for paragraph in document.paragraphs:
            if paragraph.text.startswith(('1.', '此处是同项')):
                ind = paragraph._p.find('w:pPr/w:ind', toolkit.NS)
                self.assertEqual(ind.get(toolkit.qn('w:firstLineChars')), '200')
                spacing = paragraph._p.find('w:pPr/w:spacing', toolkit.NS)
                self.assertEqual(spacing.get(toolkit.qn('w:line')), '240')
        symbol = next(r for p in document.paragraphs for r in p.runs if r.text == 'g')
        self.assertEqual(symbol.font.name, '宋体')
        self.assertEqual(symbol.font.size.pt, 14)
        self.assertFalse(symbol.italic)
        self.assertFalse(symbol.font.cs_italic)
        self.assertFalse(symbol.bold)
        self.assertEqual(symbol._r.find('w:rPr/w:vertAlign', toolkit.NS).get(toolkit.qn('w:val')), 'baseline')

    def test_footer_reference_stack_survives_and_rejects_collapsed_or_shifted_geometry(self):
        self.build()
        original = self.root / 'v1' / 'draft.docx'
        context = toolkit.validate(self.root, self.spec)
        with zipfile.ZipFile(original) as archive:
            for i, expected_count in enumerate((2, 1, 2, 3, 3), 1):
                footer = ET.fromstring(archive.read(f'word/footer{i}.xml'))
                self.assertEqual(len(footer.findall('w:p', toolkit.NS)), expected_count)
                for anchor in footer.xpath('.//wp:anchor', namespaces=toolkit.NS):
                    self.assertEqual(anchor.find('wp:positionH/wp:posOffset', toolkit.NS).text, '12700')
        for mutation in ('collapse', 'spacer_height', 'tail_font', 'line_x', 'line_y', 'distance', 'footer_style', 'abstract_tail'):
            candidate = self.root / (mutation + '.docx')
            candidate.write_bytes(original.read_bytes())
            def change(parts):
                footer = ET.fromstring(parts['word/footer4.xml'])
                paragraphs = footer.findall('w:p', toolkit.NS)
                if mutation == 'collapse':
                    footer.remove(paragraphs[-1])
                elif mutation == 'spacer_height':
                    paragraphs[1].find('w:pPr/w:spacing', toolkit.NS).set(toolkit.qn('w:line'), '180')
                elif mutation == 'tail_font':
                    paragraphs[-1].find('w:pPr/w:rPr/w:sz', toolkit.NS).set(toolkit.qn('w:val'), '28')
                elif mutation.startswith('line_'):
                    axis = 'H' if mutation == 'line_x' else 'V'
                    footer.find(f'.//wp:position{axis}/wp:posOffset', toolkit.NS).text = '0'
                elif mutation == 'footer_style':
                    styles = ET.fromstring(parts['word/styles.xml'])
                    styles.find('w:style[@w:styleId="Footer"]/w:rPr/w:sz', toolkit.NS).set(toolkit.qn('w:val'), '28')
                    parts['word/styles.xml'] = ET.tostring(styles)
                elif mutation == 'abstract_tail':
                    abstract = ET.fromstring(parts['word/footer1.xml'])
                    abstract.findall('w:p', toolkit.NS)[-1].find('w:pPr/w:spacing', toolkit.NS).set(toolkit.qn('w:line'), '240')
                    parts['word/footer1.xml'] = ET.tostring(abstract)
                else:
                    doc = ET.fromstring(parts['word/document.xml'])
                    doc.findall('.//w:sectPr', toolkit.NS)[3].find('w:pgMar', toolkit.NS).set(toolkit.qn('w:footer'), '227')
                    parts['word/document.xml'] = ET.tostring(doc)
                parts['word/footer4.xml'] = ET.tostring(footer)
            self.rewrite_docx(candidate, change)
            with self.subTest(mutation=mutation), self.assertRaises(toolkit.AuthoringError):
                toolkit.verify_document(context, candidate)

    def test_source_script_bold_subscript_superscript_and_unicode_survive(self):
        source = self.add_math()
        source_sha = toolkit.sha(source)
        result = self.build()
        verified = toolkit.verify_bundle(self.root, 'v1')
        self.assertEqual(result['counts']['main_formula_occurrences'], 1)
        self.assertEqual(verified['source_objects'][0]['tree_sha256'], self.spec['formulas'][0]['tree_sha256'])
        self.assertEqual(toolkit.sha(source), source_sha)
        with zipfile.ZipFile(self.root / 'v1' / 'draft.docx') as archive:
            root = ET.fromstring(archive.read('word/document.xml'))
        math = root.find('.//m:oMath', toolkit.NS)
        self.assertIsNotNone(math.find('m:sSub/m:e/m:r/m:rPr/m:scr', toolkit.NS))
        self.assertEqual(math.find('m:sSub/m:e/m:r/m:rPr/m:sty', toolkit.NS).get(toolkit.qn('m:val')), 'b')
        self.assertEqual(math.find('m:sSub/m:sub/m:r/m:t', toolkit.NS).text, 'g')
        self.assertEqual(math.find('m:sSup/m:sup/m:r/m:t', toolkit.NS).text, '2')
        p = math.getparent()
        self.assertEqual(p.find('w:pPr/w:ind', toolkit.NS).get(toolkit.qn('w:firstLineChars')), '0')

    def test_upright_display_and_complete_reference_preserve_source_and_reject_slant_tamper(self):
        source = self.add_math()
        def source_styles(parts):
            root = ET.fromstring(parts['word/document.xml'])
            formula = root.find('.//m:oMath', toolkit.NS)
            formula.find('.//m:sty', toolkit.NS).set(toolkit.qn('m:val'), 'bi')
            parent_props = ET.Element(toolkit.qn('w:pPr'))
            parent_run = ET.SubElement(parent_props, toolkit.qn('w:rPr'))
            ET.SubElement(parent_run, toolkit.qn('w:i')).set(toolkit.qn('w:val'), '1')
            formula.getparent().insert(0, parent_props)
            for run in formula.xpath('.//m:r', namespaces=toolkit.NS):
                props = run.find('w:rPr', toolkit.NS)
                if props is None:
                    props = ET.Element(toolkit.qn('w:rPr'))
                    run.insert(1 if run.find('m:rPr', toolkit.NS) is not None else 0, props)
                for name in ('i', 'iCs'):
                    ET.SubElement(props, toolkit.qn('w:' + name)).set(toolkit.qn('w:val'), '1')
            parts['word/document.xml'] = ET.tostring(root)
        self.rewrite_docx(source, source_styles)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'][0] = inv['source']
        self.spec['formulas'][0]['tree_sha256'] = inv['omml_objects'][0]['tree_sha256']
        self.spec['sections'][3]['blocks'].append(dict(id='upright_ref', type='paragraph', role='body', runs=[
            dict(type='text', text='其中，'),
            dict(type='source_symbol', formula_id='f1', node_path='./m:sSub[1]'),
            dict(type='text', text='表示原式中的完整花体量。')]))
        before = source.read_bytes()
        result = self.build()
        context = toolkit.validate(self.root, self.spec)
        path = self.root / 'v1' / 'draft.docx'
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read('word/document.xml'))
        math = root.findall('.//m:oMath', toolkit.NS)
        self.assertEqual(len(math), 2)
        for node, expected in zip(math, ('ℬg=x2', 'ℬg')):
            self.assertEqual(''.join(node.xpath('.//m:t/text()', namespaces=toolkit.NS)), expected)
            self.assertEqual(node.find('.//m:scr', toolkit.NS).get(toolkit.qn('m:val')), 'script')
            self.assertEqual(node.find('.//m:sty', toolkit.NS).get(toolkit.qn('m:val')), 'b')
            for run in node.xpath('.//m:r', namespaces=toolkit.NS):
                self.assertIn(run.find('m:rPr/m:sty', toolkit.NS).get(toolkit.qn('m:val')), {'p', 'b'})
                for name in ('i', 'iCs'):
                    self.assertEqual(run.find('w:rPr/w:' + name, toolkit.NS).get(toolkit.qn('w:val')), '0')
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(toolkit.verify_bundle(self.root, 'v1')['source_objects'][0]['source_tree_sha256'],
                         inv['omml_objects'][0]['tree_sha256'])
        original = path.read_bytes()
        for index in (0, 1):
            for prop, value in (('m:sty', 'i'), ('m:sty', 'bi'), ('w:i', '1'), ('w:iCs', '1')):
                with self.subTest(index=index, prop=prop, value=value):
                    path.write_bytes(original)
                    def corrupt(parts):
                        doc = ET.fromstring(parts['word/document.xml'])
                        node = doc.findall('.//m:oMath', toolkit.NS)[index]
                        node.find('.//' + prop, toolkit.NS).set(toolkit.qn(prop[0] + ':val'), value)
                        parts['word/document.xml'] = ET.tostring(doc)
                    self.rewrite_docx(path, corrupt)
                    with self.assertRaisesRegex(toolkit.AuthoringError, 'content/structure changed'):
                        toolkit.verify_document(context, path)

    def test_unicode_italic_and_dedicated_italic_font_require_reviewed_mapping(self):
        for glyph in ('\U0001d44e', '\U0001d468', '\U0001d71c', '\U0001d656', '\u210e'):
            source = ET.fromstring(MATH.replace('ℬ', glyph))
            before = toolkit.tree_signature(source)
            with self.subTest(glyph=glyph), self.assertRaisesRegex(toolkit.AuthoringError, 'Unicode italic'):
                toolkit.upright_math_display(source)
            self.assertEqual(toolkit.tree_signature(source), before)
        source = ET.fromstring(MATH.replace('Cambria Math', 'Times New Roman Italic'))
        with self.assertRaisesRegex(toolkit.AuthoringError, 'upright font mapping'):
            toolkit.upright_math_display(source)
        path = self.add_math()
        def italic_math_font(parts):
            settings = ET.fromstring(parts['word/settings.xml'])
            props = settings.find('m:mathPr', toolkit.NS)
            if props is None:
                props = ET.SubElement(settings, toolkit.qn('m:mathPr'))
            font = props.find('m:mathFont', toolkit.NS)
            if font is None:
                font = ET.SubElement(props, toolkit.qn('m:mathFont'))
            font.set(toolkit.qn('m:val'), 'Example Math Oblique')
            parts['word/settings.xml'] = ET.tostring(settings)
        self.rewrite_docx(path, italic_math_font)
        self.spec['sources'][0]['sha256'] = toolkit.sha(path)
        with self.assertRaisesRegex(toolkit.AuthoringError, 'mathPr has a dedicated italic font'):
            toolkit.validate(self.root, self.spec)

    def test_upright_display_preserves_consistent_bold_and_rejects_conflicting_weight(self):
        for math_style, word_style in [('p', '1'), ('i', '1'), ('b', '0'), ('bi', '0')]:
            for prop in ('b', 'bCs'):
                source = ET.fromstring(MATH)
                source.find('.//m:sty', toolkit.NS).set(toolkit.qn('m:val'), math_style)
                props = source.find('.//w:rPr', toolkit.NS)
                ET.SubElement(props, toolkit.qn('w:' + prop)).set(toolkit.qn('w:val'), word_style)
                before = toolkit.tree_signature(source)
                with self.subTest(math_style=math_style, prop=prop, word_style=word_style):
                    with self.assertRaisesRegex(toolkit.AuthoringError, 'reviewed weight resolution'):
                        toolkit.upright_math_display(source)
                    self.assertEqual(toolkit.tree_signature(source), before)
        for math_style in (None, 'b', 'bi'):
            source = ET.fromstring(MATH)
            style = source.find('.//m:sty', toolkit.NS)
            if math_style is None:
                style.getparent().remove(style)
            else:
                style.set(toolkit.qn('m:val'), math_style)
            ET.SubElement(source.find('.//w:rPr', toolkit.NS), toolkit.qn('w:b')).set(toolkit.qn('w:val'), '1')
            result = toolkit.upright_math_display(source)
            self.assertEqual(result.find('.//m:sty', toolkit.NS).get(toolkit.qn('m:val')), 'b')
            self.assertEqual(result.find('.//w:b', toolkit.NS).get(toolkit.qn('w:val')), '1')
            self.assertEqual(result.find('.//m:scr', toolkit.NS).get(toolkit.qn('m:val')), 'script')

    def test_figures_reference_caption_and_abstract_reuse(self):
        self.add_figure()
        self.build()
        readback = toolkit.verify_bundle(self.root, 'v1')
        self.assertEqual(readback['counts']['image_occurrences_including_abstract_reuse'], 2)
        self.assertEqual(len(readback['source_objects']), 2)
        self.assertEqual([r['abstract_reuse'] for r in readback['source_objects']], [True, False])
        figures = [b for b in readback['blocks'] if b['type'] == 'figure']
        self.assertIsNone(figures[0]['caption_text'])
        self.assertIsNone(figures[0]['caption_locator'])
        self.assertEqual(figures[1]['caption_text'], '图1')

    def test_embedded_source_docx_media_can_be_reused_without_reencoding(self):
        image = self.add_figure()
        source = self.root / '图源.docx'
        document = Document()
        document.add_picture(str(image))
        document.save(source)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'] = [dict(inv['source'], id='figure_source')]
        self.spec['figures'][0]['media_part'] = inv['media'][0]['part']
        self.spec['figures'][0]['sha256'] = inv['media'][0]['sha256']
        self.build()
        self.assertEqual(toolkit.verify_bundle(self.root, 'v1')['source_objects'][0]['source_media'], inv['media'][0]['part'])

    def test_unknown_keys_and_inline_xml_are_rejected(self):
        self.spec['sections'][3]['blocks'][2]['runs'].append({'type': 'raw_xml', 'xml': '<m:oMath/>'})
        with self.assertRaisesRegex(toolkit.AuthoringError, 'Unsupported run'):
            self.build()
        self.assertFalse((self.root / 'v1').exists())

    def test_main_formula_must_have_exactly_one_target(self):
        self.add_math()
        self.spec['sections'][3]['blocks'].append({'id': 'copy', 'type': 'formula', 'formula_id': 'f1'})
        with self.assertRaisesRegex(toolkit.AuthoringError, 'one target'):
            self.build()

    def add_repeated_source_equation(self):
        source = self.add_math()
        document = Document(source)
        document.add_paragraph('后面的图说明再次引用同一关系。')
        document.add_paragraph()._p.append(ET.fromstring(MATH))
        document.save(source)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'] = [inv['source']]
        self.spec['formulas'] = []
        for row in inv['omml_objects']:
            row.pop('linear_text_for_navigation_only')
            row['category'] = 'main_formula'
            self.spec['formulas'].append(row)
        return source

    def test_repeated_source_alias_preserves_evidence_with_only_one_target(self):
        source = self.add_repeated_source_equation()
        canonical, alias = self.spec['formulas']
        alias.update(duplicate_of=canonical['id'], duplicate_reason='同一来源同一作用域的同一完整关系在图说明中重列。')
        source_sha = toolkit.sha(source)
        result = self.build()
        verified = toolkit.verify_bundle(self.root, 'v1')
        self.assertEqual(result['counts']['main_formula_occurrences'], 1)
        self.assertEqual(verified['counts']['omml_occurrences'], 1)
        self.assertEqual(len(verified['source_formula_aliases']), 1)
        self.assertEqual(verified['source_formula_aliases'][0]['source_locator'], alias['locator'])
        self.assertEqual(len(verified['source_formula_aliases'][0]['target_occurrences']), 1)
        self.assertEqual(toolkit.sha(source), source_sha)

    def test_repeated_source_requires_explicit_identity_and_no_second_target(self):
        self.add_repeated_source_equation()
        original = deepcopy(self.spec)
        for fault in ('unmapped', 'output_alias', 'wrong_count', 'cycle', 'missing_reason', 'missing_target', 'symbol_disguise'):
            with self.subTest(fault=fault):
                self.spec = deepcopy(original)
                canonical, alias = self.spec['formulas']
                alias.update(duplicate_of=canonical['id'], duplicate_reason='同一关系重复出现')
                if fault == 'unmapped':
                    alias.pop('duplicate_of'); alias.pop('duplicate_reason')
                    self.spec['formula_policy']['source_main_count'] = 2
                elif fault == 'output_alias':
                    self.spec['sections'][3]['blocks'].append({'id':'again','type':'formula','formula_id':alias['id']})
                elif fault == 'wrong_count':
                    self.spec['formula_policy']['source_main_count'] = 2
                elif fault == 'cycle':
                    canonical.update(duplicate_of=alias['id'], duplicate_reason='不合法环')
                elif fault == 'missing_reason':
                    alias['duplicate_reason'] = ''
                elif fault == 'missing_target':
                    self.spec['sections'][3]['blocks'].pop()
                else:
                    alias.pop('duplicate_of'); alias.pop('duplicate_reason')
                    alias['category'] = 'symbol_reference'
                with self.assertRaises(toolkit.AuthoringError):
                    toolkit.validate(self.root, self.spec)

    def test_different_source_equation_cannot_be_aliased(self):
        source = self.add_repeated_source_equation()
        def change(parts):
            doc = ET.fromstring(parts['word/document.xml'])
            nodes = doc.findall('.//m:oMath', toolkit.NS)
            nodes[1].find('m:sSup/m:sup/m:r/m:t', toolkit.NS).text = '3'
            parts['word/document.xml'] = ET.tostring(doc)
        self.rewrite_docx(source, change)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'] = [inv['source']]
        self.spec['formulas'] = []
        for row in inv['omml_objects']:
            row.pop('linear_text_for_navigation_only'); row['category'] = 'main_formula'
            self.spec['formulas'].append(row)
        self.spec['formulas'][1].update(duplicate_of=self.spec['formulas'][0]['id'], duplicate_reason='错误地合并不同指数')
        with self.assertRaisesRegex(toolkit.AuthoringError, 'matching complete'):
            toolkit.validate(self.root, self.spec)

    def test_unmapped_source_formula_is_not_silently_ignored(self):
        self.add_math()
        self.spec['formulas'] = []
        with self.assertRaisesRegex(toolkit.AuthoringError, 'Every source oMath'):
            self.build()

    def test_declared_non_omml_formula_requires_manual_route(self):
        self.spec['formula_policy']['non_omml_formulas'] = 'present'
        with self.assertRaisesRegex(toolkit.AuthoringError, 'manual route'):
            self.build()

    def test_protected_name_declared_location_is_checked(self):
        self.spec['sections'][0]['blocks'][0]['text'] = '错误地改成文稿模块示例名称。'
        with self.assertRaisesRegex(toolkit.AuthoringError, 'Protected exact name'):
            self.build()

    def test_noncontiguous_figure_number_and_missing_description_fail(self):
        self.add_figure()
        self.spec['figures'][0]['number'] = 2
        with self.assertRaisesRegex(toolkit.AuthoringError, 'contiguous'):
            self.build()
        self.spec['figures'][0]['number'] = 1
        self.spec['sections'][3]['blocks'] = [b for b in self.spec['sections'][3]['blocks'] if b['type'] != 'figure_description']
        with self.assertRaisesRegex(toolkit.AuthoringError, 'figure_description'):
            self.build()

    def test_formula_hash_and_source_hash_mismatch_fail(self):
        self.add_math()
        original = self.spec['formulas'][0]['tree_sha256']
        self.spec['formulas'][0]['tree_sha256'] = '0' * 64
        with self.assertRaisesRegex(toolkit.AuthoringError, 'tree hash'):
            self.build()
        self.spec['formulas'][0]['tree_sha256'] = original
        self.spec['sources'][0]['sha256'] = '0' * 64
        with self.assertRaisesRegex(toolkit.AuthoringError, 'Source hash'):
            self.build()

    def test_invalid_zip_and_xml_fail_without_output(self):
        source = self.root / '损坏.docx'
        for content in [b'not-a-zip', self.invalid_xml_docx()]:
            source.write_bytes(content)
            self.spec['sources'] = [{'id': 'broken', 'file': source.name, 'sha256': toolkit.sha(source)}]
            with self.subTest(content=len(content)), self.assertRaises(toolkit.AuthoringError):
                self.build()
            self.assertFalse((self.root / 'v1').exists())

    @staticmethod
    def invalid_xml_docx():
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('word/document.xml', '<w:document>broken')
        return stream.getvalue()

    def test_outside_case_paths_and_absolute_paths_are_rejected(self):
        self.save_spec()
        for path in ['../other/内容.json', str(self.root / '内容.json'), 'C:/another/source.docx']:
            with self.subTest(path=path), self.assertRaises(toolkit.AuthoringError):
                toolkit.build(self.root, path, 'v1')
        with self.assertRaises(toolkit.AuthoringError):
            toolkit.build(self.root, '内容.json', '../other-output')
        self.spec['sources'] = [{'id': 'outside', 'file': '../source.docx', 'sha256': '0' * 64}]
        with self.assertRaises(toolkit.AuthoringError):
            self.build()

    def test_existing_output_survives_successful_and_failed_rebuild(self):
        self.build()
        def files():
            return {str(p.relative_to(self.root / 'v1')): p.read_bytes() for p in (self.root / 'v1').rglob('*') if p.is_file()}
        original = files()
        self.spec['sections'][0]['blocks'][0]['text'] += '新增。'
        with self.assertRaisesRegex(toolkit.AuthoringError, 'already exists'):
            self.build()
        self.assertEqual(original, files())
        self.spec['unknown'] = 'bad'
        with self.assertRaises(toolkit.AuthoringError):
            self.build('v2')
        self.assertFalse((self.root / 'v2').exists())
        self.assertEqual(original, files())

    def test_independent_readback_catches_builder_omission_before_commit(self):
        real = toolkit.render_document
        def omit(context, path):
            real(context, path)
            def change(parts):
                root = ET.fromstring(parts['word/document.xml'])
                node = root.xpath('.//w:t[text()="技术领域"]', namespaces=toolkit.NS)[0]
                node.text = '被篡改'
                parts['word/document.xml'] = ET.tostring(root)
            self.rewrite_docx(path, change)
        with patch.object(toolkit, 'render_document', omit), self.assertRaisesRegex(toolkit.AuthoringError, 'text changed'):
            self.build()
        self.assertFalse((self.root / 'v1' / 'build_manifest.json').exists())

    def test_readback_rejects_formula_glyph_corruption_even_with_updated_package_hash(self):
        self.add_math()
        self.build()
        path = self.root / 'v1' / 'draft.docx'
        def corrupt(parts):
            parts['word/document.xml'] = parts['word/document.xml'].replace('ℬ'.encode(), b'B')
        self.rewrite_docx(path, corrupt)
        manifest_file = self.root / 'v1' / 'build_manifest.json'
        manifest = json.loads(manifest_file.read_text(encoding='utf-8'))
        manifest['output_sha256'] = toolkit.sha(path)
        manifest_file.write_text(json.dumps(manifest), encoding='utf-8')
        context = toolkit.validate(self.root, self.spec)
        with self.assertRaisesRegex(toolkit.AuthoringError, 'source OMML content/structure changed'):
            toolkit.verify_document(context, path)

    def test_readback_rejects_unregistered_table_claim_reordering_and_page_font(self):
        self.build()
        original=self.root/'v1/draft.docx'
        context=toolkit.validate(self.root,self.spec)
        for mutation in ('table','order','page_font'):
            candidate=self.root/(mutation+'.docx');candidate.write_bytes(original.read_bytes())
            def change(parts):
                doc=ET.fromstring(parts['word/document.xml']);body=doc.find('w:body',toolkit.NS)
                if mutation=='table':
                    body.insert(0,toolkit.element('w:tbl'))
                elif mutation=='order':
                    claims=[p for p in body.findall('w:p',toolkit.NS) if toolkit.word_text(p).startswith(('1.','此处是同项'))]
                    body.remove(claims[1]);body.insert(body.index(claims[0]),claims[1])
                else:
                    foot=ET.fromstring(parts['word/footer4.xml'])
                    foot.find('.//w:r/w:rPr/w:sz',toolkit.NS).set(toolkit.qn('w:val'),'48')
                    parts['word/footer4.xml']=ET.tostring(foot)
                parts['word/document.xml']=ET.tostring(doc)
            self.rewrite_docx(candidate,change)
            with self.subTest(mutation=mutation),self.assertRaises(toolkit.AuthoringError):
                toolkit.verify_document(context,candidate)

    def test_readback_rejects_moving_inline_math_across_unchanged_text(self):
        self.add_math()
        formula=self.spec['sections'][3]['blocks'].pop()['formula_id']
        self.spec['sections'][3]['blocks'].append({'id':'inline','type':'paragraph','role':'body','runs':[
            {'type':'text','text':'前文'}, {'type':'formula','formula_id':formula}, {'type':'text','text':'后文'}]})
        self.build()
        candidate=self.root/'v1/draft.docx'
        def change(parts):
            doc=ET.fromstring(parts['word/document.xml']);node=doc.find('.//m:oMath',toolkit.NS)
            parent=node.getparent();parent.remove(node);parent.append(node)
            parts['word/document.xml']=ET.tostring(doc)
        self.rewrite_docx(candidate,change)
        with self.assertRaisesRegex(toolkit.AuthoringError,'interleaving'):
            toolkit.verify_document(toolkit.validate(self.root,self.spec),candidate)

    def test_builder_inspects_once_and_next_handoff_reuses_cache(self):
        with patch.object(toolkit.inspect_docx, 'inspect', wraps=toolkit.inspect_docx.inspect) as inspector:
            result = self.build()
            self.assertEqual(inspector.call_count, 1)
            again = toolkit.inspect_docx.run_inspection(self.root / 'v1' / 'draft.docx', self.root / 'v1' / 'inspection-cache',
                                                        'draft', True, True)
            self.assertEqual(inspector.call_count, 1)
        self.assertEqual(again['cache_status'], 'hit')
        self.assertEqual(again['inspection'], result['inspection']['inspection'])

    def test_bundle_rejects_damaged_extracted_inspection_media(self):
        self.add_figure()
        result = self.build()
        cache_dir = Path(result['inspection']['output_directory'])
        media = next((cache_dir / 'media').iterdir())
        media.write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, 'Cache artifact changed'):
            toolkit.verify_bundle(self.root, 'v1')

    def test_duplicate_json_keys_fail(self):
        (self.root / '内容.json').write_text('{"schema_version":1,"schema_version":2}', encoding='utf-8')
        with self.assertRaisesRegex(toolkit.AuthoringError, 'Duplicate JSON key'):
            toolkit.build(self.root, '内容.json', 'v1')

    def test_math_style_dependency_is_rejected_instead_of_lost(self):
        self.add_math()
        source = self.root / '原始公式.docx'
        def change(parts):
            root = ET.fromstring(parts['word/document.xml'])
            node = root.find('.//m:oMath/m:sSub/m:e/m:r/w:rPr', toolkit.NS)
            node.append(toolkit.element('w:rStyle', val='SpecialVector'))
            parts['word/document.xml'] = ET.tostring(root)
        self.rewrite_docx(source, change)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'][0] = inv['source']
        self.spec['formulas'][0]['tree_sha256'] = inv['omml_objects'][0]['tree_sha256']
        with self.assertRaisesRegex(toolkit.AuthoringError, 'manual source-preserving route'):
            self.build()

    def test_colored_source_math_is_not_silently_recolored_or_accepted(self):
        self.add_math()
        source = self.root / '原始公式.docx'
        def change(parts):
            root = ET.fromstring(parts['word/document.xml'])
            node = root.find('.//m:oMath/m:sSub/m:e/m:r/w:rPr', toolkit.NS)
            node.append(toolkit.element('w:color', val='FF0000'))
            parts['word/document.xml'] = ET.tostring(root)
        self.rewrite_docx(source, change)
        inv = toolkit.inventory(self.root, source.name)
        self.spec['sources'][0] = inv['source']
        self.spec['formulas'][0]['tree_sha256'] = inv['omml_objects'][0]['tree_sha256']
        with self.assertRaisesRegex(toolkit.AuthoringError, 'nonblack text'):
            self.build()


if __name__ == '__main__':
    unittest.main()
