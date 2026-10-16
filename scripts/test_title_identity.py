"""Official title identity across claims and their mirrored solution."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.oxml import parse_xml
import authoring_toolkit as toolkit
import batch_tool as batch
import inspect_docx as scan
import invention_contract as inv
import output_contract as final
from test_authoring_toolkit import minimal_spec, MATH


class TitleIdentityTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.path=self.root/'case.docx'

    def document(self,title,name=None,extra=None,cross_run=False):
        name=title if name is None else name
        claims=[f'1.{name}，其特征在于，包括：获取记录，保留所述记录；',
                '仅在所述记录有效时输出所述记录，否则不参与更新。',
                f'2.根据权利要求1所述的{name}，其特征在于，保留所述记录的原标识。']
        solution=[inv.claim_head(p)[2] if p[0].isdigit() else p for p in claims]
        doc=Document()
        for line in ['说明书摘要','内容摘要。','权利要求书',*claims,title,'技术领域','记录处理。',
                     '背景技术','背景。','发明内容',inv.purpose(title),inv.TRANSITION,*solution,
                     inv.BENEFITS,'1.保留原标识。','具体实施方式',extra or '采用一种处理方式保存记录。']:
            p=doc.add_paragraph()
            for part in (line if cross_run else [line]):p.add_run(part)
        return doc

    def inspect(self,doc):
        doc.save(self.path);before=self.path.read_bytes()
        report=scan.inspect(self.path);scan.precheck(report,'draft')
        self.assertEqual(before,self.path.read_bytes())
        return report

    def test_same_wrong_prefix_in_claim_and_solution_cannot_pass(self):
        title='空分压缩预冷与分馏的约束协调过程控制方法'
        r=self.inspect(self.document(title,'一种'+title,cross_run=True))
        hits=[i for i in r['invention_evidence']['issues'] if i['code']=='title_reference_prefix']
        self.assertEqual(len(hits),3)  # independent, dependent head, independent solution
        self.assertTrue(all(h['expected']==title and h['actual']=='一种'+title for h in hits))
        with self.assertRaisesRegex(batch.BatchError,'title_reference_prefix'):
            final.check(r,title,strict_invention=True)
        final.check(r,title)  # Frozen legacy contract has no new evidence requirement.

    def test_original_prefix_preserved_missing_and_repeated_prefix_rejected(self):
        title='一种记录处理方法'
        final.check(self.inspect(self.document(title)),title,strict_invention=True)
        for name in ['记录处理方法','一种'+title,'一种一种'+title]:
            with self.subTest(name=name),self.assertRaisesRegex(batch.BatchError,'title_reference_prefix'):
                final.check(self.inspect(self.document(title,name)),title,strict_invention=True)

    def test_valid_name_without_prefix_and_ordinary_classifier_unchanged(self):
        title='记录处理方法'
        r=self.inspect(self.document(title,extra='本发明采用一种处理方式，另一种模块保留记录。'))
        final.check(r,title,strict_invention=True)
        self.assertIsNone(inv.name_prefix_issue('一种记录处理方法扩展系统，用于记录。',title))

    def test_name_slot_in_abstract_purpose_or_implementation_is_checked(self):
        title='记录处理方法'
        for text in ['本发明提供一种'+title+'，包括步骤。','本发明公开一种'+title+'。',
                     inv.purpose('一种'+title),'一种'+title+'，包括步骤。']:
            with self.subTest(text=text),self.assertRaisesRegex(batch.BatchError,'title_reference_prefix'):
                final.check(self.inspect(self.document(title,extra=text)),title,strict_invention=True)

    def test_only_solution_changed_still_rejected_by_mirroring(self):
        title='记录处理方法';doc=self.document(title)
        p=next(p for p in doc.paragraphs if p.text.startswith(title+'，'))
        p.text='一种'+p.text
        r=self.inspect(doc)
        self.assertIn('solution_text_changed',[i['code'] for i in r['invention_evidence']['issues']])

    def spec(self,title,name=None):
        spec=minimal_spec();spec['project_title']=title;spec['protected_names']=[]
        spec['sections'][3]['blocks'][0]['text']=title
        spec['sections'][2]['blocks']=[dict(id='claim',type='paragraph',role='body',
            text='1.'+(name or title)+'，其特征在于，包括：仅保存所述记录，否则不参与更新。')]
        return spec

    def test_generator_rejects_both_manual_and_auto_mirror_routes(self):
        for auto in (False,True):
            spec=self.spec('记录处理方法','一种记录处理方法')
            if auto:
                spec['invention_from_claims']=True
                spec['sections'][3]['blocks'] += [dict(id='iv',type='paragraph',role='heading',text='发明内容'),
                    dict(id='be',type='paragraph',role='body',text=inv.BENEFITS)]
            original=deepcopy(spec)
            with self.subTest(auto=auto),self.assertRaisesRegex(toolkit.AuthoringError,'title_reference_prefix'):
                toolkit.validate(self.root,spec)
            self.assertEqual(spec,original)

    def test_fixed_name_generates_without_rewriting_conditions(self):
        for title in ['记录处理方法','一种记录处理方法']:
            spec=self.spec(title);spec['invention_from_claims']=True
            spec['sections'][3]['blocks'] += [dict(id='iv',type='paragraph',role='heading',text='发明内容'),
                dict(id='be',type='paragraph',role='body',text=inv.BENEFITS)]
            original=deepcopy(spec);context=toolkit.validate(self.root,spec)
            self.assertEqual(spec,original)
            self.assertIn(title+'，包括：仅保存所述记录，否则不参与更新。',context['block_texts'].values())
            toolkit.render_document(context,self.path);toolkit.verify_document(context,self.path)
            r=scan.inspect(self.path);scan.precheck(r,'draft');final.check(r,title,strict_invention=True)

    def test_multiple_categories_not_forced_to_composite_title(self):
        title='记录处理方法、系统及存储介质'
        for name in ['记录处理方法','一种记录系统','存储介质']:
            with self.subTest(name=name):
                final.check(self.inspect(self.document(title,name)),title,strict_invention=True)

    def test_math_bytes_and_ordinary_prose_not_rewritten_by_inspection(self):
        title='记录处理方法';doc=self.document(title)
        doc.paragraphs[-1]._p.append(parse_xml(MATH))
        before=doc.paragraphs[-1]._p.xml
        r=self.inspect(doc)
        final.check(r,title,strict_invention=True)
        self.assertEqual(doc.paragraphs[-1]._p.xml,before)
        self.assertEqual(len(r['formulas']),1)

    def test_frozen_final_path_reuses_same_inspection(self):
        title='记录处理方法';r=self.inspect(self.document(title,'一种'+title))
        prepared={'handoff':{'input_inspection':{'inspection':'sealed.json'}},'input_report':r}
        with patch('prepare_handoff.verify_prepared',return_value=prepared),patch.object(scan,'inspect',side_effect=AssertionError('duplicate parse')):
            with self.assertRaisesRegex(batch.BatchError,'title_reference_prefix'):
                final.verify({'policy':{'output_contract':final.MODE,'invention_contract':inv.MODE}},
                    {'title':title},{'input_artifact_sha256':r['sha256']},{})


if __name__=='__main__':unittest.main()
