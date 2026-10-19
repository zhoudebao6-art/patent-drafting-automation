#!/usr/bin/env python3
"""Build a case-local DOCX from explicit content; never certify its semantics.

Only source OMML can enter mathematical objects. Unsupported content fails
closed. See references/authoring-toolkit.md for the schema and review boundary.
"""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import sys
import tempfile
import unicodedata
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips, Cm, RGBColor
from lxml import etree as ET
from PIL import Image

import inspect_docx
import math_fragments
import invention_contract
import figure_description_contract
import single_line_math
import formula_uniqueness


class AuthoringError(ValueError):
    """An input or independently observed output violates the limited contract."""


NS = dict(inspect_docx.NS, wp='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
          pr='http://schemas.openxmlformats.org/package/2006/relationships')
SECTIONS = ('abstract', 'abstract_figure', 'claims', 'description', 'figures')
HEADERS = ('说明书摘要', '摘要附图', '权利要求书', '说明书', '说明书附图')
SHA = re.compile(r'^[0-9a-f]{64}$')
IDENTIFIER = re.compile(r'^[A-Za-z][A-Za-z0-9_-]{0,47}$')
LIMIT = 64 * 1024 * 1024
PROFILE = 'manual-v34-upright-equation-v8'


def fail(message):
    raise AuthoringError(message)


def digest(data):
    return hashlib.sha256(data).hexdigest()


def sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def dump(path, value):
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')


def unique_pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            fail('Duplicate JSON key: ' + key)
        result[key] = value
    return result


def read_json(path):
    try:
        if Path(path).stat().st_size > LIMIT:
            fail('JSON input is too large')
        return json.loads(Path(path).read_text(encoding='utf-8-sig'), object_pairs_hook=unique_pairs,
                          parse_constant=lambda value: fail('Non-finite JSON number: ' + value))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise AuthoringError('Cannot read JSON: ' + str(exc)) from exc


def fields(value, required, optional=(), where='object'):
    if not isinstance(value, dict):
        fail(where + ' must be an object')
    missing = set(required) - value.keys()
    unknown = value.keys() - set(required) - set(optional)
    if missing or unknown:
        fail(f'{where}: missing={sorted(missing)}, unsupported={sorted(unknown)}')


def text_value(value, where):
    if not isinstance(value, str) or not value.strip():
        fail(where + ' must be nonempty text')
    if any(ord(c) < 32 for c in value) or any(0xD800 <= ord(c) <= 0xDFFF for c in value):
        fail(where + ': control characters, tabs and embedded line breaks are unsupported; use separate paragraphs')
    return value


def visible_value(value, where):
    text_value(value,where)
    if any(c.isspace() or c in {'\u200b','\ufeff'} for c in value):
        fail(where + ': literal whitespace in visible content; rewrite prose without merging source names or changing math')
    return value


def identifier(value, where):
    if not isinstance(value, str) or not IDENTIFIER.fullmatch(value):
        fail(where + ' must be an ASCII identifier, 1..48 characters')
    return value


def sha_value(value, where):
    if not isinstance(value, str) or not SHA.fullmatch(value):
        fail(where + ' must be a lowercase SHA256')


def integer(value, where, minimum=0):
    if type(value) is not int or value < minimum:
        fail(where + ' must be an integer >= ' + str(minimum))


def bounded_path(root, relative, *, exists=True):
    if not isinstance(relative, str) or not relative:
        fail('Expected a nonempty case-relative path')
    rel = PurePosixPath(relative.replace('\\', '/'))
    if rel.is_absolute() or any(p in {'..', '.'} for p in rel.parts) or ':' in relative:
        fail('Path must be relative and cannot traverse: ' + relative)
    path = root.joinpath(*rel.parts).resolve()
    if path == root or not path.is_relative_to(root):
        fail('Path escapes the case root: ' + relative)
    if exists and not path.is_file():
        fail('Missing input file: ' + relative)
    return path


def xml(data, label):
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        fail('DTD/entities are unsupported: ' + label)
    try:
        return ET.fromstring(data, parser=ET.XMLParser(resolve_entities=False, no_network=True))
    except ET.XMLSyntaxError as exc:
        raise AuthoringError('Invalid XML in ' + label) from exc


def package(data, label):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = [e.filename for e in entries]
            if len(names) != len(set(names)) or len(names) > 10000:
                fail('Duplicate/excessive package members: ' + label)
            if sum(e.file_size for e in entries) > 4 * LIMIT or any(e.file_size > LIMIT for e in entries):
                fail('Expanded package exceeds limit: ' + label)
            for name in names:
                member = PurePosixPath(name)
                if member.is_absolute() or '..' in member.parts or '\\' in name:
                    fail('Unsafe package member: ' + name)
            parts = {e.filename: archive.read(e) for e in entries if not e.is_dir()}
            if 'word/document.xml' not in parts:
                fail('Missing word/document.xml: ' + label)
            for name, content in parts.items():
                if name.endswith(('.xml', '.rels')):
                    xml(content, name)
            return parts
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise AuthoringError('Invalid DOCX package: ' + label) from exc


def tree_signature(element):
    """Prefix/attribute-order independent; preserve every node, value and text.

    XML indentation whitespace outside leaf text is not mathematical content.
    Literal m:t and w:t whitespace remains significant, including xml:space.
    """
    def visit(node):
        if not isinstance(node.tag, str):
            fail('Comments/processing instructions in imported OMML are unsupported')
        value = node.text or ''
        if len(node) and not value.strip():
            value = ''
        tail = node.tail or ''
        if not tail.strip():
            tail = ''
        return [node.tag, sorted(node.attrib.items()), value, tail, [visit(c) for c in node]]
    return digest(json.dumps(visit(element), ensure_ascii=False, separators=(',', ':')).encode('utf-8'))


def math_nodes(parts):
    root = xml(parts['word/document.xml'], 'word/document.xml')
    # Import by stable index among all source body oMath, not flattened text.
    return root.xpath('.//m:oMath', namespaces=NS)


def inventory(case_root, source):
    root = Path(case_root).resolve(strict=True)
    path = bounded_path(root, source)
    if path.stat().st_size > LIMIT:
        fail('Source exceeds size limit')
    data = path.read_bytes()
    parts = package(data, source)
    formulas = []
    for i, node in enumerate(math_nodes(parts), 1):
        formulas.append({'id': f'f{i}', 'source_id': 'source1',
                         'locator': {'part': 'word/document.xml', 'omath_index': i},
                         'tree_sha256': tree_signature(node), 'category': None,
                         'linear_text_for_navigation_only': ''.join(node.xpath('.//m:t/text()', namespaces=NS))})
    images = [{'part': name, 'sha256': digest(data)} for name, data in parts.items()
              if name.startswith('word/media/')]
    return {'source': {'id': 'source1', 'file': source, 'sha256': digest(data)},
            'omml_objects': formulas, 'media': images,
            'boundary': 'Inventory is navigation only. Classify each object and inspect text/image/OLE formulas; null categories are not buildable. oMath count is not the semantic main-formula count.'}


def ordinary_reference_letter(text):
    """Ordinary alphabet letters only; mathematical alphabet variants stay native."""
    return bool(re.fullmatch(r'[A-Za-z\u0391-\u03a1\u03a3-\u03a9\u03b1-\u03c9]', text))


def plain_source_text(node, category, allow_greek=True):
    """Convert only an explicitly classified, flat scalar/prose example.

    This is not a math recognizer. Scripts, accents, vectors and operators
    with structural operands remain native; their linear text is insufficient.
    """
    if category not in {'symbol_reference', 'prose_example', 'prose_label'}:
        return None
    children = [n for n in node if n.tag != qn('m:oMathPr')]
    if not children or any(n.tag != qn('m:r') for n in children):
        return None
    for run in children:
        if any(n.tag not in {qn('m:rPr'), qn('w:rPr'), qn('m:t')} for n in run):
            return None
        if any(n.tag not in {qn('m:sty'), qn('m:nor')} for props in run.findall('m:rPr', NS) for n in props):
            return None
        safe_word_props = {'rFonts', 'sz', 'szCs', 'i', 'iCs', 'b', 'bCs', 'color', 'lang',
                           'vertAlign', 'position', 'spacing'}
        if any(ET.QName(n).localname not in safe_word_props for props in run.findall('w:rPr', NS) for n in props):
            return None
        if run.xpath('./m:rPr/m:scr', namespaces=NS):
            return None
        if any(n.get(qn('w:val')) != 'baseline' for n in run.findall('w:rPr/w:vertAlign', NS)):
            return None
        if any(n.get(qn('w:val')) != '0'
               for n in run.xpath('./w:rPr/w:position | ./w:rPr/w:spacing', namespaces=NS)):
            return None
        if any(n.get(qn('m:val')) not in {'p', 'i'} for n in run.findall('m:rPr/m:sty', NS)):
            return None
        if any(n.get(qn('w:val'), '1') not in {'0', 'false', 'off'}
               for n in run.xpath('./w:rPr/w:b | ./w:rPr/w:bCs', namespaces=NS)):
            return None
        if category == 'symbol_reference':
            # Symbol/custom-font encodings can change the visible identity of
            # an ordinary Unicode letter. Do not guess their rendered glyph.
            ordinary_fonts = {'Cambria Math', 'Cambria', 'Times New Roman', '宋体', 'SimSun', 'Arial', 'Calibri'}
            for fonts in run.findall('w:rPr/w:rFonts', NS):
                if any(ET.QName(key).localname not in {'ascii', 'hAnsi', 'eastAsia', 'cs', 'hint'} or
                       (ET.QName(key).localname != 'hint' and value not in ordinary_fonts)
                       for key, value in fonts.attrib.items()):
                    return None
    text = ''.join(n.text or '' for run in children for n in run.findall('m:t', NS))
    if category == 'symbol_reference':
        return text if ordinary_reference_letter(text) and (allow_greek or text.isascii()) else None
    if category == 'prose_label':
        return text if re.fullmatch(r'[A-Za-z][A-Za-z0-9]*', text) else None
    return text if re.fullmatch(r'[0-9+\-−×÷*/=.,%]+', text) and '=' in text else None


def symbol_display(formula):
    """Create an upright display copy; source identity and hashes stay immutable."""
    # A derived fragment cannot inherit a display tree for its complete parent.
    formula.pop('display_node', None)
    formula.pop('display_tree_sha256', None)
    if formula['plain_text'] is not None:
        return formula
    node = upright_math_display(formula['node'])
    # OMML scripts retain their structural scale relative to this 14pt base.
    # Bold/script identity, glyphs, operators and the source tree stay intact.
    owners = node.xpath('.//m:r | .//m:ctrlPr', namespaces=NS) if formula['category'] == 'symbol_reference' else []
    for owner in owners:
        props = owner.find('w:rPr', NS)
        if props is None:
            props = ET.Element(qn('w:rPr'))
            position = 1 if owner.find('m:rPr', NS) is not None else 0
            owner.insert(position, props)
        for name in ('sz', 'szCs'):
            prop = props.find('w:' + name, NS)
            if prop is None:
                prop = ET.SubElement(props, qn('w:' + name))
            prop.set(qn('w:val'), '28')
    formula['display_node'] = node
    formula['display_tree_sha256'] = tree_signature(node)
    return formula


def upright_math_display(source):
    """Remove only slant in a display copy, never mathematical alphabet identity.

    Math uses italic by default even when Word's w:i is false. Both layers
    must be explicit. Unicode italic alphabets and dedicated italic fonts
    require a reviewed glyph mapping; toggling attributes cannot straighten
    them, and NFKC would also destroy distinct script/bold/variant symbols.
    """
    for value in source.xpath('.//m:t/text() | .//w:t/text()', namespaces=NS):
        if any(('MATHEMATICAL' in unicodedata.name(c, '') and 'ITALIC' in unicodedata.name(c, ''))
               or c == '\u210e' for c in value):
            fail('Unicode italic mathematical glyph requires a reviewed source-bound upright mapping; '
                 'manual source-preserving route required')
    for fonts in source.xpath('.//w:rFonts', namespaces=NS):
        if any(re.search(r'italic|oblique|slanted', value, re.I) for value in fonts.attrib.values()):
            fail('Dedicated italic mathematical font requires a reviewed upright font mapping; '
                 'manual source-preserving route required')
    node = deepcopy(source)
    for owner in node.xpath('.//m:r', namespaces=NS):
        props = owner.find('m:rPr', NS)
        if props is None:
            props = ET.Element(qn('m:rPr'))
            owner.insert(0, props)
        styles = props.findall('m:sty', NS)
        if len(styles) > 1 or (styles and styles[0].get(qn('m:val')) not in {'p', 'i', 'b', 'bi'}):
            fail('Unknown mathematical style; manual source-preserving route required')
        bold = owner.xpath('./w:rPr/w:b | ./w:rPr/w:bCs', namespaces=NS)
        word_bold = {p.get(qn('w:val'), '1') not in {'0', 'false', 'off'} for p in bold}
        math_bold = styles[0].get(qn('m:val')) in {'b', 'bi'} if styles else None
        if len(word_bold) > 1 or (styles and any(value != math_bold for value in word_bold)):
            fail('Conflicting mathematical and Word bold hints require reviewed weight resolution; '
                 'manual source-preserving route required')
        if styles:
            style = styles[0]
            value = 'b' if math_bold else 'p'
        else:
            style = ET.SubElement(props, qn('m:sty'))
            value = 'b' if True in word_bold else 'p'
        style.set(qn('m:val'), value)
    for owner in node.xpath('.//m:r | .//m:ctrlPr', namespaces=NS):
        props = owner.find('w:rPr', NS)
        if props is None:
            props = ET.Element(qn('w:rPr'))
            owner.insert(1 if owner.find('m:rPr', NS) is not None else 0, props)
    for props in node.xpath('.//w:rPr', namespaces=NS):
        for name in ('i', 'iCs'):
            matches = props.findall('w:' + name, NS)
            if not matches:
                matches = [ET.SubElement(props, qn('w:' + name))]
            for prop in matches:
                prop.set(qn('w:val'), '0')
    return node


def explanation_intro(block):
    """Rewrite only the paragraph-leading introduction, including split runs."""
    if block.get('role') not in {'body', 'benefit'}:
        return
    if 'text' in block:
        block['text'] = re.sub(r'^(\s*)式中(?=[，,:：])', r'\1其中', block['text'])
        return
    prefix = []
    for run in block.get('runs', []):
        if run.get('type') != 'text':
            break
        prefix.append(run)
    text = ''.join(r.get('text', '') for r in prefix)
    found = re.match(r'^\s*式中(?=[，,:：])', text)
    if found:
        offset = text.index('式')
        for run in prefix:
            if offset < len(run['text']):
                run['text'] = run['text'][:offset] + '其' + run['text'][offset+1:]
                break
            offset -= len(run['text'])


def formula_suffix(formula):
    children = [n for n in formula['node'] if n.tag != qn('m:oMathPr')]
    # The final row of a genuine array is its display end. Internal parameter
    # punctuation remains protected by the top-level-run test below.
    while children and children[-1].tag == qn('m:eqArr'):
        rows=children[-1].findall('m:e',NS)
        children=list(rows[-1]) if rows else []
    # A punctuation character inside a delimiter or subscript is not the
    # end of the formula. Only a final top-level math run can already end it.
    text = ''
    for node in reversed(children):
        if node.tag != qn('m:r'):
            break
        text = ''.join(node.xpath('./m:t/text()', namespaces=NS)).rstrip()
        if text or any(n.tag not in {qn('m:t'), qn('m:rPr'), qn('w:rPr')} for n in node):
            break
    if text.endswith(('。', '.', '，', ',', '：', ':')):
        fail('Display formula ends with other punctuation; review source before normalizing')
    return '' if text.endswith(('；', ';')) else '；'


def source_formula(row, sources):
    fields(row, ('id', 'source_id', 'locator', 'tree_sha256', 'category'),
           ('equation_count', 'duplicate_of', 'duplicate_reason'), where='formula')
    identifier(row['id'], 'formula.id')
    if row['source_id'] not in sources:
        fail('Unknown formula source: ' + str(row['source_id']))
    source = sources[row['source_id']]
    if source['parts'] is None:
        fail('Formula source must be DOCX')
    loc = row['locator']
    fields(loc, ('part', 'omath_index'), where='formula.locator')
    if loc['part'] != 'word/document.xml':
        fail('Only word/document.xml OMML import is supported')
    integer(loc['omath_index'], 'omath_index', 1)
    nodes = source['math']
    if loc['omath_index'] > len(nodes):
        fail('OMML index out of range: ' + row['id'])
    node = nodes[loc['omath_index'] - 1]
    sha_value(row['tree_sha256'], 'formula.tree_sha256')
    if tree_signature(node) != row['tree_sha256']:
        fail('Source OMML tree hash mismatch: ' + row['id'])
    if row['category'] not in {'main_formula', 'source_example', 'symbol_reference', 'prose_example', 'prose_label'}:
        fail('Explicit formula category required: ' + row['id'])
    if 'equation_count' in row:
        integer(row['equation_count'],'equation_count',1)
        if row['category'] not in {'main_formula','source_example'}:
            fail('Only actual complete source equations have equation_count')
    if 'duplicate_of' in row:
        identifier(row['duplicate_of'], 'formula.duplicate_of')
        text_value(row.get('duplicate_reason'), 'formula.duplicate_reason')
        if row['category'] not in {'main_formula', 'source_example'}:
            fail('Only complete source equations can declare duplicate_of')
    elif 'duplicate_reason' in row:
        fail('duplicate_reason requires duplicate_of')
    if any(h['reason']=='math_cjk_requires_english_symbol' for h in inspect_docx.math_notation_candidates(node)):
        fail('Chinese mathematical text needs a reviewed source-bound English symbol mapping before import; preserve raw evidence and use the manual normalization route: '+row['id'])
    for text in node.xpath('.//m:t/text() | .//w:t/text()',namespaces=NS):
        if any(c.isspace() or c in {'\u200b','\ufeff'} for c in text):
            fail('Source mathematical whitespace needs source-preserving normalization before import: '+row['id'])
    unsupported = node.xpath('.//w:rStyle | .//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo', namespaces=NS)
    if unsupported or any(k.startswith('{' + NS['r'] + '}') or 'theme' in k.lower()
                          for n in node.iter() for k in n.attrib):
        fail('OMML depends on revisions/styles/themes/relationships; manual source-preserving route required: ' + row['id'])
    if any(n.tag in {qn('w:ins'), qn('w:del'), qn('w:moveFrom'), qn('w:moveTo')} for n in node.iterancestors()):
        fail('OMML has a revision ancestor; manual source-preserving route required: ' + row['id'])
    if any(n.get(qn('w:val'), 'auto') not in {'000000', 'auto'} for n in node.xpath('.//w:color', namespaces=NS)):
        fail('Source OMML has nonblack text; preserve source and use the reviewed manual normalization route: ' + row['id'])
    paragraph = next((n for n in node.iterancestors() if n.tag == qn('w:p')), None)
    if paragraph is not None:
        # A cloned subtree cannot carry paragraph/character style definitions.
        # Fail on observed inherited semantic dependencies rather than guessing.
        inherited = paragraph.xpath('./w:pPr/w:pStyle | ./w:pPr/w:rPr/w:b | '
                                    './w:pPr/w:rPr/w:rFonts | ./w:pPr/w:rPr/w:vertAlign', namespaces=NS)
        if inherited:
            fail('OMML has a paragraph-style/font dependency; manual source-preserving route required: ' + row['id'])
    if node.getparent() is not None and node.getparent().tag == qn('m:oMath'):
        fail('Nested oMath objects are unsupported')
    plain = plain_source_text(node, row['category'])
    if row['category'] in {'prose_example', 'prose_label'} and plain is None:
        fail('Prose example is not safely representable as flat text; use the reviewed source-preserving route: ' + row['id'])
    return symbol_display(dict(row, node=node, plain_text=plain))


def validate(case_root, spec):
    spec = deepcopy(spec)  # Generated ordinal prefixes must not mutate caller input.
    root = Path(case_root).resolve(strict=True)
    if not root.is_dir():
        fail('Case root must be a directory')
    fields(spec, ('schema_version', 'case_id', 'project_title', 'protected_names', 'sources',
                  'formula_policy', 'formulas', 'figures', 'sections'), ('invention_from_claims',), where='document')
    if type(spec['schema_version']) is not int or spec['schema_version'] != 1:
        fail('Unsupported schema_version')
    identifier(spec['case_id'], 'case_id')
    visible_value(spec['project_title'], 'project_title')
    for key in ('sources', 'protected_names', 'formulas', 'figures', 'sections'):
        if not isinstance(spec[key], list):
            fail(key + ' must be an array')
    try:invention_contract.expand(spec)
    except (ValueError,KeyError,TypeError) as exc:fail('Invention from claims: '+str(exc))
    sources = {}
    for row in spec['sources']:
        fields(row, ('id', 'file', 'sha256'), where='source')
        identifier(row['id'], 'source.id')
        if row['id'] in sources:
            fail('Duplicate source id')
        path = bounded_path(root, row['file'])
        sha_value(row['sha256'], 'source.sha256')
        if path.stat().st_size > LIMIT:
            fail('Source exceeds size limit')
        data = path.read_bytes()
        if digest(data) != row['sha256']:
            fail('Source hash mismatch: ' + row['file'])
        suffix = path.suffix.lower()
        if suffix not in {'.docx', '.png', '.jpg', '.jpeg'}:
            fail('Unsupported source type: ' + suffix)
        parts = package(data, row['file']) if suffix == '.docx' else None
        sources[row['id']] = dict(row, path=path, data=data, parts=parts,
                                  math=math_nodes(parts) if parts else [])
    if len({s['path'] for s in sources.values()}) != len(sources):
        fail('The same source file cannot be registered twice')
    policy = spec['formula_policy']
    fields(policy, ('source_inventory_complete', 'source_main_count', 'non_omml_formulas'),
           ('no_formula_reason',), 'formula_policy')
    if policy['source_inventory_complete'] is not True or policy['non_omml_formulas'] != 'none':
        fail('Complete human source inventory required; non-OMML formulas need the manual route')
    integer(policy['source_main_count'], 'source_main_count')
    formulas = {}
    for row in spec['formulas']:
        parsed = source_formula(row, sources)
        if row['id'] in formulas:
            fail('Duplicate formula id')
        formulas[row['id']] = parsed
    source_keys = [(f['source_id'], f['locator']['omath_index']) for f in formulas.values()]
    expected_keys = [(s['id'], i) for s in sources.values() for i in range(1, len(s['math']) + 1)]
    if Counter(source_keys) != Counter(expected_keys):
        fail('Every source oMath must have exactly one registry entry; do not silently drop or alias it')
    identities = {}
    for formula_id, formula in formulas.items():
        fingerprint = formula_uniqueness.equation_fingerprint(formula['node'])
        if formula['category'] == 'symbol_reference' and fingerprint is not None:
            fail('A complete equation cannot be classified as symbol_reference: ' + formula_id)
        alias = formula.get('duplicate_of')
        if alias is not None:
            canonical = formulas.get(alias)
            if (canonical is None or alias == formula_id or canonical.get('duplicate_of') is not None
                    or canonical['category'] != formula['category']
                    or canonical.get('equation_count', 1) != formula.get('equation_count', 1)
                    or fingerprint is None
                    or fingerprint != formula_uniqueness.equation_fingerprint(canonical['node'])):
                fail('duplicate_of must name one matching complete canonical source equation; unsupported identities need source review: ' + formula_id)
        elif formula['category'] in {'main_formula', 'source_example'} and fingerprint is not None:
            if fingerprint in identities:
                fail('Repeated complete source equation needs a reviewed duplicate_of mapping, not another target: '
                     + identities[fingerprint] + ', ' + formula_id)
            identities[fingerprint] = formula_id
    main_count = sum(f.get('equation_count',1) for f in formulas.values()
                     if f['category']=='main_formula' and not f.get('duplicate_of'))
    if main_count != policy['source_main_count']:
        fail('source_main_count differs from explicit main_formula classification')
    if not formulas:
        text_value(policy.get('no_formula_reason'), 'formula_policy.no_formula_reason')
    math_pr = []
    for source in sources.values():
        if source['math']:
            settings = source['parts'].get('word/settings.xml')
            node = xml(settings, 'word/settings.xml').find('m:mathPr', NS) if settings else None
            if node is not None and any(re.search(r'italic|oblique|slanted', font.get(qn('m:val'), ''), re.I)
                                        for font in node.findall('m:mathFont', NS)):
                fail('Source mathPr has a dedicated italic font; manual source-preserving route required')
            math_pr.append(node)
    fingerprints = {tree_signature(node) if node is not None else None for node in math_pr}
    if len(fingerprints) > 1:
        fail('Source mathPr settings conflict; cannot silently select a mathematical context')
    figures = {}
    for row in spec['figures']:
        fields(row, ('id', 'source_id', 'sha256', 'number', 'title', 'width_cm'), ('media_part',), 'figure')
        identifier(row['id'], 'figure.id')
        if row['id'] in figures or row['source_id'] not in sources:
            fail('Duplicate figure or unknown source')
        integer(row['number'], 'figure.number', 1)
        visible_value(row['title'], 'figure.title')
        sha_value(row['sha256'], 'figure.sha256')
        source = sources[row['source_id']]
        if source['parts'] is not None:
            part = row.get('media_part')
            if not isinstance(part, str) or not part.startswith('word/media/') or part not in source['parts']:
                fail('DOCX figure must identify an existing word/media/ member')
            data = source['parts'][part]
        else:
            if 'media_part' in row:
                fail('Standalone image cannot have media_part')
            data = source['data']
        if digest(data) != row['sha256']:
            fail('Figure media hash mismatch: ' + row['id'])
        try:
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
                image_format = image.format
                image.verify()
        except Exception as exc:
            raise AuthoringError('Invalid figure media: ' + row['id']) from exc
        if image_format not in {'PNG', 'JPEG'}:
            fail('Only PNG/JPEG figures are supported; no silent vector conversion')
        value = row['width_cm']
        if type(value) not in (int, float) or not math.isfinite(value) or not 0.1 <= value <= 16.99:
            fail('Figure width must be 0.1..16.99 cm')
        if value * height / width > 23:
            fail('Figure is too tall for image plus caption on one page; specify a smaller width')
        figures[row['id']] = dict(row, data=data, height_cm=value * height / width)
    if [f['number'] for f in figures.values()] != list(range(1, len(figures) + 1)):
        fail('Figure registry must follow contiguous final order 1..N')
    if [s.get('kind') for s in spec['sections'] if isinstance(s, dict)] != list(SECTIONS):
        fail('Exactly five sections in the documented order are required')
    blocks = []
    ids = set()
    formula_uses, figure_uses = Counter(), Counter()
    abstract_uses, title_count = [], 0
    block_texts = {}
    benefit_count = 0
    benefit_region = False
    invention_region = False
    for section in spec['sections']:
        fields(section, ('kind', 'blocks'), where='section')
        if not isinstance(section['blocks'], list):
            fail('section.blocks must be an array')
        kind = section['kind']
        for block in section['blocks']:
            if not isinstance(block, dict):
                fail('Each block must be an object')
            block_id = identifier(block.get('id'), 'block.id')
            if block_id in ids:
                fail('Duplicate block id: ' + block_id)
            ids.add(block_id)
            btype = block.get('type')
            string = ''
            if btype == 'paragraph':
                fields(block, ('id', 'type', 'role'), ('text', 'runs'), 'paragraph')
                if kind in {'abstract_figure', 'figures'}:
                    fail('Figure sections accept figure blocks only')
                if block['role'] not in {'body', 'heading', 'project_title', 'benefit'}:
                    fail('Unsupported paragraph role')
                if ('text' in block) == ('runs' in block):
                    fail('Paragraph requires exactly one of text/runs')
                explanation_intro(block)
                runs = [{'type': 'text', 'text': block['text']}] if 'text' in block else block['runs']
                if not isinstance(runs, list) or not runs:
                    fail('runs must be a nonempty array')
                for run in runs:
                    if not isinstance(run, dict):
                        fail('Run must be an object')
                    rtype = run.get('type')
                    if rtype in {'source_symbol','source_operator'}:
                        if rtype=='source_symbol':
                            fields(run,('type','formula_id'),('node_path','node_paths','representation'),where='source_symbol')
                            if ('node_path' in run)==('node_paths' in run):
                                fail('source_symbol requires exactly one of node_path/node_paths')
                            source_path=run.get('node_path',run.get('node_paths'))
                            transform='complete_symbol_copy'
                            representation=run.get('representation','auto')
                            if representation not in {'auto','math'}:fail('Unknown source_symbol representation')
                        else:
                            fields(run,('type','formula_id','name_path','argument_path'),where='source_operator')
                            source_path=[run['name_path'],run['argument_path']]
                            transform='operator_argument_placeholder'
                            representation='math'
                        parent=formulas.get(run['formula_id'])
                        if parent is None or parent.get('source_fragment_path'):
                            fail('Source symbol must refer to an original registered formula')
                        try:
                            node=(math_fragments.symbol_fragment(parent['node'],source_path) if rtype=='source_symbol'
                                  else math_fragments.operator_reference(parent['node'],*source_path))
                        except ValueError as exc:fail(str(exc))
                        fid='symbol_'+digest(json.dumps([run['formula_id'],source_path,transform,representation],ensure_ascii=False).encode())[:32]
                        if fid in formulas and not formulas[fid].get('source_fragment_path'):
                            fail('Derived symbol id collides with a registered formula')
                        formulas[fid]=symbol_display(dict(parent,id=fid,category='symbol_reference',node=node,
                            tree_sha256=tree_signature(node),plain_text=plain_source_text(node,'symbol_reference'),
                            source_formula_id=parent['id'],source_tree_sha256=parent['tree_sha256'],
                            source_fragment_path=source_path,source_transform=transform))
                        run.clear();run.update(type='formula',formula_id=fid)
                        rtype='formula'
                    if rtype == 'symbol_reference' and 'formula_id' in run:
                        fields(run, ('type', 'text', 'formula_id'), ('style',), 'bound symbol_reference')
                        original=formulas.get(run['formula_id'])
                        if (original is None or original['category']!='symbol_reference' or
                                plain_source_text(original['node'],'symbol_reference',allow_greek=True)!=run['text']):
                            fail('Bound plain symbol_reference must match the complete registered bare symbol; never replace a structured quantity with a letter')
                        if run.get('style','body') not in {'body','reference_letter'}:
                            fail('Unknown plain symbol style')
                        fid=run['formula_id']
                        run.clear();run.update(type='formula',formula_id=fid)
                        rtype='formula'
                    if rtype in {'text', 'symbol_reference'}:
                        fields(run, ('type', 'text'), ('formula_id', 'style') if rtype == 'symbol_reference' else (), 'run')
                        string += visible_value(run['text'], 'run.text')
                        if rtype == 'symbol_reference':
                            if not ordinary_reference_letter(run['text']):
                                fail('Plain symbol_reference is one ordinary Latin/Greek letter; use source OMML for structured/special glyphs')
                            if run.get('style', 'body') not in {'body', 'reference_letter'}:
                                fail('Unknown plain symbol style')
                    elif rtype == 'formula':
                        fields(run, ('type', 'formula_id'), where='formula run')
                        if run['formula_id'] not in formulas:
                            fail('Unknown inline formula')
                        string += formulas[run['formula_id']]['plain_text'] or ''
                        formula_uses[run['formula_id']] += 1
                    elif rtype == 'figure_ref':
                        fields(run, ('type', 'figure_id'), where='figure_ref')
                        if run['figure_id'] not in figures:
                            fail('Unknown figure reference')
                        string += '图' + str(figures[run['figure_id']]['number'])
                    else:
                        fail('Unsupported run type: ' + str(rtype))
                if block['role'] == 'benefit':
                    if kind != 'description' or not benefit_region:
                        fail('Benefit paragraphs must follow the invention benefit introduction')
                    if re.match(r'^\s*(?:[0-9]+[.．、]|[（(]|[一二三四五六七八九十]+[、.])', string):
                        fail('Benefit role accepts unnumbered prose; do not supply a second list prefix')
                    text_value(string, 'benefit prose')
                    if string[0].isspace():
                        fail('Benefit prose must begin immediately after its generated ordinal')
                    benefit_count += 1
                    prefix = str(benefit_count)+'.'
                    if 'text' in block:
                        block['text'] = prefix+block['text']
                    else:
                        block['runs'] = [{'type':'text', 'text':prefix}]+block['runs']
                    string = prefix+string
                    block['role'] = 'body'
                elif kind == 'description':
                    if block['role'] == 'heading':
                        invention_region = string.strip() == '发明内容'
                        benefit_region = False
                    if inspect_docx.benefit_introduction(string):
                        if not invention_region or benefit_region or benefit_count:
                            fail('Benefit introduction must occur once within invention content')
                        benefit_region = True
                    elif benefit_region and block['role'] == 'body':
                        ordinal = re.fullmatch(r'([1-9][0-9]*)\.([^\s].*)', string, re.S)
                        if not ordinal or int(ordinal[1]) != benefit_count+1:
                            fail('Use benefit role for unnumbered effects or supply consecutive literal ordinals')
                        benefit_count += 1
                if block['role'] == 'project_title':
                    title_count += 1
                    if kind != 'description' or string != spec['project_title'] or any(r['type'] != 'text' for r in runs):
                        fail('Project title must exactly equal project_title as text in description')
                if block['role'] == 'heading' and kind != 'description':
                    fail('Column headings belong to description')
            elif btype == 'formula':
                fields(block, ('id', 'type', 'formula_id'), ('display_breaks','single_line'), where='formula block')
                if kind in {'abstract_figure', 'figures'} or block['formula_id'] not in formulas:
                    fail('Unknown formula or formula in figure section')
                if formulas[block['formula_id']]['category'] == 'symbol_reference':
                    fail('Symbol references belong in paragraph runs, not complete display formula blocks')
                if formulas[block['formula_id']]['plain_text'] is not None:
                    fail('Plain symbols and prose examples belong in paragraph runs, not display formula blocks')
                formula=formulas[block['formula_id']]
                display_node=formula['node']
                if 'single_line' in block:
                    if 'display_breaks' in block:fail('single_line and display_breaks are mutually exclusive')
                    try:display_node=single_line_math.render(display_node,block['single_line'],formula.get('equation_count',1))
                    except ValueError as exc:fail(str(exc))
                if 'display_breaks' in block:
                    try:
                        display_node=math_fragments.stacked_equations(formula['node'],block['display_breaks'],formula.get('equation_count',1))
                    except ValueError as exc:fail(str(exc))
                try:
                    display_node=math_fragments.equation_group_punctuation(display_node,formula.get('equation_count',1))
                except ValueError as exc:fail(str(exc))
                string=formula_suffix({'node':display_node})
                display_children=[n for n in display_node if n.tag!=qn('m:oMathPr')]
                if len(display_children)==1 and display_children[0].tag==qn('m:eqArr') and string:
                    # A Word run after an array aligns to its vertical centre.
                    # Place display punctuation at the end of the final row.
                    display_node=deepcopy(display_node)
                    rows=display_node.findall('m:eqArr/m:e',NS)
                    if not rows:fail('Display equation array has no final row')
                    mark=ET.SubElement(rows[-1],qn('m:r'))
                    props=ET.SubElement(mark,qn('m:rPr'))
                    ET.SubElement(props,qn('m:sty')).set(qn('m:val'),'p')
                    ET.SubElement(mark,qn('m:t')).text=string
                    string=''
                if display_node is formula['node']:
                    block['display_node']=formula['display_node']
                    block['display_tree_sha256']=formula['display_tree_sha256']
                else:
                    block['display_node']=upright_math_display(display_node)
                    block['display_tree_sha256']=tree_signature(block['display_node'])
                formula_uses[block['formula_id']] += 1
            elif btype in {'figure', 'figure_description'}:
                fields(block, ('id', 'type', 'figure_id'), where=btype)
                if block['figure_id'] not in figures:
                    fail('Unknown figure')
                f = figures[block['figure_id']]
                if btype == 'figure':
                    if kind == 'figures':
                        figure_uses[block['figure_id']] += 1
                    elif kind == 'abstract_figure':
                        abstract_uses.append(block['figure_id'])
                    else:
                        fail('Images belong in the abstract figure or final figure section')
                    string = '图' + str(f['number'])
                else:
                    if kind != 'description':
                        fail('Figure description belongs in description')
                    try:string=figure_description_contract.description(f['number'],f['title'],f['number']==len(figures))
                    except ValueError as exc:fail(str(exc))
            else:
                fail('Unsupported block type: ' + str(btype))
            blocks.append(dict(block, section=kind, bookmark=f'b{len(blocks) + 1:06d}'))
            block_texts[block_id] = string
    if title_count != 1:
        fail('Exactly one protected project_title paragraph is required')
    for block_id,string in block_texts.items():
        if invention_contract.name_prefix_issue(string,spec['project_title']):
            fail('NAME-01 title_reference_prefix: use the official name unchanged in '+block_id)
    for formula_id, formula in formulas.items():
        count = formula_uses[formula_id]
        if formula.get('duplicate_of') and formula['category'] in {'main_formula', 'source_example'}:
            if count:
                fail('A duplicate source occurrence must not appear again in the target: ' + formula_id)
            continue
        if count < 1 or (formula['category'] in {'main_formula', 'source_example', 'prose_example', 'prose_label'} and count != 1):
            fail('Main/example formulas require one target; symbol references at least one: ' + formula_id)
    if figure_uses != Counter({f: 1 for f in figures}):
        fail('Each registered figure requires exactly one final figure block')
    final_order = [b['figure_id'] for b in blocks if b['type'] == 'figure' and b['section'] == 'figures']
    if final_order != list(figures):
        fail('Final images must retain the registered figure order')
    if abstract_uses and (len(abstract_uses) != 1 or figures[abstract_uses[0]]['number'] != 1):
        fail('Abstract figure can reuse only final figure 1 once')
    described = [b['figure_id'] for b in blocks if b['type'] == 'figure_description']
    if described != list(figures):
        fail('Each figure requires one ordered figure_description block')
    for block_id, string in block_texts.items():
        for match in re.finditer(r'图\s*([0-9]+)', string):
            if int(match.group(1)) not in {f['number'] for f in figures.values()}:
                fail('Unmapped figure number in ' + block_id)
    for row in spec['protected_names']:
        fields(row, ('text', 'block_ids'), where='protected_name')
        text_value(row['text'], 'protected_name.text')
        if not isinstance(row['block_ids'], list) or not row['block_ids']:
            fail('protected_name.block_ids must be a nonempty array')
        for block_id in row['block_ids']:
            if block_id not in block_texts or row['text'] not in block_texts[block_id]:
                fail('Protected exact name missing from declared block: ' + str(block_id))
    return {'root': root, 'spec': spec, 'sources': sources, 'formulas': formulas, 'figures': figures,
            'blocks': blocks, 'block_texts': block_texts, 'math_pr': math_pr[0] if math_pr else None}


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def font(run, name='宋体', size=14, bold=False, spacing=0):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.cs_italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    props = run._r.get_or_add_rPr()
    fonts = props.find(qn('w:rFonts'))
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), name)
    props.append(element('w:spacing', val=spacing))
    props.append(element('w:szCs', val=int(size * 2)))
    props.append(element('w:vertAlign', val='baseline'))


def paragraph_format(p, role='body', section='description'):
    props = p._p.get_or_add_pPr()
    indented = role in {'body', 'heading', 'project_title'}
    centered = role in {'project_title', 'formula', 'image', 'caption', 'header', 'footer'}
    props.append(element('w:jc', val='center' if centered else 'left'))
    props.append(element('w:ind', left=0, right=0,
                         firstLine=562 if role == 'project_title' else 560 if indented else 0,
                         firstLineChars=200 if indented else 0))
    spacing = element('w:spacing', before=0, after=0)
    if section == 'claims' and role == 'body':
        spacing.set(qn('w:line'), '240')
        spacing.set(qn('w:lineRule'), 'auto')
    props.append(spacing)
    props.append(element('w:snapToGrid', val='1' if indented else '0'))
    for name in ('autoSpaceDE', 'autoSpaceDN'):
        props.append(element('w:' + name, val='0'))
    # Ordinary prose must use the remaining page area and continue naturally.
    # Keep only an actual indivisible object or short heading together.
    props.append(element('w:keepLines', val=0 if role == 'body' else 1))
    props.append(element('w:keepNext', val=int(role in {'heading', 'project_title'} or
                                             (role == 'image' and section == 'figures'))))
    props.append(element('w:pageBreakBefore', val=0))
    props.append(element('w:widowControl', val=0))


def attach_section_boundary(document):
    """Reuse the preceding real paragraph instead of a full-height blank line.

    python-docx's add_section creates a separate empty paragraph. Keep a tiny
    structural carrier only when the preceding section itself has no content.
    Header/footer spacers and source object paragraphs are never removed here.
    """
    boundary = document.paragraphs[-1]._p
    props = boundary.get_or_add_pPr()
    section = props.find(qn('w:sectPr'))
    require(section is not None and not boundary.xpath('./w:r | ./m:oMath | ./m:oMathPara'),
            'unexpected section carrier')
    previous = boundary.getprevious()
    if (previous is not None and previous.tag == qn('w:p') and
            previous.find('w:pPr/w:sectPr', NS) is None):
        previous.get_or_add_pPr().append(section)
        boundary.getparent().remove(boundary)
    else:
        props.remove(section)
        for key in ('keepLines', 'keepNext', 'pageBreakBefore', 'widowControl', 'snapToGrid'):
            props.append(element('w:' + key, val=0))
        props.append(element('w:spacing', before=0, after=0, line=20, lineRule='exact'))
        mark_props = element('w:rPr')
        mark_props.append(element('w:sz', val=2))
        mark_props.append(element('w:szCs', val=2))
        props.append(mark_props)
        props.append(section)


def line_shape(paragraph, line_id, top_pt, left_pt=0):
    """One DrawingML anchor, no duplicate compatibility/VML branch."""
    cx = round(481.9 * 12700)
    fragment = f'''<w:r xmlns:w="{NS['w']}" xmlns:wp="{NS['wp']}" xmlns:a="{NS['a']}"
      xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
        relativeHeight="0" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
        <wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>{round(left_pt * 12700)}</wp:posOffset></wp:positionH>
        <wp:positionV relativeFrom="paragraph"><wp:posOffset>{round(top_pt * 12700)}</wp:posOffset></wp:positionV>
        <wp:extent cx="{cx}" cy="0"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
        <wp:docPr id="{line_id}" name="Rule {line_id}"/><wp:cNvGraphicFramePr/>
        <a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="0"/></a:xfrm>
        <a:prstGeom prst="line"><a:avLst/></a:prstGeom><a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>
        </wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>
      </wp:anchor></w:drawing></w:r>'''
    paragraph._p.append(ET.fromstring(fragment))


def configure_section(section, index):
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.top_margin, section.right_margin = Twips(1417), Twips(851)
    section.bottom_margin, section.left_margin = Twips(850), Twips(1418)
    section.gutter = Twips(0)
    section.header_distance = Twips(851 if index < 3 else 709)
    section.footer_distance = Twips(227 if index < 3 else 70)
    section.different_first_page_header_footer = False
    props = section._sectPr
    for tag in ('w:docGrid', 'w:pgNumType'):
        for child in list(props.findall(qn(tag))):
            props.remove(child)
    props.append(element('w:docGrid', type='lines', linePitch=312))
    if index >= 3:
        props.append(element('w:pgNumType', start=1))
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    paragraph_format(header, 'header')
    font(header.add_run(HEADERS[index]), name='宋体' if index == 2 else '黑体',
         bold=index == 2, spacing=0 if index == 2 else 90)
    if index == 2:
        borders = element('w:pBdr')
        borders.append(element('w:bottom', val='single', sz=6, space=1, color='000000'))
        header._p.get_or_add_pPr().append(borders)
    else:
        line_shape(header, 1001 + index, 26.3)
    footer = section.footer.paragraphs[0]
    footer.clear()
    paragraph_format(footer, 'footer')
    if index != 1:
        # The manual footer is a stack of paragraphs.  Its height is part of
        # the anchor geometry; collapsing it to one line puts the rule at the
        # paper edge even when pgMar/footer and the anchor offset are correct.
        carrier_spacing = footer._p.get_or_add_pPr().find(qn('w:spacing'))
        carrier_spacing.set(qn('w:line'), '200')
        carrier_spacing.set(qn('w:lineRule'), 'exact')
        if index < 3:
            carrier_rpr = element('w:rPr')
            carrier_rpr.append(element('w:rFonts', ascii='黑体', hAnsi='黑体', eastAsia='黑体', cs='黑体'))
            carrier_rpr.append(element('w:sz', val=18))
            carrier_rpr.append(element('w:szCs', val=18))
            footer._p.get_or_add_pPr().append(carrier_rpr)
    if index >= 3:
        begin = element('w:fldChar', fldCharType='begin', dirty='true')
        code = element('w:instrText')
        code.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        code.text = ' PAGE '
        # A proper begin/code/separate/result/end field. Empty cached text is
        # intentional; the renderer evaluates PAGE. Each carrier is 9 pt,
        # including the result and paragraph mark (not only the instruction).
        for child in (begin, code, element('w:fldChar', fldCharType='separate'),
                      element('w:t'), element('w:fldChar', fldCharType='end')):
            run = footer.add_run()
            font(run, name='黑体', size=9)
            run._r.append(child)
        footer._p.get_or_add_pPr().append(deepcopy(run._r.get_or_add_rPr()))
    if index != 1:
        line_shape(footer, 2001 + index, 9.6 if index >= 3 else -3, 1)
        font(footer.runs[-1], name='黑体', size=10, spacing=90)
        if index >= 3:
            spacer = section.footer.add_paragraph(style='Footer')
            paragraph_format(spacer, 'footer')
            spacing = spacer._p.get_or_add_pPr().find(qn('w:spacing'))
            spacing.set(qn('w:line'), '200')
            spacing.set(qn('w:lineRule'), 'exact')
        tail = section.footer.add_paragraph(style='Footer')
        paragraph_format(tail, 'footer')
        # Explicitly resolve the reference's 9 pt Footer -> Normal font chain,
        # rather than inheriting this generator's 14 pt body or Hei metrics.
        tail_props = tail._p.get_or_add_pPr()
        if index < 3:
            # Resolve the reference's empty-line layout explicitly.  A generated
            # document's different Normal/style chain otherwise adds ~2 pt to
            # this blank paragraph, even when the footer XML is copied intact.
            tail_spacing = tail_props.find(qn('w:spacing'))
            tail_spacing.set(qn('w:line'), '207')
            tail_spacing.set(qn('w:lineRule'), 'exact')
        tail_rpr = element('w:rPr')
        tail_rpr.append(element('w:rFonts', ascii='Times New Roman', hAnsi='Times New Roman',
                                eastAsia='宋体', cs='Times New Roman'))
        tail_rpr.append(element('w:sz', val=18))
        tail_rpr.append(element('w:szCs', val=18))
        tail_props.append(tail_rpr)


def mark(p, bookmark, marker_id):
    start = element('w:bookmarkStart', id=marker_id, name=bookmark)
    end = element('w:bookmarkEnd', id=marker_id)
    p._p.append(start)
    p._p.append(end)


def render_document(context, path):
    document = Document()
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = '宋体', Pt(14)
    normal.font.bold, normal.font.italic = False, False
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    normal.element.get_or_add_rPr().find(qn('w:rFonts')).set(qn('w:eastAsia'), '宋体')
    # Footer paragraph marks must not inherit 14 pt body metrics: that can move
    # the baseline down across the fixed +9.6 pt rule in LibreOffice/Word.
    footer_style = document.styles['Footer']
    footer_style.font.name, footer_style.font.size = '黑体', Pt(9)
    footer_style.paragraph_format.line_spacing = 1
    settings = document.settings.element
    for node in list(settings.findall(qn('w:evenAndOddHeaders'))):
        settings.remove(node)
    for node in list(settings.findall(qn('m:mathPr'))):
        settings.remove(node)
    if context['math_pr'] is not None:
        settings.append(deepcopy(context['math_pr']))
    settings.append(element('w:updateFields', val='true'))
    document.core_properties.title = context['spec']['project_title']
    document.core_properties.subject = 'Structured patent draft; semantic and visual review pending'
    document.core_properties.author = ''
    for index, kind in enumerate(SECTIONS):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        if index:
            attach_section_boundary(document)
        configure_section(section, index)
        for block in (b for b in context['blocks'] if b['section'] == kind):
            p = document.add_paragraph()
            mark(p, block['bookmark'], len(document.paragraphs) * 2)
            btype = block['type']
            if btype == 'paragraph':
                paragraph_format(p, block['role'], kind)
                runs = [{'type': 'text', 'text': block['text']}] if 'text' in block else block['runs']
                for run in runs:
                    if run['type'] == 'formula':
                        formula = context['formulas'][run['formula_id']]
                        if formula['plain_text'] is None:
                            p._p.append(deepcopy(formula.get('display_node',formula['node'])))
                        else:
                            font(p.add_run(formula['plain_text']), bold=block['role'] in {'heading', 'project_title'})
                    else:
                        string = '图' + str(context['figures'][run['figure_id']]['number']) if run['type'] == 'figure_ref' else run['text']
                        added = p.add_run(string)
                        font(added, bold=block['role'] in {'heading', 'project_title'})
            elif btype == 'formula':
                paragraph_format(p, 'formula', kind)
                if 'single_line' in block:
                    math_paragraph_properties=element('w:rPr')
                    for prop in ('w:sz','w:szCs'):math_paragraph_properties.append(element(prop,val=int(block['single_line']['font_size_pt']*2)))
                    p._p.get_or_add_pPr().append(math_paragraph_properties)
                # The correct reference centers an inline math object in its
                # own paragraph, keeping the following prose semicolon inline.
                p._p.append(deepcopy(block.get('display_node',context['formulas'][block['formula_id']]['node'])))
                suffix = context['block_texts'][block['id']]
                if suffix:
                    font(p.add_run(suffix),size=block.get('single_line',{}).get('font_size_pt',14))
            elif btype == 'figure_description':
                paragraph_format(p, 'body', kind)
                font(p.add_run(context['block_texts'][block['id']]))
            else:
                figure = context['figures'][block['figure_id']]
                paragraph_format(p, 'image', kind)
                p.add_run().add_picture(io.BytesIO(figure['data']), width=Cm(figure['width_cm']))
                if kind == 'figures':
                    caption = document.add_paragraph()
                    mark(caption, block['bookmark'] + 'c', len(document.paragraphs) * 2)
                    paragraph_format(caption, 'caption', kind)
                    font(caption.add_run('图' + str(figure['number'])))
    document.save(path)


def word_text(node):
    return ''.join(node.xpath('.//w:t[not(ancestor::m:oMath)]/text()', namespaces=NS))


def attributes(node):
    return {ET.QName(k).localname: v for k, v in node.attrib.items()} if node is not None else {}


def relationships(parts, part='word/document.xml'):
    prefix, name = part.rsplit('/', 1)
    relpart = prefix + '/_rels/' + name + '.rels'
    result = {}
    for rel in xml(parts[relpart], relpart):
        if rel.get('TargetMode') == 'External':
            fail('Unexpected external relationship in generated DOCX')
        target = rel.get('Target', '')
        # Generated relationships may use /word/... absolute package paths.
        result[rel.get('Id')] = target.lstrip('/') if target.startswith('/') else prefix + '/' + target
    return result


def require(condition, message):
    if not condition:
        fail('Readback verification: ' + message)


def verify_document(context, path):
    """Independently open final ZIP and reconstruct positions, counts and media.

    It consumes the input contract, never the builder's paragraph bookkeeping.
    It is intentionally not a mathematical/wording/visual acceptance function.
    """
    parts = package(Path(path).read_bytes(), str(path))
    doc = xml(parts['word/document.xml'], 'word/document.xml')
    body = doc.find('w:body', NS)
    require(body is not None, 'missing body')
    require(all(node.tag == qn('w:p') for node in list(body)[:-1]) and
            len(body) > 0 and body[-1].tag == qn('w:sectPr'), 'unexpected body object or section boundary')
    rels = relationships(parts)
    settings = xml(parts['word/settings.xml'], 'word/settings.xml')
    actual_math_pr = settings.find('m:mathPr', NS)
    wanted_math_pr = context['math_pr']
    require((tree_signature(actual_math_pr) if actual_math_pr is not None else None) ==
            (tree_signature(wanted_math_pr) if wanted_math_pr is not None else None), 'source mathPr context changed')
    require(settings.find('w:evenAndOddHeaders', NS) is None, 'unexpected even-page header mode')
    styles = xml(parts['word/styles.xml'], 'word/styles.xml')
    footer_style = styles.find('w:style[@w:styleId="Footer"]', NS)
    require(footer_style is not None and
            attributes(footer_style.find('w:rPr/w:sz', NS)).get('val') == '18' and
            attributes(footer_style.find('w:rPr/w:rFonts', NS)).get('ascii') == '黑体' and
            attributes(footer_style.find('w:pPr/w:spacing', NS)).get('line') == '240' and
            attributes(footer_style.find('w:pPr/w:spacing', NS)).get('lineRule') == 'auto',
            'footer inherited style metrics changed')
    sections = doc.xpath('.//w:sectPr', namespaces=NS)
    require(len(sections) == 5, 'expected five sections')
    section_evidence = []
    for i, section in enumerate(sections):
        require(not section.findall('w:titlePg', NS), 'unexpected first-page header mode')
        headers = section.findall('w:headerReference', NS)
        footers = section.findall('w:footerReference', NS)
        require(len(headers) == len(footers) == 1, 'each section must have explicit default header/footer')
        require(headers[0].get(qn('w:type')) == footers[0].get(qn('w:type')) == 'default', 'unexpected header/footer type')
        header_part = rels[headers[0].get(qn('r:id'))]
        footer_part = rels[footers[0].get(qn('r:id'))]
        header = xml(parts[header_part], header_part)
        footer = xml(parts[footer_part], footer_part)
        require(word_text(header) == HEADERS[i], 'section header text changed')
        expected_margins = {'top': '1417', 'right': '851', 'bottom': '850', 'left': '1418',
                            'header': '851' if i < 3 else '709', 'footer': '227' if i < 3 else '70', 'gutter': '0'}
        require(attributes(section.find('w:pgMar', NS)) == expected_margins, 'section margins changed')
        require(attributes(section.find('w:pgSz', NS)) == {'w': '11906', 'h': '16838'}, 'paper size changed')
        require(attributes(section.find('w:docGrid', NS)) == {'type': 'lines', 'linePitch': '312'}, 'line grid changed')
        page_fields = footer.xpath('.//w:instrText/text()', namespaces=NS)
        require(page_fields == ([' PAGE '] if i >= 3 else []), 'PAGE field visibility/scope changed')
        require(not word_text(footer), 'footer contains cached or visible page-number text')
        if i >= 3:
            require(footer.xpath('.//w:fldChar/@w:fldCharType', namespaces=NS) == ['begin', 'separate', 'end'],
                    'PAGE field structure changed')
            require(attributes(footer.find('.//w:pPr/w:rPr/w:sz', NS)).get('val') == '18', 'footer paragraph mark size changed')
            require(attributes(footer.find('.//w:pPr/w:spacing', NS)) == {'before': '0', 'after': '0', 'line': '200', 'lineRule': 'exact'},
                    'PAGE-only carrier line metrics changed')
            carriers = footer.xpath('.//w:r[w:fldChar or w:instrText or w:t]', namespaces=NS)
            require(len(carriers) == 5, 'PAGE carrier count changed')
            for run in carriers:
                props = run.find('w:rPr', NS)
                require(props is not None and attributes(props.find('w:sz', NS)).get('val') == '18', 'PAGE carrier size changed')
                require(attributes(props.find('w:rFonts', NS)).get('eastAsia') == '黑体' and
                        attributes(props.find('w:rFonts', NS)).get('ascii') == '黑体', 'PAGE carrier font changed')
                require(attributes(props.find('w:color', NS)).get('val') == '000000', 'PAGE carrier color changed')
        require(attributes(section.find('w:pgNumType', NS)) == ({'start': '1'} if i >= 3 else {}), 'page restart changed')
        require(len(header.xpath('.//wp:anchor', namespaces=NS)) == (0 if i == 2 else 1), 'header line count changed')
        require(len(footer.xpath('.//wp:anchor', namespaces=NS)) == (0 if i == 1 else 1), 'footer line count changed')
        footer_paragraphs = footer.findall('w:p', NS)
        require(len(footer_paragraphs) == (1 if i == 1 else 2 if i < 3 else 3),
                'footer paragraph stack changed')
        if i != 1:
            line_run = footer_paragraphs[0].find('w:r[w:drawing]', NS)
            require(line_run is not None and
                    attributes(line_run.find('w:rPr/w:sz', NS)).get('val') == '20' and
                    attributes(line_run.find('w:rPr/w:rFonts', NS)).get('eastAsia') == '黑体',
                    'footer line run metrics changed')
            require(attributes(footer_paragraphs[0].find('w:pPr/w:spacing', NS)) ==
                    {'before': '0', 'after': '0', 'line': '200', 'lineRule': 'exact'},
                    'footer line carrier metrics changed')
            require(attributes(footer_paragraphs[0].find('w:pPr/w:rPr/w:sz', NS)).get('val') == '18',
                    'footer line paragraph mark size changed')
            for spacer in footer_paragraphs[1:]:
                require(len(spacer) == 1 and spacer[0].tag == qn('w:pPr'), 'footer spacer is not empty')
                require(attributes(spacer.find('w:pPr/w:snapToGrid', NS)).get('val') == '0',
                        'footer spacer grid changed')
            if i >= 3:
                require(attributes(footer_paragraphs[1].find('w:pPr/w:spacing', NS)) ==
                        {'before': '0', 'after': '0', 'line': '200', 'lineRule': 'exact'},
                        'footer spacer line metrics changed')
            tail = footer_paragraphs[-1]
            require(attributes(tail.find('w:pPr/w:spacing', NS)) ==
                    ({'before': '0', 'after': '0', 'line': '207', 'lineRule': 'exact'} if i < 3 else {'before': '0', 'after': '0'}),
                    'footer tail line metrics changed')
            require(attributes(tail.find('w:pPr/w:rPr/w:rFonts', NS)) ==
                    {'ascii': 'Times New Roman', 'hAnsi': 'Times New Roman', 'eastAsia': '宋体', 'cs': 'Times New Roman'} and
                    attributes(tail.find('w:pPr/w:rPr/w:sz', NS)).get('val') == '18',
                    'footer tail font metrics changed')
        hr = header.find('.//w:r[w:t]', NS)
        require(hr is not None, 'missing header text run')
        props = hr.find('w:rPr', NS)
        require(attributes(props.find('w:rFonts', NS)).get('eastAsia') == ('宋体' if i == 2 else '黑体'), 'header font changed')
        require(attributes(props.find('w:spacing', NS)).get('val') == ('0' if i == 2 else '90'), 'header spacing changed')
        require(attributes(props.find('w:sz', NS)).get('val') == '28', 'header size changed')
        require(attributes(props.find('w:b', NS)).get('val', '1') == ('1' if i == 2 else '0'), 'header bold changed')
        require(attributes(props.find('w:color', NS)).get('val') == '000000', 'header color changed')
        border = header.find('.//w:pPr/w:pBdr/w:bottom', NS)
        require(attributes(border) == ({'val': 'single', 'sz': '6', 'space': '1', 'color': '000000'} if i == 2 else {}),
                'claim header border changed')
        for container, y, x in ((header, 26.3, 0), (footer, 9.6 if i >= 3 else -3, 1)):
            for anchor in container.xpath('.//wp:anchor', namespaces=NS):
                require(anchor.find('wp:positionH', NS).get('relativeFrom') == 'column' and
                        anchor.find('wp:positionV', NS).get('relativeFrom') == 'paragraph', 'line reference frame changed')
                require(anchor.find('wp:positionH/wp:posOffset', NS).text == str(round(x * 12700)) and
                        anchor.find('wp:positionV/wp:posOffset', NS).text == str(round(y * 12700)), 'line offset changed')
                require(dict(anchor.find('wp:extent', NS).attrib) == {'cx': str(round(481.9 * 12700)), 'cy': '0'}, 'line width changed')
                require(anchor.find('.//a:ln', NS).get('w') == '12700' and
                        anchor.find('.//a:ln/a:solidFill/a:srgbClr', NS).get('val') == '000000', 'line thickness/color changed')
        section_evidence.append({'kind': SECTIONS[i], 'header_part': header_part, 'footer_part': footer_part,
                                 'header_text': word_text(header), 'page_field': i >= 3})
    require(len({s['header_part'] for s in section_evidence}) == 5, 'headers unexpectedly shared')
    require(len({s['footer_part'] for s in section_evidence}) == 5, 'footers unexpectedly shared')
    paragraphs = body.findall('w:p', NS)
    observed = {}
    current_section = 0
    for index, p in enumerate(paragraphs, 1):
        names = p.xpath('./w:bookmarkStart/@w:name', namespaces=NS)
        for name in names:
            require(name not in observed, 'duplicate block bookmark')
            observed[name] = (index, p, SECTIONS[current_section])
        if p.find('w:pPr/w:sectPr', NS) is not None:
            current_section += 1
    expected_markers = {b['bookmark'] for b in context['blocks']}
    expected_markers |= {b['bookmark'] + 'c' for b in context['blocks']
                         if b['type'] == 'figure' and b['section'] == 'figures'}
    require(set(observed) == expected_markers, 'missing/unexpected block bookmarks')
    expected_order = []
    for block in context['blocks']:
        expected_order.append(block['bookmark'])
        if block['type'] == 'figure' and block['section'] == 'figures':
            expected_order.append(block['bookmark'] + 'c')
    require(list(observed) == expected_order, 'block/caption order changed')
    expected_math_count = 0
    object_rows, block_rows = [], []
    scaled_math_suffix_sizes = {}
    for block in context['blocks']:
        index, p, section = observed[block['bookmark']]
        require(section == block['section'], 'block moved to wrong section: ' + block['id'])
        locator = f'word/document.xml body/w:p[{index}]'
        row = {'id': block['id'], 'type': block['type'], 'section': section, 'bookmark': block['bookmark'], 'target_locator': locator}
        role = block.get('role', {'formula': 'formula', 'figure': 'image', 'figure_description': 'body'}.get(block['type']))
        if block['type']=='formula' and 'single_line' in block:
            half_points=str(int(block['single_line']['font_size_pt']*2))
            for prop in ('sz','szCs'):
                require(attributes(p.find('w:pPr/w:rPr/w:'+prop,NS)).get('val')==half_points,'Single-line paragraph math size changed')
            for run in p.findall('w:r',NS):
                scaled_math_suffix_sizes[run]=half_points
                require(attributes(run.find('w:rPr/w:szCs',NS)).get('val')==half_points,'Single-line suffix math size changed')
        indented = role in {'body', 'heading', 'project_title'}
        wanted_ind = {'left': '0', 'right': '0', 'firstLine': '562' if role == 'project_title' else '560' if indented else '0',
                      'firstLineChars': '200' if indented else '0'}
        require(attributes(p.find('w:pPr/w:ind', NS)) == wanted_ind, 'paragraph indentation changed: ' + block['id'])
        wanted_alignment = 'center' if role in {'project_title', 'formula', 'image'} else 'left'
        require(attributes(p.find('w:pPr/w:jc', NS)).get('val') == wanted_alignment, 'paragraph alignment changed')
        for prop in ('autoSpaceDE','autoSpaceDN'):
            require(attributes(p.find('w:pPr/w:'+prop,NS)).get('val')=='0','Automatic extra content spacing changed')
        for prop, value in {'keepLines': int(role != 'body'),
                            'keepNext': int(role in {'heading', 'project_title'} or
                                            (role == 'image' and section == 'figures')),
                            'pageBreakBefore': 0, 'widowControl': 0}.items():
            require(attributes(p.find('w:pPr/w:' + prop, NS)).get('val') == str(value),
                    'paragraph pagination changed: ' + block['id'] + '/' + prop)
        if section == 'claims' and role == 'body':
            require(attributes(p.find('w:pPr/w:spacing', NS)) == {'before': '0', 'after': '0', 'line': '240', 'lineRule': 'auto'},
                    'claim line spacing changed')
        if block['type'] == 'paragraph':
            for run in p.findall('w:r', NS):
                if run.find('w:t', NS) is not None:
                    require(attributes(run.find('w:rPr/w:b', NS)).get('val', '1') == ('1' if role in {'heading', 'project_title'} else '0'),
                            'paragraph bold changed')
        if block['type'] == 'figure':
            figure = context['figures'][block['figure_id']]
            require(not word_text(p).strip(), 'unexpected text inside image paragraph')
            blips = p.xpath('.//a:blip', namespaces=NS)
            require(len(blips) == 1, 'figure needs one embedded image')
            media_part = rels[blips[0].get(qn('r:embed'))]
            require(digest(parts[media_part]) == figure['sha256'], 'embedded image differs from source')
            if section == 'figures':
                caption_index, caption, caption_section = observed[block['bookmark'] + 'c']
                require(caption_index == index + 1 and caption_section == section, 'caption detached from image')
                require(word_text(caption) == '图' + str(figure['number']), 'figure number changed')
                forbidden_caption_tags = {qn('w:'+tag) for tag in
                    ('br','cr','tab','fldChar','fldSimple','instrText','drawing','pict','object','sym','numPr')}
                require(not any(node.tag in forbidden_caption_tags or node.tag in {qn('m:oMath'),qn('m:oMathPara')}
                                for node in caption.iter()), 'figure caption must be one plain number-only line')
                for prop, value in {'keepNext':'0', 'keepLines':'1', 'pageBreakBefore':'0', 'widowControl':'0'}.items():
                    require(attributes(caption.find('w:pPr/w:' + prop, NS)).get('val') == value,
                            'figure caption pagination changed')
                row.update(caption_locator=f'word/document.xml body/w:p[{caption_index}]', caption_text=word_text(caption),
                           same_page='requires_rendered_page_review')
            else:
                row.update(caption_locator=None, caption_text=None, same_page='abstract_image_has_no_caption')
            object_rows.append({'kind': 'figure', 'id': figure['id'], 'source_file': context['sources'][figure['source_id']]['file'],
                                'source_sha256': context['sources'][figure['source_id']]['sha256'], 'source_media': figure.get('media_part'),
                                'media_sha256': figure['sha256'], 'target_part': media_part, 'target_locator': locator,
                                'caption_locator': row['caption_locator'], 'abstract_reuse': section == 'abstract_figure'})
            expected_formulas = []
        else:
            require(word_text(p) == context['block_texts'][block['id']], 'text changed or missing in ' + block['id'])
            expected_formulas = ([block['formula_id']] if block['type'] == 'formula' else
                                 [r['formula_id'] for r in block.get('runs', []) if r['type'] == 'formula' and
                                  context['formulas'][r['formula_id']]['plain_text'] is None])
            # Preserve interleaving: concatenated text plus a separate math list
            # would miss moving an unchanged equation across explanatory text.
            actual_tokens = []
            for child in p:
                if child.tag in {qn('w:pPr'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')}: continue
                if child.tag == qn('w:r'):
                    require(all(n.tag in {qn('w:rPr'), qn('w:t')} for n in child), 'unexpected body run content')
                    actual_tokens.append(('text', word_text(child)))
                elif child.tag in {qn('m:oMath'), qn('m:oMathPara')}:
                    node = child if child.tag == qn('m:oMath') else child.find('m:oMath', NS)
                    require(node is not None, 'empty mathematical wrapper')
                    actual_tokens.append(('math', tree_signature(node)))
                else:
                    fail('Readback verification: unexpected paragraph object')
            if block['type'] == 'formula':
                expected_tokens = [('math', block.get('display_tree_sha256',context['formulas'][block['formula_id']]['tree_sha256']))]
                if context['block_texts'][block['id']]:
                    expected_tokens.append(('text', context['block_texts'][block['id']]))
            elif 'runs' in block:
                expected_tokens = []
                actual_runs = iter(p.findall('w:r', NS))
                for r in block['runs']:
                    if r['type'] == 'formula':
                        f = context['formulas'][r['formula_id']]
                        if f['plain_text'] is None:
                            expected_tokens.append(('math', f.get('display_tree_sha256',f['tree_sha256'])))
                            continue
                        value = f['plain_text']
                        object_rows.append({'kind': 'source_text', 'id': f['id'], 'category': f['category'],
                                            'source_file': context['sources'][f['source_id']]['file'],
                                            'source_sha256': context['sources'][f['source_id']]['sha256'],
                                            'source_locator': f['locator'], 'tree_sha256': f['tree_sha256'],
                                            'source_formula_id':f.get('source_formula_id',f['id']),
                                            'source_tree_sha256':f.get('source_tree_sha256',f['tree_sha256']),
                                            'source_fragment_path':f.get('source_fragment_path'),
                                            'source_transform':f.get('source_transform'),
                                            'target_locator': locator, 'text': value})
                    else:
                        value = '图' + str(context['figures'][r['figure_id']]['number']) if r['type'] == 'figure_ref' else r['text']
                    expected_tokens.append(('text', value))
                    actual_run = next(actual_runs, None)
                    require(actual_run is not None, 'ordinary source/symbol run missing')
            else:
                expected_tokens = [('text', context['block_texts'][block['id']])]
        math = p.xpath('.//m:oMath', namespaces=NS)
        require(len(math) == len(expected_formulas), 'formula occurrence count mismatch in ' + block['id'])
        for number, (node, formula_id) in enumerate(zip(math, expected_formulas), 1):
            formula = context['formulas'][formula_id]
            expected_tree=(block.get('display_tree_sha256',formula['tree_sha256']) if block['type']=='formula'
                           else formula.get('display_tree_sha256',formula['tree_sha256']))
            require(tree_signature(node) == expected_tree, 'source OMML content/structure changed: ' + formula_id)
            object_rows.append({'kind': 'formula', 'id': formula_id, 'category': formula['category'],
                                'source_file': context['sources'][formula['source_id']]['file'],
                                'source_sha256': context['sources'][formula['source_id']]['sha256'],
                                'source_locator': formula['locator'], 'tree_sha256': formula['tree_sha256'],
                                'source_formula_id':formula.get('source_formula_id',formula_id),
                                'source_tree_sha256':formula.get('source_tree_sha256',formula['tree_sha256']),
                                'source_fragment_path':formula.get('source_fragment_path'),
                                'source_transform':formula.get('source_transform'),
                                'display_tree_sha256':expected_tree,'display_breaks':block.get('display_breaks'),
                                'single_line':block.get('single_line'),
                                'equation_count':formula.get('equation_count',1) if formula['category'] in {'main_formula','source_example'} else 0,
                                'target_locator': locator + f' oMath[{number}]'})
        expected_math_count += len(expected_formulas)
        if block['type'] != 'figure':
            require(actual_tokens == expected_tokens, 'text/formula interleaving changed in ' + block['id'])
        block_rows.append(row)
    require(len(doc.xpath('.//m:oMath', namespaces=NS)) == expected_math_count, 'unmapped OMML outside blocks')
    expected_images = sum(b['type'] == 'figure' for b in context['blocks'])
    require(len(doc.xpath('.//a:blip', namespaces=NS)) == expected_images, 'unmapped body images')
    require(not doc.xpath('.//w:ins | .//w:del | .//w:commentReference', namespaces=NS), 'unexpected revisions/comments')
    marked = {id(value[1]) for value in observed.values()}
    for p in paragraphs:
        require(id(p) in marked or (not word_text(p) and p.find('w:pPr/w:sectPr', NS) is not None), 'unregistered body content')
        if id(p) not in marked:
            previous = p.getprevious()
            require(previous is None or previous.find('w:pPr/w:sectPr', NS) is not None,
                    'redundant blank section carrier after content')
            require(attributes(p.find('w:pPr/w:spacing', NS)) ==
                    {'before':'0', 'after':'0', 'line':'20', 'lineRule':'exact'},
                    'empty section carrier consumes body line height')
            require(attributes(p.find('w:pPr/w:snapToGrid', NS)).get('val') == '0',
                    'empty section carrier snaps to body grid')
        for run in p.findall('w:r', NS):
            if run.find('w:t', NS) is not None:
                props = run.find('w:rPr', NS)
                require(props is not None and attributes(props.find('w:color', NS)).get('val') == '000000', 'body text is not explicitly black')
                expected_fonts = {k:'宋体' for k in ('ascii','hAnsi','eastAsia','cs')}
                require(all(attributes(props.find('w:rFonts', NS)).get(slot) == name
                            for slot,name in expected_fonts.items()), 'ordinary body text font changed')
                require(attributes(props.find('w:sz', NS)).get('val') == scaled_math_suffix_sizes.get(run,'28'),
                        'ordinary body text size changed')
                require(attributes(props.find('w:szCs', NS)).get('val') == scaled_math_suffix_sizes.get(run,'28'),
                        'ordinary body complex-script text size changed')
                require(attributes(props.find('w:i', NS)).get('val') == '0', 'ordinary text italic changed')
                require(attributes(props.find('w:iCs', NS)).get('val') == '0', 'ordinary text complex-script italic changed')
                require(attributes(props.find('w:vertAlign', NS)).get('val') == 'baseline', 'ordinary symbol vertical alignment changed')
                require(attributes(props.find('w:position', NS)).get('val', '0') == '0',
                        'ordinary symbol baseline position changed')
    return {'verification': 'structural_readback_passed', 'profile': PROFILE, 'sections': section_evidence,
            'blocks': block_rows, 'source_objects': object_rows,
            'source_formula_aliases': [
                {'id': f['id'], 'source_file': context['sources'][f['source_id']]['file'],
                 'source_sha256': context['sources'][f['source_id']]['sha256'],
                 'source_locator': f['locator'], 'source_tree_sha256': f['tree_sha256'],
                 'duplicate_of': f['duplicate_of'], 'reason': f['duplicate_reason'],
                 'target_occurrences': [r['target_locator'] for r in object_rows
                                        if r.get('id') == f['duplicate_of']]}
                for f in context['formulas'].values()
                if f.get('duplicate_of') and f['category'] in {'main_formula', 'source_example'}],
            'counts': {'sections': 5, 'blocks': len(block_rows), 'omml_occurrences': expected_math_count,
                       'source_text_occurrences': sum(r['kind'] == 'source_text' for r in object_rows),
                       'main_formula_occurrences': sum(r.get('equation_count',1) for r in object_rows if r.get('category')=='main_formula'),
                       'image_occurrences_including_abstract_reuse': expected_images},
            'semantic_review': 'not_run', 'visual_review': 'not_run',
            'boundary': 'Structural readback is not semantic approval. N-to-N uses the supplied classification, not an inferred count of mathematical relationships. Same-page controls require final rendering and visual review.'}


def build(case_root, input_file, out_dir):
    root = Path(case_root).resolve(strict=True)
    source_json = bounded_path(root, input_file)
    output = bounded_path(root, out_dir, exists=False)
    if output.exists():
        fail('Output directory already exists; select a new revision directory')
    spec = read_json(source_json)
    context = validate(root, spec)
    output.parent.mkdir(parents=True, exist_ok=True)
    # A short sibling staging path avoids deep Chinese project/cache paths.
    staging = Path(tempfile.mkdtemp(prefix='.author-', dir=output.parent))
    try:
        docx_path = staging / 'draft.docx'
        render_document(context, docx_path)
        dump(staging / 'content.json', spec)
        for source in context['sources'].values():
            if sha(source['path']) != source['sha256']:
                fail('Source changed during build: ' + source['file'])
        if read_json(source_json) != spec:
            fail('Content JSON changed during build')
        # Exclusive reservation; no old output is ever overwritten. The manifest
        # is written LAST; absence means an incomplete commit, not approval.
        output.mkdir(exist_ok=False)
        for name in ('draft.docx', 'content.json'):
            destination = output / name
            with (staging / name).open('rb') as src, destination.open('xb') as dest:
                shutil.copyfileobj(src, dest)
        final_docx = output / 'draft.docx'
        readback = verify_document(context, final_docx)
        # One existing extractor call at the permanent artifact path. The next
        # handoff searches artifact.parent/inspection-cache and reuses this.
        summary = inspect_docx.run_inspection(final_docx, output / 'inspection-cache', 'draft', True, True)
        inspection_file = Path(summary['inspection'])
        cache_file = Path(summary['output_directory']) / 'cache.json'
        for source in context['sources'].values():
            if sha(source['path']) != source['sha256']:
                fail('Source changed during final inspection: ' + source['file'])
        if read_json(source_json) != spec:
            fail('Content JSON changed during final inspection')
        manifest = dict(readback, schema_version=1, case_id=spec['case_id'], project_title=spec['project_title'],
                        input_file=str(source_json.relative_to(root)), input_sha256=sha(source_json),
                        toolkit_sha256=sha(__file__), output_file='draft.docx', output_sha256=sha(final_docx),
                        inspection_file=inspection_file.relative_to(output).as_posix(), inspection_sha256=sha(inspection_file),
                        inspection_cache_file=cache_file.relative_to(output).as_posix(), inspection_cache_sha256=sha(cache_file),
                        content_sha256=sha(output / 'content.json'), sources=spec['sources'])
        dump(output / 'build_manifest.json', manifest)
        return {'case_id': spec['case_id'], 'output': str(output / 'draft.docx'),
                'manifest': str(output / 'build_manifest.json'), 'output_sha256': manifest['output_sha256'],
                'inspection': summary, 'counts': manifest['counts'], 'semantic_review': 'not_run', 'visual_review': 'not_run'}
    finally:
        # Only the exact temp directory created by this invocation is removed.
        shutil.rmtree(staging)


def verify_bundle(case_root, out_dir):
    root = Path(case_root).resolve(strict=True)
    output = bounded_path(root, out_dir, exists=False)
    manifest = read_json(output / 'build_manifest.json')
    for name, key in [('draft.docx', 'output_sha256'), (manifest['inspection_file'], 'inspection_sha256'),
                      (manifest['inspection_cache_file'], 'inspection_cache_sha256'), ('content.json', 'content_sha256')]:
        require(not Path(name).is_absolute() and '..' not in PurePosixPath(name.replace('\\', '/')).parts,
                'bundle member must be a safe relative path')
        path = bounded_path(root, str((output / name).relative_to(root)))
        require(path.is_relative_to(output), 'bundle evidence escaped output directory')
        require(sha(path) == manifest[key], 'bundle hash mismatch: ' + name)
    context = validate(root, read_json(output / 'content.json'))
    for key in ('case_id', 'project_title', 'sources'):
        require(manifest.get(key) == context['spec'][key], 'manifest source/project identity mismatch')
    require(manifest.get('output_file') == 'draft.docx', 'unexpected bundle filename')
    require(manifest.get('semantic_review') == manifest.get('visual_review') == 'not_run', 'builder manifest cannot certify reviews')
    cache_dir = (output / manifest['inspection_cache_file']).parent
    require(cache_dir.is_relative_to(output / 'inspection-cache') and
            output / manifest['inspection_file'] == cache_dir / 'inspection.json', 'unexpected inspection cache location')
    inspect_docx.verified_cache(cache_dir, inspect_docx.cache_identity(output / 'draft.docx', 'draft', True))
    readback = verify_document(context, output / 'draft.docx')
    for key in ('blocks', 'source_objects', 'counts', 'sections'):
        require(readback[key] == manifest[key], 'manifest disagrees with final DOCX: ' + key)
    return dict(readback, output_sha256=sha(output / 'draft.docx'))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    for command in ('inventory', 'build', 'verify'):
        p = sub.add_parser(command)
        p.add_argument('--case-root', required=True)
        if command == 'inventory':
            p.add_argument('--source', required=True)
        elif command == 'build':
            p.add_argument('--input', required=True)
            p.add_argument('--out', required=True)
        else:
            p.add_argument('--out', required=True)
    args = parser.parse_args()
    try:
        if args.command == 'inventory':
            result = inventory(args.case_root, args.source)
        elif args.command == 'build':
            result = build(args.case_root, args.input, args.out)
        else:
            result = verify_bundle(args.case_root, args.out)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except (AuthoringError, OSError, KeyError, TypeError, zipfile.BadZipFile) as exc:
        print(json.dumps({'error': str(exc), 'semantic_review': 'not_run', 'visual_review': 'not_run'}, ensure_ascii=False), file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
