"""Source-bound formula/prose gates; fixtures never claim visual acceptance."""
import os
import json
import copy
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml import parse_xml
import batch_tool as b
import inspect_docx as scan
import output_contract as contract
import parallel_batch as queue
import test_output_contract as base
import test_parallel_batch as fixtures
import formula_uniqueness as unique
import stage_bundle

W=scan.NS['w'];M=scan.NS['m']
SOURCE_HASH='1'*64


def mr(text):
    return '<m:r><m:t>'+text+'</m:t></m:r>'


def omml(text):
    return '<m:oMath>'+mr(text)+'</m:oMath>'


class FormulaProseTests(unittest.TestCase):
    setUp=base.OutputContractTests.setUp
    report=base.OutputContractTests.report

    def with_paragraph(self,content,props=''):
        self.report()
        path=self.root/'scratch.docx';doc=Document(path)
        doc._element.body.insert(len(doc._element.body)-1,parse_xml(
            f'<w:p xmlns:w="{W}" xmlns:m="{M}">{props}{content}</w:p>'))
        doc.save(path);before=b.sha(path)
        report=scan.inspect(path);scan.precheck(report,'draft')
        self.assertEqual(before,b.sha(path))
        return report

    def accept(self,report,findings=(),source_hashes=()):
        return contract.check(report,'一种处理系统',findings,source_hashes=source_hashes,strict_formulas=True)

    def finding(self,report,hit):
        return dict(hit,rule_id='FORMULA-05',classification='source_math_identity',status='closed',
                    artifact_sha256=report['sha256'],source_sha256=SOURCE_HASH,
                    registry_id='F-source-1',source_locator='源文档独立公式1的结构定义',
                    reason='独立复查已核定该位置具有来源数学身份，需要保留原表示。')

    def test_exact_leads_cross_runs_but_not_embedded_substrings(self):
        for punctuation in '，,：:':
            with self.subTest(punctuation=punctuation):
                r=self.with_paragraph('<w:r><w:t>式</w:t></w:r><w:r><w:t>中'+punctuation+'符号说明。</w:t></w:r>')
                hits=r['formula_prose_evidence']['explanation_lead_hits']
                self.assertEqual([h['text'] for h in hits],['式中'+punctuation])
                with self.assertRaisesRegex(b.BatchError,'explanation lead'):self.accept(r)
                f=self.finding(r,hits[0])
                with self.assertRaises(b.BatchError):self.accept(r,[f],[SOURCE_HASH])
        for text in ['其中，符号说明。','在实施方式中读取数据。','公式中的符号表示记录。','式中的S沿用前述定义。']:
            r=self.with_paragraph('<w:r><w:t>'+text+'</w:t></w:r>')
            self.assertFalse(r['formula_prose_evidence']['explanation_lead_hits'])
            self.accept(r)

    def test_clear_standalone_formula_needs_semicolon_inside_or_after_math(self):
        center='<w:pPr><w:jc w:val="center"/></w:pPr>'
        for content,passes in [(omml('x=y+1'),False),(omml('x=y+1；'),True),
                                (omml('x=y+1')+'<w:r><w:t>；</w:t></w:r>',True),
                                (omml('x=y+1')+'<w:r><w:t>;</w:t></w:r>',True)]:
            with self.subTest(content=content):
                r=self.with_paragraph(content,center)
                self.assertEqual(len(r['formula_prose_evidence']['standalone_formulas']),1)
                if passes:self.accept(r)
                else:
                    with self.assertRaisesRegex(b.BatchError,'terminal semicolon'):self.accept(r)
        # A source identity finding cannot waive missing block punctuation.
        r=self.with_paragraph(omml('x=y+1'),center)
        f=self.finding(r,r['formula_prose_evidence']['standalone_formula_semicolon_hits'][0])
        with self.assertRaises(b.BatchError):self.accept(r,[f],[SOURCE_HASH])

    def test_semicolon_inside_implicit_closing_fence_or_script_is_not_terminal(self):
        center='<w:pPr><w:jc w:val="center"/></w:pPr>'
        for tail in ['<m:d><m:e>'+mr('y；')+'</m:e></m:d>',
                     '<m:sSub><m:e>'+mr('y')+'</m:e><m:sub>'+mr('i；')+'</m:sub></m:sSub>']:
            r=self.with_paragraph('<m:oMath>'+mr('x=')+tail+'</m:oMath>',center)
            self.assertTrue(r['formula_prose_evidence']['standalone_formula_semicolon_hits'])
            with self.assertRaisesRegex(b.BatchError,'terminal semicolon'):self.accept(r)

    def test_inline_equality_and_multiple_objects_do_not_establish_complete_formula(self):
        r=self.with_paragraph('<w:r><w:t>当前取值</w:t></w:r>'+omml('x=1')+'<w:r><w:t>时继续处理。</w:t></w:r>')
        self.assertFalse(r['formula_prose_evidence']['standalone_formulas'])
        self.assertFalse(r['formula_prose_evidence']['review_candidates'])
        self.accept(r)
        center='<w:pPr><w:jc w:val="center"/></w:pPr>'
        for content in [omml('x=1')+omml('y=2'),omml('x=1')+'<w:r><w:t>（1）</w:t></w:r>']:
            r=self.with_paragraph(content,center)
            self.assertFalse(r['formula_prose_evidence']['standalone_formula_semicolon_hits'])
            candidates=r['formula_prose_evidence']['review_candidates']
            self.assertTrue(candidates)
            with self.assertRaisesRegex(b.BatchError,'formula/prose review'):self.accept(r)
            findings=[self.finding(r,h) for h in candidates]
            self.accept(r,findings,[SOURCE_HASH])

    def test_multiobject_review_needs_current_source_and_paragraph_path_not_empty_object_id(self):
        r=self.with_paragraph(omml('x=1')+omml('y=2'),'<w:pPr><w:jc w:val="center"/></w:pPr>')
        candidate=r['formula_prose_evidence']['review_candidates'][0]
        self.assertNotIn('object_id',candidate)
        f=self.finding(r,candidate);self.accept(r,[f],[SOURCE_HASH])
        for key,value in [('artifact_sha256','2'*64),('source_sha256','3'*64),('part','word/header1.xml'),
                          ('paragraph',candidate['paragraph']+1),('path','p[999]'),('registry_id','')]:
            with self.subTest(key=key),self.assertRaisesRegex(b.BatchError,'formula/prose review'):
                self.accept(r,[dict(f,**{key:value})],[SOURCE_HASH])
        invalid=dict(f,path='',object_id='')
        with self.assertRaisesRegex(b.BatchError,'formula/prose review'):self.accept(r,[invalid],[SOURCE_HASH])
        # Even a malformed empty object ID in evidence cannot match a global
        # empty-ID finding. A nonempty exact paragraph path remains necessary.
        candidate['object_id']=''
        with self.assertRaisesRegex(b.BatchError,'formula/prose review'):self.accept(r,[invalid],[SOURCE_HASH])

    def test_bare_latin_greek_and_numeric_inline_require_plain_text(self):
        for text,key in [('x','bare_symbol_omml_candidates'),('θ','bare_symbol_omml_candidates'),
                         ('3−0+1=4','inline_numeric_omml_candidates'),('6−1+0=5','inline_numeric_omml_candidates')]:
            with self.subTest(text=text):
                r=self.with_paragraph('<w:r><w:t>本例中</w:t></w:r>'+omml(text)+'<w:r><w:t>表示结果。</w:t></w:r>')
                h=r['formula_prose_evidence'][key][0]
                self.assertTrue(h['path'].endswith('/oMath[1]'));self.assertTrue(h['object_id'].startswith('F'))
                with self.assertRaisesRegex(b.BatchError,'plain text'):self.accept(r)
                # Literal text carrying exactly the same glyphs is acceptable.
                plain=self.with_paragraph('<w:r><w:t>本例中'+text+'表示结果。</w:t></w:r>')
                self.accept(plain)

    def test_source_exception_requires_both_hashes_exact_location_and_registry(self):
        r=self.with_paragraph('<w:r><w:t>定义</w:t></w:r>'+omml('x'))
        h=r['formula_prose_evidence']['bare_symbol_omml_candidates'][0];f=self.finding(r,h)
        self.accept(r,[f],[SOURCE_HASH])
        # Either a precise path or current object ID is a valid locator.
        for locator in ['path','object_id']:
            located=dict(f);located.pop(locator);self.accept(r,[located],[SOURCE_HASH])
        for key,value in [('artifact_sha256','2'*64),('source_sha256','3'*64),('part','word/header1.xml'),
                          ('paragraph',h['paragraph']+1),('registry_id',''),('source_locator',''),
                          ('reason',''),('status','resolved'),('classification','ignore'),('rule_id','WORDING-06')]:
            with self.subTest(key=key),self.assertRaises(b.BatchError):
                self.accept(r,[dict(f,**{key:value})],[SOURCE_HASH])
        with self.assertRaises(b.BatchError):self.accept(r,[dict(f,path='p[999]/oMath[1]',object_id='F9999')],[SOURCE_HASH])
        with self.assertRaises(b.BatchError):self.accept(r,[f],[])
        # A waiver for one object must never cover another identical glyph.
        r2=self.with_paragraph('<w:r><w:t>定义</w:t></w:r>'+omml('x')+omml('x'))
        first=self.finding(r2,r2['formula_prose_evidence']['bare_symbol_omml_candidates'][0])
        with self.assertRaises(b.BatchError):self.accept(r2,[first],[SOURCE_HASH])

    def test_real_scripts_and_styled_unicode_symbols_are_not_flattening_candidates(self):
        structured='<m:oMath><m:sSub><m:e>'+mr('r')+'</m:e><m:sub>'+mr('seen')+'</m:sub></m:sSub>'+mr('(x)')+'</m:oMath>'
        r=self.with_paragraph('<w:r><w:t>读取</w:t></w:r>'+structured+omml('𝒳'))
        e=r['formula_prose_evidence']
        self.assertFalse(e['bare_symbol_omml_candidates']);self.assertFalse(e['inline_numeric_omml_candidates'])
        self.assertIn('sSub',r['formulas'][0]['math_evidence']['structure_elements']);self.accept(r)

    def test_flat_ascii_label_is_a_semantic_candidate_not_a_guessed_label(self):
        for text in ['base','new','seen','live','sin','argmin']:
            with self.subTest(text=text):
                r=self.with_paragraph('<w:r><w:t>定义</w:t></w:r>'+omml(text))
                h=r['formula_prose_evidence']['plain_label_omml_candidates'][0]
                self.assertEqual(h['text'],text)
                with self.assertRaisesRegex(b.BatchError,'plain text'):self.accept(r)
                self.accept(r,[self.finding(r,h)],[SOURCE_HASH])

    def test_incomplete_or_stale_evidence_is_rejected_but_old_policy_unchanged(self):
        r=self.with_paragraph('<w:r><w:t>式中，说明。</w:t></w:r>')
        contract.check(r,'一种处理系统')
        for key,value in [('schema_version',2),('artifact_sha256','2'*64),('extraction_complete',False),('complete',False)]:
            invalid=copy.deepcopy(r);invalid['formula_prose_evidence'][key]=value
            with self.subTest(key=key),self.assertRaisesRegex(b.BatchError,'inspection'):self.accept(invalid)
        r.pop('formula_prose_evidence')
        contract.check(r,'一种处理系统')
        with self.assertRaisesRegex(b.BatchError,'inspection'):self.accept(r)

    def test_complete_equation_identity_preserves_semantics_not_display(self):
        def fp(content):
            return unique.equation_fingerprint(parse_xml(f'<m:oMath xmlns:m="{M}" xmlns:w="{W}">{content}</m:oMath>'))
        plain=fp(mr('x=y+1'))
        self.assertIsNotNone(plain)
        self.assertEqual(plain,fp(mr('x=')+mr('y')+mr('+1；')))
        display='<m:r><w:rPr><w:sz w:val="18"/><w:color w:val="FF0000"/></w:rPr><m:rPr><m:sty m:val="p"/></m:rPr><m:t>x=y+1</m:t></m:r>'
        self.assertEqual(plain,fp(display))
        math_bold='<m:rPr><m:sty m:val="b"/></m:rPr>'
        word_plain='<w:rPr><w:b w:val="0"/></w:rPr>'
        bold=fp('<m:r>'+math_bold+'<m:t>x</m:t></m:r>'+mr('=y+1'))
        for properties in [math_bold+word_plain,word_plain+math_bold]:
            self.assertEqual(bold,fp('<m:r>'+properties+'<m:t>x</m:t></m:r>'+mr('=y+1')))
        self.assertNotEqual(plain,bold)
        for change in [mr('x=y+2'),mr('y+1=x'),mr('x=f(y)'),mr('𝒙=y+1'),
                       '<m:r><m:rPr><m:sty m:val="b"/></m:rPr><m:t>x</m:t></m:r>'+mr('=y+1'),
                       '<m:sSup><m:e>'+mr('x')+'</m:e><m:sup>'+mr('2')+'</m:sup></m:sSup>'+mr('=y+1'),
                       '<m:sSub><m:e>'+mr('x')+'</m:e><m:sub>'+mr('j')+'</m:sub></m:sSub>'+mr('=y+1'),
                       '<m:acc><m:accPr><m:chr m:val="̂"/></m:accPr><m:e>'+mr('x')+'</m:e></m:acc>'+mr('=y+1')]:
            with self.subTest(change=change):self.assertNotEqual(plain,fp(change))
        for nonformula in [mr('x'),mr('x='),mr('3−0+1=4'),
                           '<m:sSub><m:e>'+mr('min')+'</m:e><m:sub>'+mr('x∈A')+'</m:sub></m:sSub>']:
            with self.subTest(nonformula=nonformula):self.assertIsNone(fp(nonformula))
        group='<m:eqArr><m:e>'+mr('x=y')+'</m:e><m:e>'+mr('a=b')+'</m:e></m:eqArr>'
        punctuated=group.replace('x=y','x=y，').replace('a=b','a=b；')
        self.assertEqual(fp(group),fp(punctuated))
        self.assertNotEqual(fp(group),fp(mr('x=y=a=b')))

    def test_duplicate_complete_formula_requires_exact_distinct_scope_evidence(self):
        r=self.with_paragraph('<w:r><w:t>按</w:t></w:r>'+omml('x=y+1')+
                              '<w:r><w:t>计算，随后再次按</w:t></w:r>'+omml('x=y+1；'))
        hits=r['formula_prose_evidence']['duplicate_formula_candidates']
        self.assertEqual(len(hits),1);candidate=hits[0]
        self.assertEqual(candidate['artifact_sha256'],r['sha256'])
        self.assertEqual([p['object_id'] for p in candidate['occurrences']],['F0001','F0002'])
        with self.assertRaisesRegex(b.BatchError,'duplicate complete formula'):self.accept(r)
        generic=[self.finding(r,p) for p in candidate['occurrences']]
        with self.assertRaisesRegex(b.BatchError,'duplicate complete formula'):self.accept(r,generic,[SOURCE_HASH])
        finding=dict(rule_id='FORMULA-01',classification='distinct_formula_scopes',status='closed',
                     artifact_sha256=r['sha256'],fingerprint=candidate['fingerprint'],
                     reason='Independent source review established two different defined quantities.',
                     occurrences=[dict(p,source_sha256=SOURCE_HASH,source_locator='source paragraph '+str(i),
                                       registry_id='F'+str(i),scope_id='scope'+str(i),
                                       scope_definition='source definition '+str(i))
                                  for i,p in enumerate(candidate['occurrences'],1)])
        self.accept(r,[finding],[SOURCE_HASH])
        for key,value in [('artifact_sha256','2'*64),('fingerprint','3'*64),
                          ('classification','source_math_identity'),('reason',''),('status','open')]:
            with self.subTest(key=key),self.assertRaisesRegex(b.BatchError,'duplicate complete formula'):
                self.accept(r,[dict(finding,**{key:value})],[SOURCE_HASH])
        for key,value in [('object_id','F9999'),('source_sha256','2'*64),('source_locator',''),
                          ('scope_definition',''),('scope_id',''),('registry_id','')]:
            invalid=copy.deepcopy(finding);invalid['occurrences'][1][key]=value
            with self.subTest(key=key),self.assertRaisesRegex(b.BatchError,'duplicate complete formula'):
                self.accept(r,[invalid],[SOURCE_HASH])
        for key in ['source_locator','scope_id','scope_definition']:
            invalid=copy.deepcopy(finding);invalid['occurrences'][1][key]=invalid['occurrences'][0][key]
            with self.subTest(key=key),self.assertRaisesRegex(b.BatchError,'duplicate complete formula'):
                self.accept(r,[invalid],[SOURCE_HASH])
        invalid=copy.deepcopy(finding);invalid['occurrences'].pop()
        with self.assertRaisesRegex(b.BatchError,'duplicate complete formula'):self.accept(r,[invalid],[SOURCE_HASH])
        clean=self.with_paragraph('<w:r><w:t>按</w:t></w:r>'+omml('x=y+1')+'<w:r><w:t>计算，后续使用上述关系。</w:t></w:r>')
        self.assertFalse(clean['formula_prose_evidence']['duplicate_formula_candidates']);self.accept(clean)
        clean['formula_prose_evidence'].pop('duplicate_formula_candidates')
        with self.assertRaisesRegex(b.BatchError,'inspection'):self.accept(clean)

    def test_duplicate_scan_covers_body_table_header_but_excludes_deleted_content(self):
        r=self.with_paragraph('<w:r><w:t>按</w:t></w:r>'+omml('x=y+1'))
        path=self.root/'scratch.docx';doc=Document(path)
        for paragraph in [doc.add_table(rows=1,cols=1).cell(0,0).paragraphs[0],
                          doc.sections[0].header.paragraphs[0]]:
            paragraph._p.append(parse_xml(f'<m:oMath xmlns:m="{M}">'+mr('x=y+1')+'</m:oMath>'))
        doc.add_paragraph()._p.append(parse_xml(f'<w:del xmlns:w="{W}" xmlns:m="{M}">'+omml('x=y+1')+'</w:del>'))
        doc.save(path);before=b.sha(path);r=scan.inspect(path)
        self.assertEqual(before,b.sha(path))
        hits=r['formula_prose_evidence']['duplicate_formula_candidates'];self.assertEqual(len(hits),1)
        self.assertEqual(len(hits[0]['occurrences']),3)
        self.assertEqual({h['part'] for h in hits[0]['occurrences']},{'word/document.xml','word/header1.xml'})
        self.assertIn('formula_uniqueness.py',scan.extractor_dependencies())
        identity={'extractor_sha256':scan.file_sha(Path(scan.__file__)),
                  'extractor_dependencies':scan.extractor_dependencies()}
        self.assertTrue(scan.extractor_unchanged(identity))
        changed=self.root/'changed-helper.py';changed.write_text('# different helper bytes',encoding='utf-8')
        with patch.object(unique,'__file__',str(changed)):
            self.assertFalse(scan.extractor_unchanged(identity))
        self.assertIn(Path(unique.__file__).resolve(),stage_bundle.selected_files(Path(scan.__file__).resolve().parents[1]))

    def test_policy_requires_frozen_output_contract_and_known_formula_mode(self):
        p=copy.deepcopy(fixtures.POLICY);queue.validate_policy(p)
        for mode in [contract.FORMULA_MODE,'unknown']:
            with self.subTest(mode=mode),self.assertRaises(b.BatchError):queue.validate_policy(dict(p,formula_contract=mode))
        p.update(worker_preparation='frozen_stage_packet_v1',output_contract=contract.MODE,
                 delivery={'folder_prefix':'专利撰写完成','layout':'flat_final_files','timing':'after_each_acceptance'},
                 formula_contract=contract.FORMULA_MODE)
        queue.validate_policy(p)

    def test_verify_reads_frozen_flag_and_reuses_existing_inspection(self):
        r=self.with_paragraph('<w:r><w:t>式中，说明。</w:t></w:r>')
        report_path=self.root/'inspection.json';b.atomic(report_path,r)
        handoff_path=self.root/'handoff.json';b.atomic(handoff_path,{'input_inspection':{'inspection':str(report_path)}})
        attempt={'prepared_handoff':{'path':str(handoff_path)},'input_artifact_sha256':r['sha256']}
        policy={'output_contract':contract.MODE}
        verified={'handoff':{'input_inspection':{'inspection':str(report_path)}},'input_report':r}
        with patch('prepare_handoff.verify_prepared',return_value=verified),patch.object(scan,'inspect',side_effect=AssertionError('second scan')):
            old=contract.verify({'policy':policy},{'title':'一种处理系统'},attempt,{'findings':[]})
            self.assertNotIn('formula_contract',old)
            self.assertEqual(old['scope'],'Objective title, located wording and enabled numbering/page-flow checks only; original six reviews and final-page acceptance remain required')
            with self.assertRaisesRegex(b.BatchError,'explanation lead'):
                contract.verify({'policy':dict(policy,formula_contract=contract.FORMULA_MODE)},
                                {'title':'一种处理系统'},attempt,{'findings':[]})

    def test_user_source_documents_detection_is_read_only(self):
        fixture_path=os.environ.get('PATENT_FORMULA_REFERENCE_FIXTURES')
        if not fixture_path:self.skipTest('local reference fixture manifest not configured')
        fixture_rows=json.loads(Path(fixture_path).read_text(encoding='utf8'))
        sources=[(Path(r['path']),r['leads'],r['missing'],r['blocks']) for r in fixture_rows]
        for row in fixture_rows:self.assertEqual(b.sha(Path(row['path'])),row['sha256'])
        for path,leads,missing,blocks in sources:
            with self.subTest(source=path.name):
                before=b.sha(path);r=scan.inspect(path);e=r['formula_prose_evidence']
                self.assertEqual(before,b.sha(path));self.assertEqual(e['artifact_sha256'],before)
                self.assertEqual([h['paragraph'] for h in e['explanation_lead_hits']],leads)
                self.assertEqual([h['paragraph'] for h in e['standalone_formula_semicolon_hits']],missing)
                self.assertEqual(len(e['standalone_formulas']),blocks)
                if missing:
                    self.assertEqual([h['text'] for h in e['inline_numeric_omml_candidates']],['6−1+0=5','3−0+1=4'])
                    self.assertEqual([h['text'] for h in e['plain_label_omml_candidates']],['base','new','seen','live'])
                    self.assertTrue(any(h['text']=='x' and h['paragraph']==143 for h in e['bare_symbol_omml_candidates']))


if __name__=='__main__':unittest.main()
