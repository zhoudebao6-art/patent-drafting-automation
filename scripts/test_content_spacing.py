"""Synthetic spacing evidence/gates only; no source cases or visual claims."""
import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from docx import Document
from docx.oxml import parse_xml
import batch_tool as b
import inspect_docx as scan
import output_contract as contract
import prepare_handoff as handoff

W, M, A = (scan.NS[key] for key in ('w', 'm', 'a'))
C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
TITLE = '一种记录处理方法'


class ContentSpacingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root/'scratch.docx'

    def document(self, text='保存对应记录。'):
        doc = Document()
        for line in ['说明书摘要', '摘要内容。', '权利要求书',
                     '1.一种记录处理方法，其特征在于，保存记录。', TITLE,
                     '技术领域', '本发明涉及记录处理。', '具体实施方式', text]:
            doc.add_paragraph(line)
        return doc

    def inspect(self, doc=None, extras=None):
        (doc or self.document()).save(self.path)
        if extras:
            with zipfile.ZipFile(self.path) as archive:
                files = {name: archive.read(name) for name in archive.namelist()}
            files.update(extras)
            with zipfile.ZipFile(self.path, 'w') as archive:
                for name, content in files.items():
                    archive.writestr(name, content)
        before = b.sha(self.path)
        report = scan.inspect(self.path)
        scan.precheck(report, 'draft')
        self.assertEqual(before, b.sha(self.path))
        return report

    def xml_paragraph(self, doc, contents):
        node = parse_xml(f'<w:p xmlns:w="{W}" xmlns:m="{M}" xmlns:a="{A}">{contents}</w:p>')
        doc._element.body.insert(len(doc._element.body)-1, node)

    def accept(self, report, **kwargs):
        return contract.check(report, TITLE, strict_spacing=True, **kwargs)

    def test_every_unicode_whitespace_plus_zero_width_padding_has_exact_locator(self):
        chars = ' \t\r\n\u0085\u00a0\u1680\u2000\u2001\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200a\u2028\u2029\u202f\u205f\u3000\u200b\ufeff'
        for char in chars:
            with self.subTest(codepoint=f'U+{ord(char):04X}'):
                root = ET.Element(f'{{{W}}}document')
                body = ET.SubElement(root, f'{{{W}}}body')
                p = ET.SubElement(body, f'{{{W}}}p')
                r = ET.SubElement(p, f'{{{W}}}r')
                ET.SubElement(r, f'{{{W}}}t').text = '甲'+char+'乙'
                hits = scan.content_spacing_hits(root, 'word/document.xml')
                self.assertEqual(len(hits), 1)
                self.assertEqual((hits[0]['offset'], hits[0]['text'], hits[0]['paragraph']), (1, char, 1))
                self.assertEqual(hits[0]['codepoint'], f'U+{ord(char):04X}')
                self.assertTrue(hits[0]['path'].endswith('p[1]/r[1]/t[1]'))
        self.assertFalse(scan.literal_content_whitespace('\u2061'))
        self.assertFalse(scan.literal_content_whitespace('\u200c'))

    def test_cross_run_spaces_are_not_lost_and_clean_runs_are_not_concatenated(self):
        doc = self.document()
        p = doc.add_paragraph()
        for text in ['甲', ' ', '乙 ', ' 丙', '丁']:
            p.add_run(text)
        report = self.inspect(doc)
        hits = report['content_spacing_evidence']['hits']
        self.assertEqual([h['offset'] for h in hits], [0, 1, 0])
        self.assertEqual(len({h['path'] for h in hits}), 3)
        self.assertEqual(len({h['paragraph'] for h in hits}), 1)
        self.assertTrue(any(i['code']=='CONTENT_LITERAL_WHITESPACE' for i in report['issues']))
        with self.assertRaisesRegex(b.BatchError, 'content spacing contract'):
            self.accept(report)
        self.assertEqual(''.join(r.text for r in p.runs), '甲 乙  丙丁')

    def test_headers_footers_tables_nested_textbox_math_and_drawing_are_scanned_once(self):
        doc = self.document()
        doc.sections[0].header.paragraphs[0].text = '页 眉'
        doc.sections[0].footer.paragraphs[0].text = '页\u00a0脚'
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '表\u3000格'
        self.xml_paragraph(doc, '<m:oMath><m:r><m:t>x\u2009y</m:t></m:r></m:oMath>')
        self.xml_paragraph(doc, '<w:r><w:drawing><a:p><a:r><a:t>图\u200b字</a:t></a:r></a:p></w:drawing></w:r>')
        self.xml_paragraph(doc, '<w:r><w:pict><w:txbxContent><w:p><w:r><w:t>框\ufeff内</w:t></w:r></w:p></w:txbxContent></w:pict></w:r>')
        report = self.inspect(doc)
        hits = report['content_spacing_evidence']['hits']
        self.assertEqual(len(hits), 6)
        self.assertEqual(len({(h['part'], h['path'], h['offset']) for h in hits}), 6)
        self.assertEqual({h['part'] for h in hits}, {'word/document.xml', 'word/header1.xml', 'word/footer1.xml'})
        self.assertTrue(all(h['paragraph']>0 for h in hits))
        self.assertEqual(sum(h['representation']=='math' for h in hits), 1)
        self.assertEqual(sum(h['representation']=='drawing_text' for h in hits), 1)

    def test_chart_diagram_and_vml_text_but_not_chart_formula_are_covered(self):
        extras = {
            'word/charts/chart99.xml': f'<c:chartSpace xmlns:c="{C}"><c:tx><c:v>Label A</c:v></c:tx><c:f> Sheet 1 ! A1 </c:f></c:chartSpace>',
            'word/charts/chart100.xml': '<cx:chartSpace xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex"><cx:v>Modern Text</cx:v><cx:f> Sheet 2 ! A1 </cx:f></cx:chartSpace>',
            'word/diagrams/data99.xml': f'<a:txBody xmlns:a="{A}"><a:p><a:r><a:t>Diagram\u202fText</a:t></a:r></a:p></a:txBody>',
            'word/drawings/text99.vml': f'<v:shape xmlns:v="{scan.NS["v"]}"><v:textpath string="VML Text"/></v:shape>',
        }
        report = self.inspect(extras=extras)
        hits = report['content_spacing_evidence']['hits']
        self.assertEqual(len(hits), 4)
        self.assertEqual({h['part'] for h in hits}, set(extras))
        self.assertEqual({h['representation'] for h in hits}, {'chart_text', 'drawing_text', 'vml_text'})
        self.assertTrue(all(h['paragraph'] is None for h in hits))

    def test_removed_revisions_field_instructions_and_style_dimensions_are_excluded(self):
        doc = self.document()
        self.xml_paragraph(doc, '<w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs><w:spacing w:before="90"/></w:pPr>'
            '<w:r><w:rPr><w:spacing w:val="90"/></w:rPr><w:t>普通文字</w:t></w:r>'
            '<w:del><w:r><w:t>已 删除</w:t><w:tab/></w:r><w:p><w:r><w:t>删除段 落</w:t></w:r></w:p></w:del>'
            '<w:moveFrom><w:r><w:t>原 位置</w:t></w:r></w:moveFrom>'
            '<w:r><w:instrText xml:space="preserve"> PAGE \\* MERGEFORMAT </w:instrText></w:r>'
            '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
            '<m:oMath><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f><m:r><m:t>x\u2061y</m:t></m:r></m:oMath>')
        report = self.inspect(doc)
        self.assertFalse(report['content_spacing_evidence']['hits'])
        self.accept(report)

    def test_actual_content_tab_current_insertions_and_field_results_are_not_excluded(self):
        doc = self.document()
        self.xml_paragraph(doc, '<w:pPr><w:tabs><w:tab w:val="left" w:pos="720"/></w:tabs></w:pPr>'
            '<w:r><w:tab/></w:r><w:ins><w:r><w:t>新 文字</w:t></w:r></w:ins>'
            '<w:moveTo><w:r><w:t>新\u00a0位置</w:t></w:r></w:moveTo>'
            '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1 2</w:t></w:r></w:fldSimple>')
        report = self.inspect(doc)
        hits = report['content_spacing_evidence']['hits']
        self.assertEqual(len(hits), 4)
        self.assertEqual([h['codepoint'] for h in hits], ['U+0009', 'U+0020', 'U+00A0', 'U+0020'])
        self.assertEqual(hits[0]['representation'], 'text_tab')

    def test_complete_current_hash_bound_empty_evidence_passes(self):
        report = self.inspect()
        evidence = report['content_spacing_evidence']
        self.assertEqual(evidence['schema_version'], 1)
        self.assertTrue(evidence['complete'] and evidence['extraction_complete'])
        self.assertEqual(evidence['artifact_sha256'], report['sha256'])
        self.assertFalse(evidence['hits'])
        self.assertEqual(self.accept(report)['spacing_contract'], contract.SPACING_MODE)

    def test_missing_incomplete_or_forged_evidence_binding_cannot_pass(self):
        report = self.inspect()
        mutations = [lambda r: r.pop('content_spacing_evidence')]
        for key, value in [('schema_version', True), ('schema_version', 2), ('complete', False),
                           ('complete', 1), ('extraction_complete', False), ('hits', None),
                           ('hits', {}), ('scope', ''), ('artifact_sha256', 'f'*64)]:
            mutations.append(lambda r, k=key, v=value: r['content_spacing_evidence'].__setitem__(k, v))
        mutations.extend([lambda r: r.__setitem__('sha256', 'a'*64), lambda r: r.__setitem__('sha256', '')])
        for index, mutate in enumerate(mutations):
            with self.subTest(mutation=index):
                modified = copy.deepcopy(report)
                mutate(modified)
                with self.assertRaisesRegex(b.BatchError, 'content spacing inspection'):
                    self.accept(modified)
        modified = copy.deepcopy(report)
        modified['content_spacing_evidence']['hits'] = ['forged']
        with self.assertRaisesRegex(b.BatchError, 'content spacing contract'):
            self.accept(modified)

    def test_plain_source_formula_findings_never_waive_literal_spaces(self):
        for mathematical in (False, True):
            with self.subTest(mathematical=mathematical):
                doc = self.document('使用x y计算。' if not mathematical else '使用下式计算。')
                if mathematical:
                    self.xml_paragraph(doc, '<m:oMath><m:r><m:t>x y</m:t></m:r></m:oMath>')
                report = self.inspect(doc)
                source = 'e'*64
                findings = [dict(h, rule_id='WORDING-06', classification='source_formula', status='closed',
                                 artifact_sha256=report['sha256'], source_sha256=source,
                                 source_locator='源公式', registry_id='F1', reason='来源存在空格')
                            for h in report['content_spacing_evidence']['hits']]
                with self.assertRaisesRegex(b.BatchError, 'content spacing contract'):
                    self.accept(report, findings=findings, source_hashes=[source])

    def test_legacy_checks_do_not_require_or_enable_spacing_evidence(self):
        report = self.inspect(self.document('保留 原有记录。'))
        report.pop('content_spacing_evidence')
        self.assertNotIn('spacing_contract', contract.check(report, TITLE))
        with patch('prepare_handoff.verify_prepared', side_effect=AssertionError('legacy should not inspect')):
            self.assertIsNone(contract.verify({'policy': {}}, {}, {}, {}))

    def test_verify_uses_frozen_policy_and_existing_sealed_report_without_docx_reparse(self):
        report = self.inspect(self.document('现有 空格。'))
        verified = {'handoff': {'input_inspection': {'inspection': 'sealed.json'}}, 'input_report': report}
        attempt = {'input_artifact_sha256': report['sha256']}
        with patch('prepare_handoff.verify_prepared', return_value=verified), patch.object(scan, 'inspect', side_effect=AssertionError('duplicate ZIP parse')):
            legacy = contract.verify({'policy': {'output_contract': contract.MODE}}, {'title': TITLE}, attempt, {})
            self.assertNotIn('spacing_contract', legacy)
            with self.assertRaisesRegex(b.BatchError, 'content spacing contract'):
                contract.verify({'policy': {'output_contract': contract.MODE, 'spacing_contract': contract.SPACING_MODE}}, {'title': TITLE}, attempt, {})
            with self.assertRaisesRegex(b.BatchError, 'another artifact'):
                contract.verify({'policy': {'output_contract': contract.MODE}}, {'title': TITLE}, {'input_artifact_sha256': 'f'*64}, {})

    def test_spacing_policy_requires_frozen_output_contract_and_rejects_unknown_modes(self):
        import parallel_batch as queue
        from test_parallel_batch import POLICY
        legacy = copy.deepcopy(POLICY)
        queue.validate_policy(legacy)
        with self.assertRaisesRegex(b.BatchError, 'spacing contract requires'):
            queue.validate_policy(dict(legacy, spacing_contract=contract.SPACING_MODE))
        policy = dict(legacy, worker_preparation='frozen_stage_packet_v1',
                      output_contract=contract.MODE, spacing_contract=contract.SPACING_MODE,
                      delivery={'folder_prefix': '专利撰写完成', 'layout': 'flat_final_files', 'timing': 'after_each_acceptance'})
        queue.validate_policy(policy)
        for mode in ['unknown', '', True, {}]:
            with self.subTest(mode=mode), self.assertRaisesRegex(b.BatchError, 'Invalid content spacing contract'):
                queue.validate_policy(dict(policy, spacing_contract=mode))
        with self.assertRaisesRegex(b.BatchError, 'frozen preparation'):
            queue.validate_policy(dict(policy, worker_preparation=None))

    def test_sealed_inspection_rejects_deleted_or_cleared_hit_report_bytes(self):
        self.document('保留 空格。').save(self.path)
        cache = self.root/'cache'
        summary = scan.run_inspection(self.path, cache, kind='draft', extract_media=True, reuse=True)
        metadata = Path(summary['output_directory'])/'cache.json'
        summary['cache_evidence_sha256'] = scan.file_sha(metadata)
        report_path = Path(summary['inspection'])
        original = report_path.read_bytes()
        digest = b.sha(self.path)
        accepted = handoff.verify_inspection(summary, self.path, digest, 'draft', [cache])
        self.assertTrue(accepted['content_spacing_evidence']['hits'])
        for mutation in ['delete', 'clear', 'rebind']:
            with self.subTest(mutation=mutation):
                data = json.loads(original)
                if mutation == 'delete':
                    data.pop('content_spacing_evidence')
                elif mutation == 'clear':
                    data['content_spacing_evidence']['hits'] = []
                else:
                    data['content_spacing_evidence']['artifact_sha256'] = 'a'*64
                report_path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
                with self.assertRaises(ValueError):
                    handoff.verify_inspection(summary, self.path, digest, 'draft', [cache])
                report_path.write_bytes(original)


if __name__ == '__main__':
    unittest.main()
