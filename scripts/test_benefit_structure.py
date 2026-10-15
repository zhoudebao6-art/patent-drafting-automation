"""Source-safe effect numbering; no model calls or rendered-page claims."""
from copy import deepcopy
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml import parse_xml
from lxml import etree as ET
from PIL import Image

import authoring_toolkit as toolkit
import batch_tool as batch
import inspect_docx as scan
import output_contract as contract
import parallel_batch as queue
import test_authoring_toolkit as authoring_fixtures
import test_parallel_batch as queue_fixtures


TITLE = '一种记录处理方法'
INTRO = '与现有技术相比，本发明的有益效果在于：'
W = scan.NS['w']
M = scan.NS['m']
DELIMITED_MATH = f'''<m:oMath xmlns:m="{M}"><m:d><m:e>
<m:r><m:t>x</m:t></m:r></m:e></m:d><m:r><m:t>=</m:t></m:r>
<m:r><m:t>y</m:t></m:r></m:oMath>'''


def paragraph(block_id, text, role='body'):
    return {'id': block_id, 'type': 'paragraph', 'role': role, 'text': text}


def effect_document(effects=('1.保留记录，可供回查。', '2.检查状态，限制无效输入。'),
                    *, intro=INTRO, end='附图说明'):
    doc = Document()
    doc.add_paragraph('说明书摘要')
    doc.add_paragraph('摘要中的1.数字不属于效果编号。')
    doc.add_paragraph('权利要求书')
    doc.add_paragraph('1.一种记录处理方法，其特征在于，保存记录。')
    doc.add_paragraph('2.根据权利要求1所述的方法，其特征在于，检查记录。')
    doc.add_paragraph(TITLE)
    doc.add_paragraph('技术领域')
    doc.add_paragraph('本发明涉及记录处理。')
    doc.add_paragraph('背景技术')
    doc.add_paragraph('背景中的有益效果字样不是本次编号区域。')
    doc.add_paragraph('发明内容')
    doc.add_paragraph('一种记录处理方法，包括读取和保存记录。')
    if intro is not None:
        doc.add_paragraph(intro)
    for text in effects:
        p = doc.add_paragraph()
        # Deliberately split labels and words across runs.
        for char in text:
            p.add_run(char)
    if end is not None:
        doc.add_paragraph(end)
    if end != '具体实施方式':
        doc.add_paragraph('具体实施方式')
    doc.add_paragraph('1.实施方式中的编号不进入效果组。')
    return doc


class BenefitInspectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='benefit-结构-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def inspect(self, doc):
        path = self.root / '审计.docx'
        doc.save(path)
        before = batch.sha(path)
        report = scan.inspect(path)
        scan.precheck(report, 'draft')
        self.assertEqual(batch.sha(path), before)
        return report

    def check(self, report):
        return contract.check(report, TITLE, strict_benefits=True)

    def test_one_and_multiple_effects_are_consecutive_and_other_sections_are_outside_scope(self):
        for effects in [('1.保留记录。',), ('1.保留记录。', '2.检查状态。', '3.保存索引。')]:
            with self.subTest(effects=effects):
                report = self.inspect(effect_document(effects))
                evidence = report['benefit_numbering_evidence']
                self.assertTrue(evidence['scope_resolved'])
                self.assertFalse(evidence['issues'])
                self.assertEqual([p['number'] for p in evidence['items']], list(range(1, len(effects)+1)))
                self.assertEqual([p['text'] for p in evidence['items']], list(effects))
                self.assertEqual(self.check(report)['benefit_contract'], contract.BENEFIT_MODE)

    def test_noncanonical_missing_duplicate_skipped_empty_and_spaced_ordinals_fail(self):
        cases = {
            'missing': ['保留记录。'], 'missing_second': ['1.保留记录。', '检查状态。'],
            'duplicate': ['1.保留记录。', '1.检查状态。'], 'skip': ['1.保留记录。', '3.检查状态。'],
            'starts_at_two': ['2.保留记录。'], 'zero': ['0.保留记录。'],
            'parenthesis': ['（1）保留记录。'], 'ascii_parenthesis': ['(1)保留记录。'],
            'ideographic_comma': ['1、保留记录。'], 'wide_dot': ['1．保留记录。'],
            'space': ['1. 保留记录。'], 'tab': ['1.\t保留记录。'], 'empty': [],
            'no_prose': ['1.'],
        }
        for name, effects in cases.items():
            with self.subTest(case=name):
                report = self.inspect(effect_document(effects))
                self.assertTrue(report['benefit_numbering_evidence']['issues'])
                with self.assertRaisesRegex(batch.BatchError, 'benefit numbering'):
                    self.check(report)

    def test_scope_requires_one_invention_heading_one_intro_and_an_end(self):
        for mutation in ['no_heading', 'two_headings', 'no_intro', 'two_intros', 'no_end']:
            with self.subTest(mutation=mutation):
                doc = effect_document()
                if mutation == 'no_heading':
                    next(p for p in doc.paragraphs if p.text == '发明内容').text = '内容'
                elif mutation == 'two_headings':
                    next(p for p in doc.paragraphs if p.text == '发明内容').insert_paragraph_before('发明内容')
                elif mutation == 'no_intro':
                    next(p for p in doc.paragraphs if p.text == INTRO).text = '效果说明'
                elif mutation == 'two_intros':
                    next(p for p in doc.paragraphs if p.text == INTRO).insert_paragraph_before(INTRO)
                else:
                    for p in doc.paragraphs:
                        if p.text in ['附图说明', '具体实施方式']:
                            p.text = '普通文字'
                with self.assertRaisesRegex(batch.BatchError, 'benefit numbering'):
                    self.check(self.inspect(doc))

    def test_embodiments_can_close_scope_without_figure_description(self):
        report = self.inspect(effect_document(end='具体实施方式'))
        self.assertEqual([p['number'] for p in report['benefit_numbering_evidence']['items']], [1, 2])
        self.check(report)

    def test_native_list_label_is_not_accepted_as_a_second_numbering_system(self):
        doc = effect_document(('1.保留记录。',))
        effect = next(p for p in doc.paragraphs if p.text == '1.保留记录。')
        effect.style = doc.styles['List Number']
        report = self.inspect(doc)
        self.assertIn('native_list_label_requires_literal_conversion',
                      {x['code'] for x in report['benefit_numbering_evidence']['issues']})
        with self.assertRaises(batch.BatchError):
            self.check(report)

    def test_inline_and_standalone_source_equations_keep_tree_and_ordinal_scope(self):
        doc = effect_document()
        first = next(p for p in doc.paragraphs if p.text == '1.保留记录，可供回查。')
        inline = parse_xml(DELIMITED_MATH)
        first._p.append(inline)
        after = doc.add_paragraph()
        after._p.append(parse_xml(DELIMITED_MATH))
        first._p.addnext(after._p)
        expected = [toolkit.tree_signature(x) for x in doc._element.findall('.//m:oMath', toolkit.NS)]
        report = self.inspect(doc)
        self.check(report)
        evidence = report['benefit_numbering_evidence']
        self.assertEqual([p['number'] for p in evidence['items']], [1, 2])
        self.assertEqual(len(evidence['formula_paragraphs']), 1)
        actual = Document(self.root/'审计.docx')._element.findall('.//m:oMath', toolkit.NS)
        self.assertEqual([toolkit.tree_signature(x) for x in actual], expected)
        self.assertFalse(report['punctuation_evidence']['hits'])
        self.assertTrue(report['punctuation_evidence']['formula_protected_candidates'])

    def test_formula_cannot_create_a_first_effect_without_prose(self):
        doc = effect_document()
        intro = next(p for p in doc.paragraphs if p.text == INTRO)
        formula = doc.add_paragraph()
        formula._p.append(parse_xml(DELIMITED_MATH))
        intro._p.addnext(formula._p)
        report = self.inspect(doc)
        self.assertIn('formula_without_effect', {x['code'] for x in report['benefit_numbering_evidence']['issues']})
        with self.assertRaises(batch.BatchError):
            self.check(report)

    def test_table_effect_is_flagged_for_review_instead_of_silently_skipped(self):
        doc = effect_document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = '3.嵌套效果段。'
        # The final 2. effect, rather than the 2. claim, is the intended anchor.
        anchor = next(p for p in doc.paragraphs if p.text == '2.检查状态，限制无效输入。')
        anchor._p.addnext(table._tbl)
        report = self.inspect(doc)
        self.assertIn('effect_container_requires_rewrite_or_review',
                      {x['code'] for x in report['benefit_numbering_evidence']['issues']})
        with self.assertRaises(batch.BatchError):
            self.check(report)

    def test_old_contract_does_not_require_new_numbering_evidence(self):
        report = self.inspect(effect_document(('尚未编号的历史效果段。',)))
        report.pop('benefit_numbering_evidence')
        result = contract.check(report, TITLE)
        self.assertNotIn('benefit_contract', result)
        with self.assertRaises(batch.BatchError):
            self.check(report)

    def formula_paragraph_finding(self, report, text, digest='a'*64):
        evidence = report['benefit_numbering_evidence']
        row = next(p for p in evidence['items']+evidence['issues'] if p.get('text') == text)
        return dict(row, rule_id='STYLE-07', classification='source_formula_paragraph', status='closed',
                    artifact_sha256=report['sha256'], source_sha256=digest,
                    source_locator='源公式段1', registry_id='F1', reason='原式独立段及源公式号，非普通效果文字。')

    def test_plain_formula_paragraph_is_source_bound_and_remaining_effect_numbers_are_rechecked(self):
        formula = 'x=(y+1) (1)'
        report = self.inspect(effect_document(['1.保留记录。', formula, '2.检查状态。']))
        with self.assertRaises(batch.BatchError):
            self.check(report)
        finding = self.formula_paragraph_finding(report, formula)
        self.assertEqual(finding['path'], f'paragraph[{finding["paragraph"]}]/visible_text')
        contract.check(report, TITLE, [finding], source_hashes=['a'*64], strict_benefits=True)
        # The ordinary prose punctuation contract is separate from whole-formula
        # paragraph scope, and needs its original character-bound evidence.
        with self.assertRaisesRegex(batch.BatchError, 'quotation/bracket'):
            contract.check(report, TITLE, [finding], strict_prose=True,
                           source_hashes=['a'*64], strict_benefits=True)
        punctuation = [dict(h, rule_id='WORDING-06', classification='source_formula', status='closed',
                            artifact_sha256=report['sha256'], source_sha256='a'*64,
                            source_locator='源公式段1', registry_id='F1', reason='原公式函数括号或原公式号。')
                       for h in report['punctuation_evidence']['hits']]
        contract.check(report, TITLE, [finding]+punctuation, strict_prose=True,
                       source_hashes=['a'*64], strict_benefits=True)
        bad_report = self.inspect(effect_document(['1.保留记录。', formula, '3.检查状态。']))
        bad_finding = self.formula_paragraph_finding(bad_report, formula)
        with self.assertRaisesRegex(batch.BatchError, 'benefit numbering'):
            contract.check(bad_report, TITLE, [bad_finding], source_hashes=['a'*64], strict_benefits=True)

    def test_formula_whole_paragraph_waiver_rejects_wrong_source_hash_locator_or_classification(self):
        text = 'x=y+1'
        report = self.inspect(effect_document(['1.保留记录。', text, '2.检查状态。']))
        finding = self.formula_paragraph_finding(report, text)
        mutations = {'artifact_sha256':'b'*64, 'source_sha256':'b'*64, 'part':'word/header1.xml',
                     'paragraph':finding['paragraph']+1, 'path':'paragraph[999]/visible_text', 'offset':1,
                     'text':text+'。', 'source_locator':'', 'registry_id':'', 'reason':'',
                     'classification':'source_formula', 'rule_id':'WORDING-06', 'status':'unresolved'}
        for field, value in mutations.items():
            with self.subTest(field=field), self.assertRaises(batch.BatchError):
                contract.check(report, TITLE, [dict(finding, **{field:value})],
                               source_hashes=['a'*64], strict_benefits=True)
        for hashes in ([], ['b'*64]):
            with self.subTest(source_hashes=hashes), self.assertRaises(batch.BatchError):
                contract.check(report, TITLE, [finding], source_hashes=hashes, strict_benefits=True)

    def test_equation_number_outside_omml_is_preserved_as_original_source_text(self):
        doc = effect_document()
        first = next(p for p in doc.paragraphs if p.text == '1.保留记录，可供回查。')
        formula = doc.add_paragraph()
        formula._p.append(parse_xml(DELIMITED_MATH))
        formula.add_run('(1)')
        first._p.addnext(formula._p)
        before = toolkit.tree_signature(formula._p.find('m:oMath', toolkit.NS))
        report = self.inspect(doc)
        text = next(p['text'] for p in report['benefit_numbering_evidence']['items'] if '(1)' in p['text'])
        finding = self.formula_paragraph_finding(report, text)
        contract.check(report, TITLE, [finding], source_hashes=['a'*64], strict_benefits=True)
        saved = Document(self.root/'审计.docx')
        containing = next(p for p in saved.paragraphs if p.text == '(1)')
        self.assertEqual(toolkit.tree_signature(containing._p.find('m:oMath', toolkit.NS)), before)
        self.assertEqual(containing.text, '(1)')

    def test_image_formula_is_preserved_and_cannot_hide_an_unrelated_native_list_issue(self):
        doc = effect_document()
        first = next(p for p in doc.paragraphs if p.text == '1.保留记录，可供回查。')
        image = io.BytesIO(); Image.new('RGB', (30, 12), 'white').save(image, format='PNG'); image.seek(0)
        formula = doc.add_paragraph(); formula.add_run().add_picture(image)
        first._p.addnext(formula._p)
        report = self.inspect(doc)
        finding = self.formula_paragraph_finding(report, '')
        contract.check(report, TITLE, [finding], source_hashes=['a'*64], strict_benefits=True)
        original_media = report['media']
        second = next(p for p in doc.paragraphs if p.text == '2.检查状态，限制无效输入。')
        second.style = doc.styles['List Number']
        bad = self.inspect(doc)
        self.assertEqual(bad['media'], original_media)
        bad_finding = self.formula_paragraph_finding(bad, '')
        with self.assertRaises(batch.BatchError):
            contract.check(bad, TITLE, [bad_finding], source_hashes=['a'*64], strict_benefits=True)

    def test_only_formula_paragraphs_cannot_be_accepted_as_an_empty_effect_group(self):
        text = 'x=y+1'
        report = self.inspect(effect_document([text]))
        finding = self.formula_paragraph_finding(report, text)
        with self.assertRaises(batch.BatchError):
            contract.check(report, TITLE, [finding], source_hashes=['a'*64], strict_benefits=True)

    def test_inline_formula_does_not_remove_ordinary_effect_prose_number_requirement(self):
        doc = effect_document(['通过保留关系，可供回查。', '2.检查状态。'])
        effect = next(p for p in doc.paragraphs if p.text == '通过保留关系，可供回查。')
        effect._p.append(parse_xml(DELIMITED_MATH))
        report = self.inspect(doc)
        self.assertEqual(len(report['benefit_numbering_evidence']['items']), 2)
        self.assertFalse(report['benefit_numbering_evidence']['formula_paragraphs'])
        row = report['benefit_numbering_evidence']['items'][0]
        character_finding = dict(row, rule_id='WORDING-06', classification='source_formula', status='closed',
                                 artifact_sha256=report['sha256'], source_sha256='a'*64,
                                 source_locator='原公式', registry_id='F1', reason='只确认行内数学对象。')
        with self.assertRaisesRegex(batch.BatchError, 'benefit numbering'):
            contract.check(report, TITLE, [character_finding], source_hashes=['a'*64], strict_benefits=True)

    def test_ole_formula_carrier_requires_a_whole_source_bound_paragraph_finding(self):
        # This is an OOXML object-reference fixture, not a rendered OLE equation.
        doc = effect_document()
        first = next(p for p in doc.paragraphs if p.text == '1.保留记录，可供回查。')
        formula = doc.add_paragraph()
        formula._p.append(parse_xml(f'''<w:r xmlns:w="{W}"
            xmlns:o="urn:schemas-microsoft-com:office:office"
            xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
            <w:object><o:OLEObject r:id="rIdFixtureEquation" ProgID="Equation.3"/></w:object></w:r>'''))
        first._p.addnext(formula._p)
        report = self.inspect(doc)
        self.assertIn('effect_container_requires_rewrite_or_review',
                      {x['code'] for x in report['benefit_numbering_evidence']['issues']})
        with self.assertRaises(batch.BatchError):
            self.check(report)
        finding = self.formula_paragraph_finding(report, '')
        contract.check(report, TITLE, [finding], source_hashes=['a'*64], strict_benefits=True)
        with zipfile.ZipFile(self.root/'审计.docx') as archive:
            self.assertIn(b'rIdFixtureEquation', archive.read('word/document.xml'))


class BenefitAuthoringTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='benefit-author-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.spec = authoring_fixtures.minimal_spec()
        self.spec['sections'][3]['blocks'] += [
            paragraph('invention', '发明内容', 'heading'),
            paragraph('scheme', '本发明读取并保存已有记录。'),
            paragraph('benefit_intro', INTRO),
            paragraph('effect1', '保留记录，可供回查。', 'benefit'),
            {'id': 'effect2', 'type': 'paragraph', 'role': 'benefit', 'runs': [
                {'type': 'text', 'text': '检查'}, {'type': 'text', 'text': '状态，限制无效输入。'}]},
            paragraph('figures_heading', '附图说明', 'heading'),
            paragraph('embodiments_heading', '具体实施方式', 'heading'),
            paragraph('example', '1.此编号属于实施步骤。')]

    def build(self):
        path = self.root/'内容.json'
        path.write_text(json.dumps(self.spec, ensure_ascii=False), encoding='utf-8')
        return toolkit.build(self.root, path.name, 'output')

    def block(self, block_id):
        return next(b for section in self.spec['sections'] for b in section['blocks'] if b['id'] == block_id)

    def test_benefit_role_is_idempotent_without_mutating_spec_or_other_numbers(self):
        before = deepcopy(self.spec)
        one = toolkit.validate(self.root, self.spec)
        two = toolkit.validate(self.root, self.spec)
        self.assertEqual(self.spec, before)
        self.assertEqual(one['block_texts'], two['block_texts'])
        self.assertEqual(one['block_texts']['effect1'], '1.保留记录，可供回查。')
        self.assertEqual(one['block_texts']['effect2'], '2.检查状态，限制无效输入。')
        self.assertEqual(one['block_texts']['example'], '1.此编号属于实施步骤。')
        self.assertEqual(one['block_texts']['claim1'], before['sections'][2]['blocks'][0]['text'])
        self.assertTrue(all(b['role'] == 'body' for b in one['blocks'] if b['id'].startswith('effect')))
        self.build()
        self.assertEqual(json.loads((self.root/'output'/'content.json').read_text(encoding='utf-8')), before)

    def test_text_and_runs_output_use_ordinary_body_style(self):
        self.build()
        doc = Document(self.root/'output'/'draft.docx')
        effects = [p for p in doc.paragraphs if p.text in ['1.保留记录，可供回查。', '2.检查状态，限制无效输入。']]
        self.assertEqual(len(effects), 2)
        for p in effects:
            props = p._p.find('w:pPr', toolkit.NS)
            self.assertIsNone(props.find('w:numPr', toolkit.NS))
            self.assertEqual(props.find('w:ind', toolkit.NS).get(toolkit.qn('w:firstLineChars')), '200')
            self.assertEqual(props.find('w:ind', toolkit.NS).get(toolkit.qn('w:firstLine')), '560')
            self.assertEqual(props.find('w:jc', toolkit.NS).get(toolkit.qn('w:val')), 'left')
            for run in p.runs:
                self.assertEqual(run.font.name, '宋体')
                self.assertEqual(run.font.size.pt, 14)
                self.assertFalse(run.bold)
                self.assertFalse(run.italic)
        report = scan.inspect(self.root/'output'/'draft.docx')
        scan.precheck(report, 'draft')
        contract.check(report, self.spec['project_title'], strict_benefits=True)

    def test_benefit_role_rejects_wrong_region_and_supplied_prefixes(self):
        for mutation in ['before_intro', 'after_invention', 'in_abstract', 'intro_outside_invention',
                         '1.已有编号。', '2、已有编号。', '（1）已有编号。', '(1)已有编号。',
                         '一、已有编号。', ' 前置空格。']:
            with self.subTest(mutation=mutation):
                spec = deepcopy(self.spec)
                blocks = spec['sections'][3]['blocks']
                target = next(b for b in blocks if b['id'] == 'effect1')
                if mutation == 'before_intro':
                    blocks.remove(target); blocks.insert(0, target)
                elif mutation == 'after_invention':
                    blocks.remove(target); blocks.append(target)
                elif mutation == 'in_abstract':
                    blocks.remove(target); spec['sections'][0]['blocks'].append(target)
                elif mutation == 'intro_outside_invention':
                    next(b for b in blocks if b['id'] == 'invention')['text'] = '背景技术'
                else:
                    target['text'] = mutation
                with self.assertRaises(toolkit.AuthoringError):
                    toolkit.validate(self.root, spec)

    def test_literal_body_ordinals_can_precede_generated_roles_without_double_prefix(self):
        effect = self.block('effect1')
        effect['role'] = 'body'; effect['text'] = '1.'+effect['text']
        context = toolkit.validate(self.root, self.spec)
        self.assertEqual(context['block_texts']['effect2'], '2.检查状态，限制无效输入。')
        for text in ['未编号。', '2.起号错误。', '1. 编号后有空格。']:
            effect['text'] = text
            with self.subTest(text=text), self.assertRaises(toolkit.AuthoringError):
                toolkit.validate(self.root, self.spec)

    def test_generated_prefix_preserves_inline_and_standalone_registered_math(self):
        source = self.root/'公式来源.docx'
        doc = Document()
        doc.add_paragraph()._p.append(ET.fromstring(authoring_fixtures.MATH))
        doc.add_paragraph()._p.append(ET.fromstring(DELIMITED_MATH))
        doc.save(source)
        inventory = toolkit.inventory(self.root, source.name)
        self.spec['sources'].append(inventory['source'])
        rows = []
        for raw in inventory['omml_objects']:
            row = dict(raw); row.pop('linear_text_for_navigation_only')
            row['category'] = 'main_formula'; rows.append(row)
        self.spec['formulas'] += rows
        self.spec['formula_policy']['source_main_count'] = 2
        effect = self.block('effect1')
        effect.pop('text')
        effect['runs'] = [{'type': 'text', 'text': '保留'}, {'type': 'formula', 'formula_id': rows[0]['id']},
                          {'type': 'text', 'text': '关系，可供回查。'}]
        blocks = self.spec['sections'][3]['blocks']
        position = blocks.index(effect)+1
        blocks.insert(position, {'id':'effect_equation','type':'formula','formula_id':rows[1]['id']})
        before_source = toolkit.sha(source)
        before_spec = deepcopy(self.spec)
        result = self.build()
        self.assertEqual(toolkit.sha(source), before_source)
        self.assertEqual(self.spec, before_spec)
        with zipfile.ZipFile(result['output']) as archive:
            root = ET.fromstring(archive.read('word/document.xml'))
        actual = root.findall('.//m:oMath', toolkit.NS)
        self.assertEqual([toolkit.tree_signature(x) for x in actual], [r['tree_sha256'] for r in rows])
        report = scan.inspect(Path(result['output'])); scan.precheck(report, 'draft')
        contract.check(report, self.spec['project_title'], strict_benefits=True, strict_prose=True)
        evidence = report['benefit_numbering_evidence']
        self.assertEqual([p['number'] for p in evidence['items']], [1, 2])
        self.assertEqual(len(evidence['formula_paragraphs']), 1)


class BenefitQueueGateTests(unittest.TestCase):
    """Real frozen cache and queue transitions; dispatch/render are fixtures."""
    reserve = queue_fixtures.ParallelTests.reserve
    receipt = queue_fixtures.ParallelTests.receipt

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='benefit-gate-')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.src = self.root/'input'; self.src.mkdir()
        doc = Document(); doc.add_paragraph('来源记录'); doc.save(self.src/'source.docx')

    def init_batch(self, benefit_mode=True):
        policy = deepcopy(queue_fixtures.POLICY)
        policy.update(worker_preparation='frozen_stage_packet_v1', output_contract=contract.MODE,
                      delivery={'folder_prefix':'专利撰写完成','layout':'flat_final_files','timing':'after_each_acceptance'})
        if benefit_mode:
            policy['benefit_contract'] = contract.BENEFIT_MODE
        policy_path = self.root/'policy.json'; batch.atomic(policy_path, policy)
        selection = {'projects':[{'title':TITLE,'source_files':['source.docx'],'primary_source':'source.docx'}]}
        with patch.object(queue, 'POLICY', policy_path):
            self.manifest = queue.init(self.src, self.root/'work', selection, self.root/'desktop')
        self.batch = self.manifest.parent

    def review_attempt(self, effects=('1.保留记录。', '2.检查状态。')):
        draft = self.reserve(1); receipt = self.receipt(draft)
        effect_document(effects).save(receipt['artifact'])
        receipt['artifact_sha256'] = batch.sha(receipt['artifact'])
        queue.record(self.manifest, draft['case_id'], draft['attempt']['id'], receipt)
        review = self.reserve(1)
        return review, self.receipt(review, 'pass')

    def test_acceptance_and_delivery_reuse_the_hash_sealed_inspection(self):
        self.init_batch()
        review, receipt = self.review_attempt()
        with patch.object(scan, 'inspect', side_effect=AssertionError('duplicate DOCX parse')):
            queue.record(self.manifest, review['case_id'], review['attempt']['id'], receipt)
            destination = Path(queue.deliver(self.manifest))
        self.assertEqual([p.name for p in destination.glob('*.docx')], [TITLE+'.docx'])
        case = batch.read(self.manifest)['projects'][0]
        self.assertEqual(case['status'], 'delivered')
        self.assertEqual(case['receipts'][-1]['document_contract']['benefit_contract'], contract.BENEFIT_MODE)

    def test_incorrect_ordinals_cannot_be_accepted_or_delivered(self):
        self.init_batch()
        review, receipt = self.review_attempt(('1.保留记录。', '3.检查状态。'))
        with patch.object(scan, 'inspect', side_effect=AssertionError('duplicate DOCX parse')):
            with self.assertRaisesRegex(batch.BatchError, 'benefit numbering'):
                queue.record(self.manifest, review['case_id'], review['attempt']['id'], receipt)
        self.assertNotEqual(batch.read(self.manifest)['projects'][0]['status'], 'accepted')
        with self.assertRaises(batch.BatchError):
            queue.deliver(self.manifest)
        self.assertFalse(list((self.root/'desktop').rglob('*.docx')))

    def test_modified_sealed_benefit_report_cannot_forge_acceptance(self):
        self.init_batch()
        review, receipt = self.review_attempt()
        prepared = batch.read(review['attempt']['prepared_handoff']['path'])
        path = Path(prepared['input_inspection']['inspection'])
        report = batch.read(path)
        report['benefit_numbering_evidence']['items'][1]['number'] = 7
        batch.atomic(path, report)
        with self.assertRaises(batch.BatchError):
            queue.record(self.manifest, review['case_id'], review['attempt']['id'], receipt)

    def test_delivery_rechecks_saved_benefit_contract(self):
        self.init_batch()
        review, receipt = self.review_attempt()
        queue.record(self.manifest, review['case_id'], review['attempt']['id'], receipt)
        manifest = batch.read(self.manifest)
        manifest['projects'][0]['receipts'][-1]['document_contract'].pop('benefit_contract')
        batch.atomic(self.manifest, manifest)
        with patch.object(scan, 'inspect', side_effect=AssertionError('duplicate DOCX parse')):
            with self.assertRaises(batch.BatchError):
                queue.deliver(self.manifest)
        self.assertFalse(list((self.root/'desktop').rglob('*.docx')))

    def test_legacy_policy_keeps_acceptance_compatible(self):
        self.init_batch(benefit_mode=False)
        review, receipt = self.review_attempt(('历史未编号效果段。',))
        with patch.object(scan, 'inspect', side_effect=AssertionError('duplicate DOCX parse')):
            queue.record(self.manifest, review['case_id'], review['attempt']['id'], receipt)
            destination = Path(queue.deliver(self.manifest))
        self.assertTrue((destination/(TITLE+'.docx')).is_file())
        case = batch.read(self.manifest)['projects'][0]
        self.assertNotIn('benefit_contract', case['receipts'][-1]['document_contract'])

    def test_benefit_policy_requires_supported_mode_and_final_output_contract(self):
        for mode, output in [('unknown', contract.MODE), (contract.BENEFIT_MODE, None)]:
            policy = deepcopy(queue_fixtures.POLICY)
            policy['benefit_contract'] = mode
            if output:
                policy.update(output_contract=output, worker_preparation='frozen_stage_packet_v1',
                              delivery={'folder_prefix':'专利撰写完成','layout':'flat_final_files','timing':'after_each_acceptance'})
            with self.subTest(mode=mode, output=output), self.assertRaises(batch.BatchError):
                queue.validate_policy(policy)


if __name__ == '__main__':
    unittest.main()
